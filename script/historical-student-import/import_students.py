#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Local-only historical teaching enrollment with class/planner pair matching."""
import argparse
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal
import hashlib
import json
import os
from pathlib import Path
import re
import secrets
import subprocess
import sys

import openpyxl
import pymysql

ROOT = Path(__file__).resolve().parents[2]
POLICY = 'historical-teaching-20261026-v5'
LEGACY_POLICY = 'historical-teaching-20261026-v1'
PRIOR_POLICIES = (LEGACY_POLICY, 'historical-teaching-20261026-v2', 'historical-teaching-20261026-v3', 'historical-teaching-20261026-v4')
CONFIRMED_SOURCE_ROWS = {(0, 11), (0, 25), (0, 28), (0, 95)}
SOURCE_SHA256 = '47fac89e8f785b178f3782daa1803b62ffbde287448d65bdab93f48e5d21baaa'
NOTE = '历史学员导入，未记录'
TENANT = 1
SHEETS = ('注册营养师-学员导入填写', '学会公共营养师-学员导入填写', '人社公共营养师-学员导入填写')
REVIEW = ({31, 45, 80, 127}, {30, 72, 58, 128, 185, 188, 131}, {7, 57, 28})
MISSING_SKU = '公共营养师｜人社部技能等级｜一级｜超值精品班'
TABLES = ('zsjos_product', 'zsjos_product_category', 'zsjos_product_sku', 'zsjos_product_attr',
          'zsjos_product_attr_value', 'zsjos_delivery_class', 'system_users', 'system_role', 'system_user_role',
          'zsjos_person', 'zsjos_person_contact_claim', 'zsjos_person_no_daily_counter',
          'zsjos_service_relation', 'zsjos_collaboration_group', 'zsjos_student_contact_record', 'zsjos_business_event',
          'system_dept', 'zsjos_order')
CATALOG_TABLES = (*TABLES[:9], 'system_dept')


def check(ok, message):
    if not ok:
        raise ValueError(message)


def json_default(value):
    if isinstance(value, (date, datetime, Decimal)):
        return str(value)
    if isinstance(value, bytes):
        return value.hex()
    raise TypeError(type(value).__name__)


def packed(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'), default=json_default)


def digest(value):
    result = hashlib.sha256()
    encoder = json.JSONEncoder(ensure_ascii=False, sort_keys=True, separators=(',', ':'), default=json_default)
    for chunk in encoder.iterencode(value):
        result.update(chunk.encode('utf-8'))
    return result.hexdigest()


def save_private(path, value):
    path = Path(path).resolve()
    check(not path.is_relative_to(ROOT), 'Private evidence must be outside the repository')
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('x', encoding='utf-8') as out:
        json.dump(value, out, ensure_ascii=False, sort_keys=True, separators=(',', ':'), default=json_default)
    return path


def connect(database, container='yudao-mysql'):
    marker = Path('/etc/zsjos/agent-environment')
    environment = os.environ.get('ZSJOS_AGENT_ENV') or (marker.read_text().strip() if marker.exists() else 'local')
    check(environment in ('local', 'test', 'production'), 'Invalid ZSJOS_AGENT_ENV value')
    check(environment == 'local', 'Only the explicitly authorized local environment is supported')
    # This adapter intentionally permits only the verified local development container.
    check(container == 'yudao-mysql', 'Only local development container is supported')
    check(database == 'ruoyi-vue-pro' or re.fullmatch(r'historical_student_verify_\d+', database), 'Unapproved local database')
    meta = json.loads(subprocess.check_output(['docker', 'inspect', container]))[0]
    env = dict(x.split('=', 1) for x in meta['Config']['Env'])
    password = env.get('MYSQL_ROOT_PASSWORD')
    if not password and env.get('MYSQL_ROOT_PASSWORD_FILE'):
        password = subprocess.check_output(['docker', 'exec', container, 'cat', env['MYSQL_ROOT_PASSWORD_FILE']]).decode().strip()
    check(bool(password), 'Existing container authentication unavailable')
    conn = pymysql.connect(host='127.0.0.1', port=3306, user='root', password=password,
                           database=database, charset='utf8mb4', autocommit=False,
                           cursorclass=pymysql.cursors.DictCursor)
    with conn.cursor() as q:
        q.execute('SET NAMES utf8mb4')
        q.execute("SET time_zone='+08:00'")
        q.execute('SET SESSION TRANSACTION ISOLATION LEVEL READ COMMITTED')
        q.execute('SET SESSION innodb_lock_wait_timeout=10')
    return conn


def text_sha(value):
    return hashlib.sha256(value.encode('utf-8')).hexdigest()


def read_source(path, confirmations_path=None):
    path = Path(path)
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    check(sha == SOURCE_SHA256, 'Source workbook changed; row-specific decisions require a new review')
    wb = openpyxl.load_workbook(path, data_only=False)
    check(tuple(wb.sheetnames[:3]) == SHEETS, 'Unexpected source sheets; policy review required')
    records, excluded, duplicates, seen = [], [], [], {}
    for index, name in enumerate(SHEETS):
        sheet = wb[name]
        check(sheet.max_column == 44, 'Source column structure changed')
        check([sheet.cell(1, c).value for c in (1, 2, 4, 5, 6, 7)] ==
              ['姓名 *', '手机号码 †', '课程名称 *', '课程规格名称 t', '班级名称 *', '学习规划师姓名 *'], 'Header changed')
        for n, row in enumerate(sheet.iter_rows(min_row=2), 2):
            values = [c.value for c in row]
            if not any(x is not None for x in values):
                continue
            check(not any(c.data_type == 'f' for c in row), 'Formula in source data')
            check(all(v is None for v in values[7:]) and values[2] is None, 'Additional source fields need review')
            clean = [str(v).strip() if v is not None else '' for v in values[:7]]
            clean[6] = {'周老师': '周老师【梁颖】', '周老师【梁颖】【梁颖】': '周老师【梁颖】'}.get(clean[6], clean[6])
            check(re.fullmatch(r'1[3-9]\d{9}', clean[1]) is not None, f'Invalid contact: {index + 1}:{n}')
            record = dict(sheet=name, sheetIndex=index, row=n, values=values, name=clean[0], mobile=clean[1],
                          course=clean[3], sku=clean[4], className=clean[5], sourcePlanner=clean[6])
            check(all(clean[i] for i in (0, 3, 4, 5, 6)), f'Missing source field: {index + 1}:{n}')
            reason = None
            if n in REVIEW[index]:
                reason = '姓名/联系方式待确认' if (index, n) in {(0, 31), (0, 45), (0, 80), (0, 127), (1, 131), (2, 28)} else '同手机号规格或规划师冲突，待确认'
            elif '代理' in clean[4]:
                reason = '代理信息不是课程规格，实际班型待确认'
            elif clean[4] == MISSING_SKU:
                reason = '规格未配置，按确认不新增'
            if reason:
                excluded.append({**record, 'reason': reason})
                continue
            key = tuple(clean)
            if key in seen:
                duplicates.append({**record, 'duplicateOf': seen[key]})
                continue
            seen[key] = f'{name}:{n}'
            record['key'] = f'hist-student:{sha[:24]}:{index + 1}:{n}'
            record['fingerprint'] = digest({'policy': POLICY, 'source': clean})
            record['legacyFingerprint'] = digest({'policy': LEGACY_POLICY, 'source': clean})
            record['priorFingerprints'] = [digest({'policy': policy, 'source': clean}) for policy in PRIOR_POLICIES]
            records.append(record)
    wb.close()
    check(len(records) + len(excluded) + len(duplicates) == 525, 'Unexpected source population; policy review required')
    result = dict(policy=POLICY, source=str(path.resolve()), sourceSha256=sha, candidates=records,
                  exceptions=excluded, duplicates=duplicates)
    if confirmations_path is not None:
        confirmation_path = Path(confirmations_path).resolve()
        check(not confirmation_path.is_relative_to(ROOT), 'Private confirmations must be outside repository')
        raw = confirmation_path.read_bytes()
        config = json.loads(raw.decode('utf-8'))
        check(config.get('tenantId') == TENANT and config.get('sourceSha256') == sha, 'Confirmation source/tenant changed')
        seen = set()
        for entry in config['decisions']:
            key = (entry['sheetIndex'], entry['row'])
            check(key in CONFIRMED_SOURCE_ROWS and key not in seen, 'Unexpected or duplicate identity confirmation')
            seen.add(key)
            record = next((r for r in records if (r['sheetIndex'], r['row']) == key), None)
            check(record is not None and entry['sourceNameHash'] == text_sha(record['name']) and
                  entry['sourceContactHash'] == text_sha(record['mobile']) and
                  entry['contactField'] in ('mobile', 'wechat_id') and entry.get('confirmed') is True,
                  'Identity confirmation no longer matches the source row')
            record['identityConfirmation'] = dict(entry, confirmationFileSha256=hashlib.sha256(raw).hexdigest())
        check(seen == CONFIRMED_SOURCE_ROWS, 'All four reviewed identity decisions are required')
        result['identityConfirmationsPath'] = str(confirmation_path)
        result['identityConfirmationsSha256'] = hashlib.sha256(raw).hexdigest()
    return result


def fetch(q, sql, args=()):
    q.execute(sql, args or None)
    return q.fetchall()


def one(q, sql, args=()):
    data = fetch(q, sql, args)
    check(len(data) == 1, 'Expected one authoritative configuration record')
    return data[0]


def catalog(q, lock=False):
    suffix = ' FOR UPDATE' if lock else ''
    return {table: fetch(q, f'SELECT * FROM `{table}` WHERE tenant_id=1 AND deleted=0 ORDER BY id' + suffix)
            for table in (*TABLES[:8], 'system_dept')} | {
                'system_user_role': fetch(q, 'SELECT * FROM system_user_role WHERE deleted=0 ORDER BY id' + suffix)}


def resolve_class(c, record, allow_unscheduled=False):
    # Exact authoritative display names are required; no fuzzy nickname/user-ID guesses.
    owners = [u for u in c['system_users'] if u['nickname'] == record['sourcePlanner']]
    check(len(owners) == 1 and owners[0]['status'] == 0, '表中规划师无法唯一匹配有效账号')
    owner = owners[0]
    roles = {r['id'] for r in c['system_role'] if r['code'] == 'study_planner' and r['status'] == 0}
    check(any(r['user_id'] == owner['id'] and r['role_id'] in roles for r in c['system_user_role']), '表中规划师无有效班主任角色资格')
    departments = [d for d in c['system_dept'] if d['id'] == owner['dept_id'] and d['status'] == 0]
    check(len(departments) == 1, '表中规划师所属部门无效')
    classes = [v for v in c['zsjos_delivery_class'] if v['class_name'] == record['className'] and v['homeroom_user_id'] == owner['id']]
    check(len(classes) <= 1, '同名且同班主任的班级有多个，待确认')
    key = digest({'className': record['className'], 'homeroomUserId': owner['id']})
    if classes:
        cls = classes[0]
        allowed_kinds = ('NORMAL', 'UNSCHEDULED') if allow_unscheduled else ('NORMAL',)
        check(cls['status'] == 'SERVING' and cls['class_kind'] in allowed_kinds, '对应班级不可入班')
        return dict(classKey=key, classId=cls['id'], ownerId=owner['id'], classCreate=None)
    check(len(record['className']) <= 100, '班级名称超过长度限制')
    return dict(classKey=key, classId=None, ownerId=owner['id'], classCreate=dict(
        classKey=key, className=record['className'], ownerId=owner['id'], ownerName=owner['nickname'],
        deptId=departments[0]['id'], deptName=departments[0]['name']))


def create_class(q, spec, actor, now):
    # Match the established BJ+timestamp+four-digit number format without a new counter table.
    for _ in range(20):
        number = 'BJ' + now.strftime('%Y%m%d%H%M%S') + f'{secrets.randbelow(10000):04d}'
        if not fetch(q, 'SELECT id FROM zsjos_delivery_class WHERE tenant_id=1 AND class_no=%s', (number,)):
            break
    else:
        raise ValueError('班级编号冲突，请重新预检')
    cid = insert(q, 'zsjos_delivery_class', dict(class_no=number, class_name=spec['className'],
                 homeroom_user_id=spec['ownerId'], homeroom_user_name_snapshot=spec['ownerName'],
                 dept_id=spec['deptId'], dept_name_snapshot=spec['deptName'], system_class=0,
                 class_kind='NORMAL', status='SERVING', version=0))
    insert(q, 'zsjos_business_event', dict(event_type='student_historical_class_created', aggregate_type='delivery-class',
           aggregate_id=cid, operator_user_id=actor, occurred_at=now,
           idempotency_key='hist-class:' + spec['classKey'], related_object_refs=packed(dict(spec, classId=cid, source=POLICY))))
    return cid


def snapshot(c, sku):
    product = next(p for p in c['zsjos_product'] if p['id'] == sku['spu_id'])
    check(product['status'] == 0, 'Disabled course')
    categories = {v['id']: v for v in c['zsjos_product_category']}
    path, parent = [], product['category_id']
    while parent:
        check(parent in categories and len(path) < 10 and parent not in [v['id'] for v in path], 'Invalid category chain')
        node = categories[parent]
        check(node['status'] == 0, 'Disabled category')
        path.insert(0, {'id': node['id'], 'name': node['name']})
        parent = node['parent_id']
    check(bool(path), 'Missing category')
    attrs = json.loads(sku['attr_values_json'])
    definitions = [a for a in c['zsjos_product_attr'] if a['spu_id'] == product['id']]
    check(set(attrs) <= {a['attr_key'] for a in definitions}, 'Unknown specification attribute')
    specs = []
    for a in sorted(definitions, key=lambda a: (a['sort'], a['id'])):
        val = attrs.get(a['attr_key'])
        check(a['status'] == 0 and (not int.from_bytes(a['required'], 'big') or val is not None), 'Inactive/missing specification attribute')
        if val is None:
            continue
        matches = [v for v in c['zsjos_product_attr_value'] if v['attr_id'] == a['id'] and v['value'] == val and v['status'] == 0]
        check(len(matches) == 1, 'Invalid specification value')
        specs.append({'attrKey': a['attr_key'], 'attrName': a['attr_name'], 'value': val, 'label': matches[0]['label'], 'labelMissing': False})
    second = path[1] if len(path) > 1 else {'id': None, 'name': None}
    return dict(productRef=product['product_ref'], name=product['name'], categoryId=path[-1]['id'],
                categoryName=path[-1]['name'], categoryPath=path, level1CategoryId=path[0]['id'],
                level1CategoryName=path[0]['name'], level2CategoryId=second['id'], level2CategoryName=second['name'],
                skuRef=sku['sku_ref'], skuName=sku['sku_name'], selectedAttrValuesJson=sku['attr_values_json'],
                price=sku['price'], spuUnknown=False, skuUnknown=False, specs=specs)


def person_name_matches(stored, requested):
    return name_match_kind(stored, requested) is not None


def name_match_kind(stored, requested):
    stored = stored.strip()
    if stored == requested:
        return 'exact'
    # Preserve the full stored name; recognize only bounded annotation syntax.
    normalized = stored.translate(str.maketrans({'（': '(', '）': ')'}))
    match = re.fullmatch(re.escape(requested) + r'\(([^()]+)\)', normalized)
    if match:
        note = match.group(1).strip()
        if note == '老学员':
            return 'historical_student_suffix'
        if note in ('复购', '代理'):
            return 'historical_name_note'
        if len(note) <= 30 and re.fullmatch(r'[\u4e00-\u9fffA-Za-z0-9· ]{1,28}推荐', note):
            return 'historical_referral_note'
    return None


def identity_match(q, person, record):
    if person['mobile'] == record['mobile']:
        contact_field = 'mobile'
    elif person['wechat_id'] == record['mobile']:
        # Same unified contact value already belongs to this unique identity. Do not move fields.
        contact_field = 'wechat_id'
    else:
        return None
    kind = name_match_kind(person['name'], record['name'])
    evidence = []
    if kind is None and record.get('identityConfirmation'):
        decision = record['identityConfirmation']
        if (decision['contactField'] != contact_field or decision['storedNameHash'] != text_sha(person['name'])
                or decision['sourceNameHash'] != text_sha(record['name']) or decision['sourceContactHash'] != text_sha(record['mobile'])):
            return None
        kind = 'user_confirmed_identity'
        evidence = [dict(type='user-confirmation', id=decision['decisionId'], fileSha256=decision['confirmationFileSha256'])]
    if kind is None:
        services = fetch(q, 'SELECT id FROM zsjos_service_relation WHERE tenant_id=1 AND deleted=0 '
                         'AND person_id=%s AND student_name_snapshot=%s ORDER BY id', (person['id'], record['name']))
        orders = fetch(q, "SELECT id FROM zsjos_order WHERE tenant_id=1 AND deleted=0 AND person_id=%s "
                       "AND status='effective' AND student_name=%s AND student_mobile=%s ORDER BY id",
                       (person['id'], record['name'], record['mobile']))
        # SQL collation equality is not a name-normalization policy; compare exact snapshot text too.
        services = [s for s in services if one(q, 'SELECT student_name_snapshot value FROM zsjos_service_relation WHERE tenant_id=1 AND id=%s', (s['id'],))['value'] == record['name']]
        orders = [o for o in orders if one(q, 'SELECT student_name value FROM zsjos_order WHERE tenant_id=1 AND id=%s', (o['id'],))['value'] == record['name']]
        evidence = [dict(type='student-service', id=s['id']) for s in services] + [dict(type='effective-order', id=o['id']) for o in orders]
        if not evidence:
            return None
        kind = 'persisted_student_name'
    return dict(kind=kind, contactField=contact_field, evidence=evidence)


def person_match(q, record):
    values = fetch(q, 'SELECT id,name,mobile,wechat_id,identity_status,version FROM zsjos_person '
                   'WHERE tenant_id=1 AND deleted=0 AND (mobile=%s OR wechat_id=%s)', (record['mobile'], record['mobile']))
    claims = fetch(q, 'SELECT person_id,reservation_key,deleted+0 deleted FROM zsjos_person_contact_claim '
                   'WHERE tenant_id=1 AND contact_value=%s', (record['mobile'],))
    ids = {v['id'] for v in values} | {v['person_id'] for v in claims if v['person_id'] is not None}
    if len(ids) > 1 or claims and (claims[0]['deleted'] or claims[0]['reservation_key']):
        return None, '目标库联系方式存在多档案/占用冲突'
    if ids:
        person = fetch(q, 'SELECT id,name,mobile,wechat_id,identity_status,version FROM zsjos_person '
                       'WHERE tenant_id=1 AND deleted=0 AND id=%s', (next(iter(ids)),))
        if len(person) != 1:
            return None, '目标库姓名或联系方式与表格不一致，保留原档案待确认'
        decision = identity_match(q, person[0], record)
        if decision is None:
            return None, '联系方式命中但姓名无可核实的同人依据，待确认'
        person[0]['matchEvidence'] = decision
        return person[0], None
    return None, None


def prepare(conn, source, actor_username, lock=False):
    check(bool(source.get('identityConfirmationsSha256') or source.get('historicalSkuAuthorization')), 'This policy requires the private reviewed identity confirmation file or scoped historical SKU authorization')
    with conn.cursor() as q:
        target = one(q, 'SELECT DATABASE() db,@@server_uuid serverUuid,VERSION() version')
        c = catalog(q, lock)
        actors = [u for u in c['system_users'] if u['username'] == actor_username and u['status'] == 0]
        check(len(actors) == 1, 'Import audit operator must resolve to one enabled local System account')
        actor = actors[0]['id']
        ready, excluded, replay, existing_rows = [], list(source['exceptions']), [], []
        for r in source['candidates']:
            previous = fetch(q, 'SELECT event_type,aggregate_id,related_object_refs FROM zsjos_business_event WHERE tenant_id=1 AND idempotency_key=%s AND deleted=0', (r['key'],))
            if previous:
                check(len(previous) == 1 and previous[0]['event_type'] == 'student_historical_import' and
                      json.loads(previous[0]['related_object_refs'])['fingerprint'] in (r['fingerprint'], *r['priorFingerprints']), 'Import key content conflict')
                # v1 owner choices are a preserved historical decision, never silently migrated.
                replay.append({**r, 'serviceId': previous[0]['aggregate_id']})
                continue
            try:
                if source.get('historicalSkuAuthorization'):
                    from import_disabled_skus import authorized_sku
                    skus = [authorized_sku(c, r, source)]
                else:
                    skus = [s for s in c['zsjos_product_sku'] if s['sku_name'] == r['sku'] and s['status'] == 0]
                check(len(skus) == 1, '课程规格未启用或名称不能唯一匹配')
                snap = snapshot(c, skus[0])
                if source.get('historicalSkuAuthorization'):
                    snap.update(historicalSku=True, skuDisabledAtImport=True, catalogPriceIsPlaceholder=True)
                selected = resolve_class(c, r, allow_unscheduled=bool(source.get('historicalSkuAuthorization')))
                if source.get('historicalSkuAuthorization'):
                    check(selected['classId'] is not None, '本批历史规格导入要求复用已有班级，缺失班级需先复核')
                person, error = person_match(q, r)
                check(error is None, error)
                if person and selected['classId'] is not None:
                    existing = fetch(q, 'SELECT id,owner_user_id,student_name_snapshot FROM zsjos_service_relation WHERE tenant_id=1 AND deleted=0 AND person_id=%s '
                                     "AND class_id=%s AND JSON_UNQUOTE(JSON_EXTRACT(service_snapshot,'$.skuRef'))=%s",
                                     (person['id'], selected['classId'], skus[0]['sku_ref']))
                    if existing:
                        check(len(existing) == 1, '目标库同班同规格服务有多条，待确认')
                        check(existing[0]['owner_user_id'] == selected['ownerId'] and existing[0]['student_name_snapshot'] == r['name'], '已有服务姓名或负责人不同，保留原服务待确认')
                        existing_rows.append({**r, 'personId': person['id'], 'serviceId': existing[0]['id'],
                                              'reason': '已有同人同班同规格且姓名负责人一致的服务，无需新增'})
                        continue
                ready.append({**r, 'personId': person['id'] if person else None, 'personState': person,
                              **selected, 'snapshot': snap,
                              'correctedCourse': snap['name']})
            except ValueError as e:
                excluded.append({**r, 'reason': str(e)})
        creates = {r['classKey']: r['classCreate'] for r in ready if r['classCreate'] is not None}
        return dict(policy=POLICY, target=target, sourceSha256=source['sourceSha256'], source=source['source'],
                    identityConfirmationsSha256=source.get('identityConfirmationsSha256'),
                    historicalSkuAuthorization=source.get('historicalSkuAuthorization'),
                    actorUsername=actor_username, actorId=actor, configHash=digest(c), ready=ready,
                    classCreates=[creates[k] for k in sorted(creates)], exceptions=excluded, duplicates=source['duplicates'], replay=replay, existing=existing_rows)


def insert(q, table, fields):
    fields = dict(fields, tenant_id=TENANT, creator='historical-student-import', updater='historical-student-import')
    q.execute(f'INSERT INTO `{table}` (' + ','.join(f'`{k}`' for k in fields) + ') VALUES (' + ','.join(['%s'] * len(fields)) + ')', tuple(fields.values()))
    return q.lastrowid


def apply(conn, plan, source, fail_after=None):
    check(not plan.get('reportOnly'), 'Review-only exception report cannot be applied')
    check(plan['policy'] == POLICY, 'Import policy changed; generate a new target-bound preflight plan')
    changed = []
    lock_name = 'hist-student:' + digest(plan['target']['db'])[:40]
    conn.rollback()
    try:
        with conn.cursor() as q:
            # Lock the class catalog including insertion gaps while resolving absent pairs.
            # No class-name/teacher unique index exists in the established schema.
            q.execute('SET TRANSACTION ISOLATION LEVEL REPEATABLE READ')
            check(one(q, 'SELECT GET_LOCK(%s,0) locked', (lock_name,))['locked'] == 1, 'Another local import is running')
        current = prepare(conn, source, plan['actorUsername'], lock=True)
        prior_keys = {r['key'] for r in plan['ready']}
        replayed = {r['key']: r for r in current['replay']}
        if prior_keys and prior_keys <= replayed.keys():
            check(current['target'] == plan['target'] and current['sourceSha256'] == plan['sourceSha256'], 'Replay target/source changed')
            verify(conn, [replayed[k]['serviceId'] for k in prior_keys])
            conn.rollback()
            return []
        # A prepared plan binds the exact server, source, configuration, identities and row decisions.
        changed_fields = [k for k in current if digest(current[k]) != digest(plan.get(k))]
        check(not changed_fields, 'Preflight changed (' + ','.join(changed_fields) + '); generate a new plan and review exceptions')
        with conn.cursor() as q:
            now = one(q, 'SELECT NOW() value')['value']
            created = {spec['classKey']: create_class(q, spec, plan['actorId'], now) for spec in plan['classCreates']}
            for r in plan['ready']:
                class_id = r['classId'] if r['classId'] is not None else created[r['classKey']]
                q.execute('INSERT INTO zsjos_person_contact_claim(contact_value,reservation_key,tenant_id,creator,updater) '
                          "VALUES(%s,%s,1,'historical-student-import','historical-student-import') "
                          'ON DUPLICATE KEY UPDATE id=LAST_INSERT_ID(id)', (r['mobile'], r['key'][:64]))
                claim = one(q, 'SELECT person_id,reservation_key,deleted+0 deleted FROM zsjos_person_contact_claim WHERE tenant_id=1 AND contact_value=%s FOR UPDATE', (r['mobile'],))
                check(not claim['deleted'], 'Deleted contact claim needs manual review')
                matches = fetch(q, 'SELECT id,name,mobile,wechat_id,identity_status,version FROM zsjos_person WHERE tenant_id=1 AND deleted=0 AND (mobile=%s OR wechat_id=%s) FOR UPDATE', (r['mobile'], r['mobile']))
                if matches:
                    check(len(matches) == 1 and identity_match(q, matches[0], r) is not None, 'Concurrent contact conflict')
                    pid = matches[0]['id']
                    check(claim['person_id'] in (None, pid), 'Concurrent identity ownership changed')
                    if matches[0]['identity_status'] != 'student':
                        q.execute("UPDATE zsjos_person SET identity_status='student',version=version+1,updater='historical-student-import' WHERE tenant_id=1 AND id=%s", (pid,))
                else:
                    check(claim['person_id'] is None and claim['reservation_key'] == r['key'][:64], 'Unbound contact conflict')
                    q.execute("INSERT INTO zsjos_person_no_daily_counter(sequence_date,current_value,tenant_id,creator,updater) VALUES(%s,1,1,'historical-student-import','historical-student-import') "
                              'ON DUPLICATE KEY UPDATE current_value=IF(current_value>=9999,1,current_value+1)', (now.date(),))
                    seq = one(q, 'SELECT current_value FROM zsjos_person_no_daily_counter WHERE tenant_id=1 AND sequence_date=%s FOR UPDATE', (now.date(),))['current_value']
                    pid = insert(q, 'zsjos_person', dict(person_no='XY' + now.strftime('%Y%m%d%H%M%S') + f'{seq:04d}', name=r['name'], mobile=r['mobile'], identity_status='student', first_seen_at=now, last_seen_at=now))
                q.execute("UPDATE zsjos_person_contact_claim SET person_id=%s,reservation_key=NULL WHERE tenant_id=1 AND contact_value=%s AND (person_id IS NULL OR person_id=%s)", (pid, r['mobile'], pid))
                history = dict(source='historical_import', remark=NOTE, historyIncomplete=True,
                               actualOccurredAt=None, importedAt=str(now), sourceSheet=r['sheet'], sourceRow=r['row'],
                               originalCourse=r['course'], correctedCourse=r['correctedCourse'], sourcePlanner=r['sourcePlanner'])
                if r['personState']:
                    history['identityNameMatch'] = r['personState']['matchEvidence']['kind']
                    history['contactMatchField'] = r['personState']['matchEvidence']['contactField']
                    history['identityMatchEvidence'] = r['personState']['matchEvidence']['evidence']
                    if r['personState']['name'].strip() != r['name']:
                        history['storedIdentityName'] = r['personState']['name']
                progress = dict(source='historical_import', historyIncomplete=True, items=[], studyPlanCompletedAtFirstContact=False)
                sid = insert(q, 'zsjos_service_relation', dict(person_id=pid, student_name_snapshot=r['name'], student_name_source='independent_entry',
                             teaching_delivery=1, media_ready=0, class_id=class_id, owner_user_id=r['ownerId'], status='active',
                             acceptance_status='accepted', accepted_by_user_id=plan['actorId'], accepted_at=now, activated_at=now,
                             service_snapshot=packed(r['snapshot']), delivery_stage='supervision', delivery_data_json=packed(history),
                             first_contact_progress_json=packed(progress), teaching_closed=0))
                gid = insert(q, 'zsjos_collaboration_group', dict(source_service_relation_id=sid, student_person_id=pid, status='active', version=0))
                q.execute('UPDATE zsjos_service_relation SET collaboration_group_id=%s WHERE tenant_id=1 AND id=%s', (gid, sid))
                for kind in ('student_first_contact', 'student_study_plan'):
                    insert(q, 'zsjos_student_contact_record', dict(service_relation_id=sid, contact_type=kind, successful=1, remark=NOTE,
                           attachment_file_ids_json='[]', checklist_result_json='[]', delivery_data_json=packed(history),
                           checklist_snapshot_json=packed(progress) if kind == 'student_first_contact' else None,
                           study_plan_completed_at_first_contact=0, operator_user_id=plan['actorId'], submitted_at=now,
                           idempotency_key=r['key'] + ':' + kind, request_fingerprint=r['fingerprint']))
                insert(q, 'zsjos_business_event', dict(event_type='student_historical_import', aggregate_type='student-service', aggregate_id=sid,
                       operator_user_id=plan['actorId'], occurred_at=now, idempotency_key=r['key'],
                       related_object_refs=packed(dict(personId=pid, serviceRelationId=sid, fingerprint=r['fingerprint'],
                                                      sourceSha256=plan['sourceSha256'], **history))))
                changed.append(sid)
                if fail_after and len(changed) == fail_after:
                    raise RuntimeError('Injected rollback verification failure')
            verify(conn, changed)
        conn.commit()
        return changed
    except Exception:
        conn.rollback()
        raise
    finally:
        with conn.cursor() as q:
            q.execute('SELECT RELEASE_LOCK(%s)', (lock_name,))


def verify(conn, ids):
    with conn.cursor() as q:
        for sid in ids:
            r = one(q, 'SELECT *,HEX(student_name_snapshot) nameHex FROM zsjos_service_relation WHERE tenant_id=1 AND id=%s', (sid,))
            check(r['nameHex'] == r['student_name_snapshot'].encode('utf-8').hex().upper(), 'UTF-8 identity verification failed')
            check(r['delivery_stage'] == 'supervision' and r['acceptance_status'] == 'accepted' and r['media_ready'] == b'\x00' and r['teaching_delivery'] == b'\x01', 'Service state verification failed')
            check(all(r[k] is None for k in ('order_id', 'order_item_id', 'registration_case_id', 'content_director_user_id', 'operator_user_id', 'career_planner_user_id', 'exam_date')), 'Unexpected business relationship')
            check(one(q, 'SELECT homeroom_user_id FROM zsjos_delivery_class WHERE tenant_id=1 AND id=%s', (r['class_id'],))['homeroom_user_id'] == r['owner_user_id'], 'Owner does not match class')
            history = json.loads(r['delivery_data_json'])
            check(history['remark'] == NOTE and history['actualOccurredAt'] is None, 'Historical provenance lost')
            contacts = fetch(q, 'SELECT contact_type,remark,HEX(remark) h,task_id FROM zsjos_student_contact_record WHERE tenant_id=1 AND service_relation_id=%s AND deleted=0', (sid,))
            check(len(contacts) == 2 and {c['contact_type'] for c in contacts} == {'student_first_contact', 'student_study_plan'} and
                  all(c['remark'] == NOTE and c['h'] == NOTE.encode('utf-8').hex().upper() and c['task_id'] is None for c in contacts), 'Historical steps verification failed')
            group = one(q, 'SELECT * FROM zsjos_collaboration_group WHERE tenant_id=1 AND id=%s', (r['collaboration_group_id'],))
            check(all(group[k] is None for k in ('director_user_id', 'operator_user_id')), 'Unexpected media collaborators')
    return len(ids)


def summary(plan):
    return dict(ready=len(plan['ready']), exceptions=len(plan['exceptions']), duplicates=len(plan['duplicates']),
                existing=len(plan['existing']),
                plannedNewClasses=len(plan['classCreates']),
                replay=len(plan['replay']), classCounts=dict(Counter(r['className'] for r in plan['ready'])),
                exceptionReasons=dict(Counter(r['reason'] for r in plan['exceptions'])))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('command', choices=('prepare', 'apply', 'verify'))
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--identity-confirmations', type=Path, required=True)
    parser.add_argument('--database', required=True)
    parser.add_argument('--actor-username', required=True, help='Enabled System account used only as import audit operator')
    parser.add_argument('--plan', type=Path, required=True)
    parser.add_argument('--sha256')
    parser.add_argument('--result', type=Path)
    args = parser.parse_args()
    marker = Path('/etc/zsjos/agent-environment')
    environment = os.environ.get('ZSJOS_AGENT_ENV') or (marker.read_text().strip() if marker.exists() else 'local')
    check(environment == 'local', 'Only the explicitly authorized local environment is supported')
    source = read_source(args.source, args.identity_confirmations)
    conn = connect(args.database)
    try:
        if args.command == 'prepare':
            plan = prepare(conn, source, args.actor_username)
            save_private(args.plan, plan)
            print(packed(dict(summary(plan), planSha256=hashlib.sha256(args.plan.read_bytes()).hexdigest())))
        elif args.command == 'apply':
            check(args.sha256 and hashlib.sha256(args.plan.read_bytes()).hexdigest() == args.sha256, 'Exact reviewed plan SHA256 required')
            plan = json.loads(args.plan.read_text(encoding='utf-8'))
            check(plan['sourceSha256'] == source['sourceSha256'] and plan['target']['db'] == args.database and plan['actorUsername'] == args.actor_username, 'Source/target/operator changed')
            check(args.result is not None and not args.result.exists() and not args.result.resolve().is_relative_to(ROOT), 'New private result path required before writing')
            changed = apply(conn, plan, source)
            result = dict(database=args.database, sourceSha256=source['sourceSha256'], serviceIds=changed, imported=len(changed), **summary(plan))
            save_private(args.result, result)
            print(packed({k: v for k, v in result.items() if k != 'serviceIds'}))
        else:
            plan = prepare(conn, source, args.actor_username)
            verify(conn, [r['serviceId'] for r in plan['replay']])
            print(packed(summary(plan)))
    finally:
        conn.rollback()
        conn.close()


if __name__ == '__main__':
    try:
        main()
    except Exception as e:
        # Database exception text can contain personal values; evidence stays in private files.
        print(type(e).__name__ + ': ' + (str(e) if isinstance(e, ValueError) else 'Local import failed; transaction rolled back when uncommitted'), file=sys.stderr)
        sys.exit(1)
