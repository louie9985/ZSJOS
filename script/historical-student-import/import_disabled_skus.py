#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Import only the 33 reviewed sole-SKU exceptions; keep the five SKUs disabled."""
import argparse
import hashlib
import json
import re
import sys
from pathlib import Path
import openpyxl
import import_students as imp
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'sql/mysql/tools'))
import historical_disabled_skus as config

POLICY='historical-disabled-students-20261027-v1'
SHEETS={2:'人社公共营养师-学员导入填写',4:'心理咨询师-学员导入填写',5:'执业中西药师-学员导入填写'}
APPROVED={(2,99):0,(5,85):4,(4,62):2}
APPROVED.update({(4,n):3 for n in [45,46,47,*range(49,61),*range(77,85)]})
APPROVED.update({(4,n):1 for n in [61,*range(63,69)]})

def read_source(path,exceptions):
    path,exceptions=Path(path),Path(exceptions)
    sha=hashlib.sha256(path.read_bytes()).hexdigest()
    review_sha=hashlib.sha256(exceptions.read_bytes()).hexdigest()
    wb=openpyxl.load_workbook(path,read_only=True,data_only=False)
    review=openpyxl.load_workbook(exceptions,read_only=True,data_only=True)
    try:
        sheet=review['异常主表']
        rows=list(sheet.values)
        imp.check(list(rows[0][:9])==['异常原因','姓名','手机号','课程规格','班级','规划师','待确认内容','原表行号','处理意见'],'Exception headers changed')
        records=[]
        for (index,n),spec_index in sorted(APPROVED.items()):
            name=SHEETS[index]
            imp.check(wb.sheetnames[index]==name,'Source sheet identity changed')
            headers=next(wb[name].iter_rows(min_row=1,max_row=1,values_only=True))
            imp.check([headers[i] for i in (0,1,3,4,5,6)]==['姓名 *','手机号码 †','课程名称 *','课程规格名称 t','班级名称 *','学习规划师姓名 *'],'Source headers changed')
            cells=next(wb[name].iter_rows(min_row=n,max_row=n))
            values=[c.value for c in cells]
            imp.check(len(values)==44 and not any(c.data_type=='f' for c in cells),'Source structure/formula requires review')
            imp.check(values[2] is None and all(v is None for v in values[7:]),'Additional source fields require review')
            clean=[str(v).strip() if v is not None else '' for v in values[:7]]
            clean[6]={'周老师':'周老师【梁颖】','周老师【梁颖】【梁颖】':'周老师【梁颖】'}.get(clean[6],clean[6])
            imp.check(clean[4]==config.SPECS[spec_index][1] and all(clean[i] for i in (0,1,3,4,5,6)), 'Reviewed source row/spec changed')
            imp.check(re.fullmatch(r'1[3-9]\d{9}',clean[1]) is not None,'Source contact requires review')
            coord=name.split('-')[0]+'第'+str(n)+'行'
            rr=[r for r in rows[1:] if len(r)>7 and r[7]==coord]
            imp.check(len(rr)==1 and rr[0][0]=='课程规格无法匹配系统','Row has additional/unreviewed exception or was removed')
            imp.check([str(rr[0][i]).strip() for i in (1,2,3,4,5)]==[clean[i] for i in (0,1,4,5,6)],'Exception/source content changed')
            record=dict(sheet=name,sheetIndex=index,row=n,values=values,name=clean[0],mobile=clean[1],course=clean[3],sku=clean[4],className=clean[5],sourcePlanner=clean[6],
                        key=f'hist-student:{sha[:24]}:{index+1}:{n}',fingerprint=imp.digest({'policy':POLICY,'source':clean}),priorFingerprints=[])
            records.append(record)
        imp.check(len(records)==33,'Reviewed population changed')
        imp.check(len({(r['mobile'],r['sku'],r['className'],r['sourcePlanner']) for r in records})==33,'Duplicate teaching target in reviewed source; review before importing')
        auth=dict(policy=POLICY,sourceSha256=sha,exceptionSha256=review_sha,rows={r['key']:imp.digest(r) for r in records})
        return dict(policy=imp.POLICY,source=str(path.resolve()),sourceSha256=sha,candidates=records,exceptions=[],duplicates=[],historicalSkuAuthorization=auth)
    finally:wb.close();review.close()

def authorized_sku(c,r,source):
    auth=source['historicalSkuAuthorization']
    imp.check(auth['policy']==POLICY and auth['sourceSha256']==source['sourceSha256'] and auth['rows'].get(r['key'])==imp.digest(r),'Historical SKU source/row authorization mismatch')
    index=APPROVED.get((r['sheetIndex'],r['row']))
    imp.check(index is not None and r['sheet']==SHEETS[r['sheetIndex']],'Unreviewed historical SKU row')
    product,name,value,unit=config.SPECS[index]
    imp.check(r['sku']==name,'Unreviewed historical SKU')
    ps=[p for p in c['zsjos_product'] if p['name']==product and p['status']==0]
    imp.check(len(ps)==1,'Historical SKU product missing/ambiguous')
    ss=[s for s in c['zsjos_product_sku'] if s['sku_name']==name]
    imp.check(len(ss)==1,'Historical SKU missing/ambiguous; run configuration prepare/apply first')
    s=ss[0]
    imp.check(s['spu_id']==ps[0]['id'] and s['status']==1 and s['price_unit']==unit and s['price']==0
              and s['pricing_note']==config.NOTE and json.loads(s['attr_values_json'])=={'sales_spec':value}, 'Historical SKU must retain exact disabled configuration')
    return s

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('command',choices=('prepare','apply','verify'))
    p.add_argument('--source',type=Path,required=True)
    p.add_argument('--exceptions',type=Path,required=True)
    p.add_argument('--database',required=True)
    p.add_argument('--connection-file',type=Path)
    p.add_argument('--actor-username',required=True)
    p.add_argument('--plan',type=Path,required=True)
    p.add_argument('--sha256')
    p.add_argument('--result',type=Path)
    a=p.parse_args()
    source=read_source(a.source,a.exceptions)
    c=config.connect(a.database,a.connection_file)
    try:
        if a.command=='prepare':
            config.inspect(c)
            plan=imp.prepare(c,source,a.actor_username)
            imp.save_private(a.plan,plan)
            print(imp.packed(dict(imp.summary(plan),planSha256=hashlib.sha256(a.plan.read_bytes()).hexdigest())))
        elif a.command=='apply':
            imp.check(a.sha256 and hashlib.sha256(a.plan.read_bytes()).hexdigest()==a.sha256,'Exact reviewed plan SHA256 required')
            imp.check(a.result and not a.result.exists() and not a.result.resolve().is_relative_to(imp.ROOT),'New private result path required')
            plan=json.loads(a.plan.read_text(encoding='utf-8'))
            imp.check(not plan['exceptions'],'Scoped preflight still has exceptions; resolve them before applying')
            imp.check(plan['actorUsername']==a.actor_username and plan['target']['db']==a.database,'Target/audit operator changed')
            imp.check(plan.get('historicalSkuAuthorization')==source['historicalSkuAuthorization'],'Source/review authorization changed')
            config.inspect(c)
            ids=imp.apply(c,plan,source)
            imp.save_private(a.result,dict(imported=len(ids),serviceIds=ids,target=plan['target'],sourceSha256=source['sourceSha256']))
            print(imp.packed({'imported':len(ids)}))
        else:
            original=json.loads(a.plan.read_text(encoding='utf-8'))
            imp.check(original.get('historicalSkuAuthorization')==source['historicalSkuAuthorization']
                      and original['actorUsername']==a.actor_username,'Verification source/operator changed')
            plan=imp.prepare(c,source,a.actor_username)
            imp.check(plan['target']==original['target'],'Verification target changed')
            replay={r['key']:r for r in plan['replay']}
            expected={r['key'] for r in original['ready']}
            imp.check(expected<=replay.keys(),'Some planned records have not been imported')
            verified=imp.verify(c,[replay[k]['serviceId'] for k in expected])
            print(imp.packed(dict(imp.summary(plan),verified=verified)))
    finally:c.rollback();c.close()

if __name__=='__main__':
    try:main()
    except Exception as e:
        print(str(e) if isinstance(e,ValueError) else type(e).__name__+': scoped import failed; inspect audit before recovery',file=sys.stderr)
        sys.exit(1)
