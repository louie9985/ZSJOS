#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Resolve the user's four retained enrollments and seven source exclusions.

Test tenant 1 only; reuses guarded historical-import transactions and fixtures.
No existing service is deleted. Private backups retain all original source data.
"""
import argparse
from copy import copy
import json
import os
from pathlib import Path
import shutil
import zipfile

import openpyxl
import import_education_exceptions_20261027 as x
from import_two_confirmed_students_20261028 import session
from archive_matching_education_students_20261028 import clone_cell

ROOT=Path('/opt/zsjos-runtime/backups/education-pending-resolved-20261028')
MARKER='edu-pending-resolved-20261028'
PENDING=x.ROOT/'output/教务学员_多条报名待确认_20261027.xlsx'
REPORT=x.ROOT/'output/教务学员_多条报名处理结果_20261028.xlsx'
TARGETS={('学会公共营养师',53),('人社健康管理师',128),('心理咨询师',19),('执业中西药师',39)}


def prepare():
    x.h.require(not (ROOT/'plan.json').exists(),'Frozen plan exists')
    _,records=x.source_rows(x.EXCEPTIONS)
    selected=[r for r in records if (r['sheet'],r['row']) in TARGETS]
    x.h.require(len(records)==67 and len(selected)==4,'Source scope changed')
    source_hashes={p.name:x.sha(p) for p in (x.MASTER,x.EXCEPTIONS,PENDING)}
    for p in (x.MASTER,x.EXCEPTIONS,PENDING):
        target=ROOT/p.name
        x.h.require(not target.exists(),'Existing backup requires review')
        shutil.copy2(p,target);target.chmod(0o600)
    master=openpyxl.load_workbook(x.MASTER,data_only=True)
    pending=openpyxl.load_workbook(PENDING,data_only=True)['待确认报名']
    headings={c.value:c.column for c in pending[1]}
    snapshot=x.h.snapshot();s=session();options=x.h.api(s,'GET','/zsjos/student-entry/options')
    pending_refs={(r[1],r[2]) for r in list(pending.values)[1:]}
    x.h.require(len(pending_refs)==11 and TARGETS <= pending_refs,'Pending scope differs')
    for r in selected:
        raw=list(master[r['sheet']+'-学员导入填写'].values)[r['row']-1]
        x.h.require([x.norm(v) for v in raw[:7]]==[r['name'],r['mobile'],None,r['product'],r['sku'],r['class_name'],r['planner']],
                    'Master row differs from confirmation')
        matched=[row for row in list(pending.values)[1:] if (row[headings['总表工作表']-1],row[headings['总表行号']-1])==(r['sheet'],r['row'])]
        x.h.require(len(matched)==1 and list(matched[0][5:])==list(raw),'Full-field pending row differs')
        products=[p for p in snapshot['products'] if p['name']==r['product'] and p['status']==0]
        x.h.require(len(products)==1,'Product unavailable')
        r['product_id'],r['product_ref']=products[0]['id'],products[0]['product_ref']
        skus=[k for k in snapshot['skus'] if k['spu_id']==r['product_id'] and k['sku_name']==r['sku'] and k['status']==0]
        x.h.require(len(skus)==1 and skus[0]['sku_ref'] in {k['skuRef'] for k in options['courses']},'SKU unavailable')
        r['sku_id'],r['skuRef']=skus[0]['id'],skus[0]['sku_ref']
        owners=[u for u in snapshot['users'] if u['nickname']==r['planner'] and u['status']==0]
        x.h.require(len(owners)==1,'Planner ambiguous')
        r['owner_id']=owners[0]['id']
        classes=[c for c in snapshot['classes'] if c['class_name']==r['class_name'] and c['homeroom_user_id']==r['owner_id'] and c['status']=='SERVING']
        x.h.require(len(classes)==1 and classes[0]['id'] in {c['id'] for c in options['classes']},'Class unavailable')
        r['class_id']=classes[0]['id']
        people=[p for p in snapshot['people'] if p['mobile']==r['mobile']]
        x.h.require(len(people)<=1 and all(p['name']==r['name'] for p in people),'Mobile/name changed')
        if people:r['person_id']=people[0]['id']
        check=x.h.api(s,'POST','/zsjos/student-entry/check',{'mobile':r['mobile'],'wechatId':None})
        x.h.require(check['status'] in ('new','reuse') and check.get('personId')==r.get('person_id'),'Live identity conflict')
        r['same_services']=[];r['other_services']=[]
        existing=[rel for rel in snapshot['services'] if rel['person_id']==r.get('person_id')]
        same=[rel for rel in existing if json.loads(rel['service_snapshot'] or '{}').get('productRef')==r['product_ref']
              or json.loads(rel['service_snapshot'] or '{}').get('name')==r['product']]
        if (r['sheet'],r['row'])==('学会公共营养师',53):
            x.h.require(len(same)==1 and json.loads(same[0]['service_snapshot'])['skuName']==
                '公共营养师｜人社部技能等级｜三级｜超值精品班','Previously confirmed enrollment changed')
            r['confirmed_same_product_ids']=[same[0]['id']]
        else:x.h.require(not same,'Unexpected existing same-product service')
        x.check_existing_courses(r,existing)
        r['other_services']=[rel['id'] for rel in existing if rel not in same]
        r['issues']=[]
    x.h.require(all(x.sha(p)==source_hashes[p.name] for p in (x.MASTER,x.EXCEPTIONS,PENDING)),'Source changed during preflight')
    x.prepare(selected,snapshot,sync=False,missing_sku_count=0,extra=dict(mobile_only=True,
        fixture_failure_points=[1,8,16,24,35],wrapper_sha256=x.sha(Path(__file__)),pending_sha256=x.sha(PENDING),source_count=67, pending_refs=sorted(pending_refs), excluded_refs=sorted(pending_refs-TARGETS)))


def verify_database(plan):
    c=x.h.connection()
    try:
        with c.cursor() as q:
            q.execute('START TRANSACTION WITH CONSISTENT SNAPSHOT, READ ONLY')
            after,result=x.verify(q,plan)
    finally:c.rollback();c.close()
    x.h.save(ROOT/'independent-after.json',after)
    s=session()
    for row in result:
        rid=row['service_id']
        context=x.h.api(s,'GET',f'/zsjos/student/service/{rid}/contact-context')
        stages={a['code']:a for a in context['deliveryStages']}
        x.h.require(context['acceptanceStatus']=='accepted' and context['deliveryStage']=='supervision'
            and context['currentTask'] is None and stages['first_contact']['status']==stages['study_plan']['status']=='done'
            and all(a['completed'] for a in context['firstContactProgress']['items']),'API stage incomplete')
        history=x.h.api(s,'GET',f'/zsjos/student/service/{rid}/contact-records?pageNo=1&pageSize=100')
        x.h.require(history['total']==2 and {r['contactType'] for r in history['list']}=={t[0] for t in x.TYPES}
            and all(r['successful'] and r['remark']==x.REMARK for r in history['list']),'API completion records differ')
    x.h.save(ROOT/'api-verification.json',dict(contexts=4,histories=4,passed=True))


def finish(plan):
    if (ROOT/'files.json').exists():
        hashes=x.h.read(ROOT/'files.json')
        x.h.require(all(x.sha(p)==hashes[p.name] for p in (x.MASTER,x.EXCEPTIONS,PENDING,REPORT)),
                    'Delivered workbook changed; review required')
        print('File replay: no changes');return
    verify_database(plan)
    x.h.require(x.sha(x.MASTER)==plan['master_sha256'] and x.sha(x.EXCEPTIONS)==plan['exception_sha256']
                and x.sha(PENDING)==plan['pending_sha256'],'Source workbook changed')
    refs={tuple(r) for r in plan['pending_refs']};excluded={tuple(r) for r in plan['excluded_refs']}
    master=openpyxl.load_workbook(x.MASTER);original=openpyxl.load_workbook(x.MASTER)
    pending=openpyxl.load_workbook(PENDING);pending_rows=list(pending['待确认报名'].values)
    exceptions=openpyxl.load_workbook(x.EXCEPTIONS);old_exc=openpyxl.load_workbook(x.EXCEPTIONS)
    report=openpyxl.Workbook();report.active.title='汇总'
    for row in [('处理项','数量','说明'),('确认导入',4,'常规督学；首联和学习计划完成'),('排除报名',7,'从总表清空，保留原始行号；完整原信息见明细'),
                ('剩余待确认报名',0,'本次11条全部处理'),('异常表剩余',57,'本次从异常表移出10条；另1条原本不在异常表')]:report.active.append(row)
    full=report.create_sheet('完整原表与处理结果');full.append(['组号','总表工作表','总表行号','处理结果','备注']+list(pending_rows[0][5:]))
    for outrow,r in enumerate(pending_rows[1:],2):
        ref=(r[1],r[2]);source=original[r[1]+'-学员导入填写']
        x.h.require(tuple(c.value for c in source[r[2]])==r[5:],'Pending/master full fields differ')
        kept=ref in TARGETS
        for col,value in enumerate([r[0],r[1],r[2],'已导入' if kept else '用户确认排除',x.REMARK if kept else '仅排除本批报名，既有系统服务保留'],1):full.cell(outrow,col,value)
        for col in range(1,45):clone_cell(source.cell(r[2],col),full.cell(outrow,col+5))
    for sheet in report:x.format_sheet(sheet)
    for sheet,row in excluded:
        for cell in master[sheet+'-学员导入填写'][row]:cell.value=None
    ps=pending['待确认报名'];ps.delete_rows(2,ps.max_row-1);ps.auto_filter.ref='A1:AW1'
    resolved=pending['已确认导入']
    for sheet,row in sorted(TARGETS):resolved.append([sheet,row,'用户已确认，已导入并进入常规督学',x.REMARK])
    x.format_sheet(resolved)
    archived=pending.create_sheet('已排除报名');archived.append(pending_rows[0])
    for r in pending_rows[1:]:
        if (r[1],r[2]) in excluded:
            archived.append(r)
            archived.cell(archived.max_row,5,'用户确认删除，已同步总表；原始信息保留')
    x.format_sheet(archived)
    es=exceptions['异常与处理总表']
    expected=[r for r in list(es.values)[1:] if (r[7],r[8]) not in refs]
    for row in range(es.max_row,1,-1):
        if (es.cell(row,8).value,es.cell(row,9).value) in refs:es.delete_rows(row)
    x.h.require(es.max_row==58 and len(expected)==57,'Exception partition changed')
    es.tables['ExceptionTable'].ref='A1:U58';es.auto_filter.ref='A1:U58'
    cs=exceptions['关联记录对比'];removed_comparisons=0
    for row in range(cs.max_row,1,-1):
        if (cs.cell(row,4).value,cs.cell(row,5).value) in refs:cs.delete_rows(row);removed_comparisons+=1
    cs.tables['ComparisonTable'].ref=f'A1:K{cs.max_row}';cs.auto_filter.ref=cs.tables['ComparisonTable'].ref
    summary=exceptions['汇总']
    for row in range(2,summary.max_row+1):
        key=summary.cell(row,1).value
        if key=='当前保留异常':summary.cell(row,2,57)
        if key=='多条报名待确认':summary.cell(row,2,0);summary.cell(row,3,'4条已导入、7条排除；本批无待确认报名')
        if key=='当前异常构成':summary.cell(row,2,57);summary.cell(row,3,'其他异常57条；多条报名清单已处理完毕')
    summary.append(['2026-10-28多条报名最终处理',11,'确认4条导入、排除7条；本次异常主表移出10条'])
    summary.tables['SummaryTable'].ref=f'A1:C{summary.max_row}';summary.auto_filter.ref=summary.tables['SummaryTable'].ref
    books=((x.MASTER,master),(x.EXCEPTIONS,exceptions),(PENDING,pending),(REPORT,report))
    for path,book in books:
        temp=path.with_suffix('.tmp.xlsx');book.save(temp)
        with zipfile.ZipFile(temp) as z:x.h.require(z.testzip() is None,'Invalid workbook ZIP')
    saved={p:openpyxl.load_workbook(p.with_suffix('.tmp.xlsx')) for p,_ in books}
    for sheet in original:
        for row in sheet:
            for cell in row:
                new=saved[x.MASTER][sheet.title].cell(cell.row,cell.column)
                blank=(sheet.title.replace('-学员导入填写',''),cell.row) in excluded
                x.h.require(new.value==(None if blank else cell.value),'Unexpected master cell change')
    x.h.require(list(saved[x.EXCEPTIONS]['异常与处理总表'].values)[1:]==expected,'Retained exception cells changed')
    x.h.require(list(saved[x.EXCEPTIONS]['聊天处理依据'].values)==list(old_exc['聊天处理依据'].values),'Historical decisions changed')
    x.h.require(saved[PENDING]['待确认报名'].max_row==1 and saved[PENDING]['已排除报名'].max_row==8,'Pending partition differs')
    for row in list(saved[REPORT]['完整原表与处理结果'].rows)[1:]:
        sheet=original[row[1].value+'-学员导入填写'];source_row=row[2].value
        for col in range(1,45):
            a,b=sheet.cell(source_row,col),row[col+4]
            x.h.require((a.value,a.data_type,a.number_format)==(b.value,b.data_type,b.number_format),'Archived source cell differs')
    x.h.require(x.sha(x.MASTER)==plan['master_sha256'] and x.sha(x.EXCEPTIONS)==plan['exception_sha256']
                and x.sha(PENDING)==plan['pending_sha256'],'Source changed before replacement')
    for path,_ in books:path.with_suffix('.tmp.xlsx').replace(path)
    x.h.save(ROOT/'files.json',{p.name:x.sha(p) for p,_ in books})
    x.h.save(ROOT/'verification.json',dict(imported=4,excluded=7,remaining_exceptions=57,pending=0,
        source_fields=44,completed_tasks=8,records=8,comparison_removed=removed_comparisons))
    print('PASS: 4 imported, 7 excluded; 57 exceptions remain; pending 0; full 44 source fields archived')


def main():
    env=os.getenv('ZSJOS_AGENT_ENV')
    if env is None:env=Path('/etc/zsjos/agent-environment').read_text().strip()
    x.h.require(env=='test','Test-only execution')
    os.umask(0o077);ROOT.mkdir(mode=0o700,parents=True,exist_ok=True)
    x.EVIDENCE=ROOT;x.MARKER=MARKER
    parser=argparse.ArgumentParser();parser.add_argument('mode',choices=['prepare','fixture','apply','finish']);parser.add_argument('--sha256');args=parser.parse_args()
    if args.mode=='prepare':prepare();return
    plan=x.loads((ROOT/'plan.json').read_text());x.h.require(args.sha256==x.h.digest(plan),'Plan digest differs')
    if args.mode!='finish':x.h.require(plan['wrapper_sha256']==x.sha(Path(__file__)),'Mutation wrapper changed')
    if args.mode=='apply':x.h.require(x.sha(PENDING)==plan['pending_sha256'],'Pending workbook changed')
    if args.mode=='finish':finish(plan)
    else:getattr(x,args.mode)(plan)


if __name__=='__main__':main()
