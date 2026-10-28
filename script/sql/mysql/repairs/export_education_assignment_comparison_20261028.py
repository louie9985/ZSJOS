#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Read-only comparison export; all decisions remain blank for user review."""
from collections import Counter
from datetime import datetime
import json
import os
from pathlib import Path
import zipfile

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
import import_education_exceptions_20261027 as x
from archive_matching_education_students_20261028 import clone_cell

EVIDENCE=Path('/opt/zsjos-runtime/backups/education-assignment-comparison-20261028')
OUTPUT=x.ROOT/'output/教务学员_22条班级班主任差异对比_20261028.xlsx'
PENDING=x.ROOT/'output/教务学员_多条报名待确认_20261027.xlsx'


def run():
    env=os.getenv('ZSJOS_AGENT_ENV')
    if env is None:env=Path('/etc/zsjos/agent-environment').read_text().strip()
    x.h.require(env=='test','Test-only export')
    os.umask(0o077);EVIDENCE.mkdir(mode=0o700,parents=True,exist_ok=True)
    sources=(x.MASTER,x.EXCEPTIONS,PENDING)
    hashes={p.name:x.sha(p) for p in sources}
    if (EVIDENCE/'delivery.json').exists():
        old=x.h.read(EVIDENCE/'delivery.json')
        x.h.require(old['source_hashes']==hashes and old['output_sha256']==x.sha(OUTPUT),'Review file or sources changed; do not overwrite')
        print('Replay: no changes');return
    x.h.require(not OUTPUT.exists(),'Existing review workbook must not be overwritten')
    _,records=x.source_rows(x.EXCEPTIONS);st=x.h.snapshot()
    classes={c['id']:c for c in st['classes']};users={u['id']:u for u in st['users']}
    selected=[];counts=Counter()
    for r in records:
        people=[p for p in st['people'] if p['mobile']==r['mobile']]
        if len(people)!=1:continue
        person=people[0]
        same=[]
        for service in st['services']:
            snap=json.loads(service['service_snapshot'] or '{}')
            if service['person_id']==person['id'] and snap.get('name',snap.get('productName'))==r['product'] and snap.get('skuName')==r['sku']:
                same.append((service,snap))
        if not same:continue
        details=[]
        for service,snap in same:
            cl=classes.get(service['class_id']);owner=users.get(service['owner_user_id'])
            x.h.require(cl is not None and owner is not None and cl['homeroom_user_id'] in users,'Unresolved current system assignment')
            homeroom=users[cl['homeroom_user_id']]
            differences=[]
            if cl['class_name']!=r['class_name']:differences.append('班级不同')
            if homeroom['nickname']!=r['planner']:differences.append('班主任不同')
            if owner['nickname']!=r['planner'] and owner['id']!=homeroom['id']:differences.append('服务负责人不同')
            x.h.require(differences,'Exact match unexpectedly present')
            counts['、'.join(differences)]+=1
            details.append(dict(service=service,snapshot=snap,cl=cl,owner=owner,homeroom=homeroom,difference='、'.join(differences)))
        selected.append(dict(source=r,person=person,details=details))
    x.h.require(len(selected)==22 and sum(len(a['details']) for a in selected)==22,'Reviewed comparison scope changed')
    master=openpyxl.load_workbook(x.MASTER);exceptions=openpyxl.load_workbook(x.EXCEPTIONS)
    book=openpyxl.Workbook();summary=book.active;summary.title='说明与汇总'
    for row in [('项目','数量或说明'),('核对时间',datetime.now().strftime('%Y-%m-%d %H:%M:%S 北京时间')),
                ('待确认原表记录',22),('系统同课程同规格服务',22),*sorted(counts.items()),
                ('阅读顺序','先看逐项对比；原表完整信息保留全部44列；原异常记录保留原处理说明'),
                ('班主任口径','系统班主任取当前班级实际班主任；服务负责人单独展示，防止混淆'),
                ('确认方式','在逐项对比的“您的处理意见”中填写以系统为准、按原表调整，或具体要求；备注栏可补充说明'),
                ('本次操作','只读核对并生成此文件；不调整归属、不删除异常、不新增服务')]:summary.append(row)
    comp=book.create_sheet('逐项对比')
    comp.append(['序号','总表工作表','总表行号','差异结论','您的处理意见','补充说明',
        '表格姓名','系统档案姓名','系统服务显示姓名','表格手机号','系统手机号','系统微信号',
        '表格课程','系统课程','表格课程规格','系统课程规格',
        '表格班级','系统班级','表格班主任（规划师）','系统班级班主任','系统服务负责人',
        '班级是否一致','班主任是否一致','服务负责人是否一致','姓名是否一致','手机号是否一致',
        '系统服务状态（原值）','系统接收状态（原值）','系统交付阶段（原值）','系统教学是否关闭','系统服务创建时间'])
    full=book.create_sheet('原表完整信息')
    headers=[c.value for c in master[selected[0]['source']['sheet']+'-学员导入填写'][1]]
    x.h.require(len(headers)==44,'Master field count changed')
    full.append(['序号','总表工作表','总表行号','差异结论']+headers)
    history=book.create_sheet('原异常记录');history.append(['序号']+[c.value for c in exceptions['异常与处理总表'][1]])
    for i,item in enumerate(selected,1):
        r=item['source'];person=item['person'];src=master[r['sheet']+'-学员导入填写']
        x.h.require([c.value for c in src[1]]==headers,'Master headings differ')
        x.h.require([x.norm(src.cell(r['row'],c).value) for c in (1,2,4,5,6,7)]==[r['name'],r['mobile'],r['product'],r['sku'],r['class_name'],r['planner']], 'Master/exception differs')
        for d in item['details']:
            s=d['service'];cl=d['cl'];owner=d['owner'];homeroom=d['homeroom'];snap=d['snapshot']
            eq=lambda a,b:'是' if a==b else '否'
            comp.append([i,r['sheet'],r['row'],d['difference'],None,None,r['name'],person['name'],s['student_name_snapshot'],
                r['mobile'],person['mobile'],person['wechat_id'],r['product'],snap.get('name',snap.get('productName')),r['sku'],snap['skuName'],
                r['class_name'],cl['class_name'],r['planner'],homeroom['nickname'],owner['nickname'],
                eq(r['class_name'],cl['class_name']),eq(r['planner'],homeroom['nickname']),eq(r['planner'],owner['nickname']),
                '是' if person['name'].strip()==r['name']==s['student_name_snapshot'] else '否',eq(r['mobile'],person['mobile']),
                s['status'],s['acceptance_status'],s['delivery_stage'],'是' if s['teaching_closed'] in (1,b'\1') else '否',s['create_time']])
        for col,val in enumerate([i,r['sheet'],r['row'],'；'.join(d['difference'] for d in item['details'])],1):full.cell(i+1,col,val)
        for col in range(1,45):clone_cell(src.cell(r['row'],col),full.cell(i+1,col+4))
        history.cell(i+1,1,i)
        for col in range(1,22):clone_cell(exceptions['异常与处理总表'].cell(r['exception_row'],col),history.cell(i+1,col+1))
    for ws in book:
        x.format_sheet(ws);ws.sheet_view.zoomScale=85;ws.sheet_properties.pageSetUpPr.fitToPage=True
    summary.column_dimensions['A'].width=26;summary.column_dimensions['B'].width=105
    for ws,name in ((comp,'AssignmentCompare'),(full,'FullSource'),(history,'OriginalExceptions')):
        table=Table(displayName=name,ref=ws.dimensions);table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(table)
    comp.freeze_panes='G2';full.freeze_panes='G2';history.freeze_panes='D2'
    for row in comp.iter_rows(min_row=2):
        comp.row_dimensions[row[0].row].height=66
        for col in (5,6):row[col-1].fill=PatternFill('solid',fgColor='E2F0D9')
        pairs=((17,18),(19,20),(19,21),(7,8),(7,9),(10,11),(15,16))
        for left,right in pairs:
            if row[left-1].value!=row[right-1].value:
                for col in (left,right):row[col-1].fill=PatternFill('solid',fgColor='FFF2CC');row[col-1].font=Font(color='9C0006')
        for col in (10,11,12):row[col-1].number_format='@'
    for col in ('A','C'):comp.column_dimensions[col].width=9
    for col in ('E','F'):comp.column_dimensions[col].width=32
    for col in ('O','P','Q','R'):comp.column_dimensions[col].width=40
    for col in ('S','T','U'):comp.column_dimensions[col].width=28
    temp=OUTPUT.with_suffix('.tmp.xlsx');book.save(temp)
    with zipfile.ZipFile(temp) as z:x.h.require(z.testzip() is None,'Invalid workbook ZIP')
    saved=openpyxl.load_workbook(temp)
    x.h.require(list(saved['逐项对比'].values)==list(comp.values),'Comparison saved values differ')
    for i,item in enumerate(selected,2):
        r=item['source'];src=master[r['sheet']+'-学员导入填写']
        for col in range(1,45):
            a,b=src.cell(r['row'],col),saved['原表完整信息'].cell(i,col+4)
            x.h.require((a.value,a.data_type,a.number_format)==(b.value,b.data_type,b.number_format),'Full source field differs')
        x.h.require(tuple(c.value for c in saved['原异常记录'][i])[1:]==tuple(c.value for c in exceptions['异常与处理总表'][r['exception_row']]),'Original exception differs')
    x.h.require(all(x.sha(p)==hashes[p.name] for p in sources),'Source changed during export')
    x.h.save(EVIDENCE/'snapshot.json',st);x.h.save(EVIDENCE/'selection.json',selected)
    temp.replace(OUTPUT)
    x.h.save(EVIDENCE/'delivery.json',dict(source_hashes=hashes,output_sha256=x.sha(OUTPUT),source_rows=22,comparisons=22,categories=dict(counts),database_writes=0))
    print(x.j(dict(source_rows=22,comparisons=22,categories=dict(counts),full_source_columns=44,database_writes=0)))


if __name__=='__main__':run()
