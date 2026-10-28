#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Test-only historical assignment correction; preserve closed and existing services.

Frozen preimages, actual MySQL temporary fixtures, exact plan and audit keys gate
writes. No new service, contact record, stage reset or notification is produced.
"""
import argparse
from datetime import datetime
import json
import os
from pathlib import Path
import shutil
import zipfile

import openpyxl
import import_education_exceptions_20261027 as x
from export_education_assignment_comparison_20261028 import EVIDENCE as REVIEW_EVIDENCE, OUTPUT as REVIEW, PENDING
from archive_matching_education_students_20261028 import clone_cell
from import_two_confirmed_students_20261028 import session

ROOT=Path('/opt/zsjos-runtime/backups/education-unscheduled-20261028')
MARKER='edu-unscheduled-20261028'
REPORT=x.ROOT/'output/教务学员_考期班级处理结果_20261028.xlsx'
REQUEST='zsjos_class_transfer_request'
TABLES=tuple(t for t in x.TABLES if t!=x.COUNTER)+(REQUEST,)
TASK_TYPES={'student_first_contact','student_study_plan','student_contact','student_first_contact_assistance'}


def capture(q,plan,lock=False):
    return x.capture(q,plan,lock)


def prepare():
    x.h.require(not (ROOT/'plan.json').exists(),'Frozen plan exists')
    _,exceptions=x.source_rows(x.EXCEPTIONS);x.h.require(len(exceptions)==41,'Exception scope changed')
    lookup={(r['sheet'],r['row']):r for r in exceptions}
    inputs=(x.MASTER,x.EXCEPTIONS,REVIEW,PENDING);hashes={p.name:x.sha(p) for p in inputs}
    st=x.h.snapshot();old_review=x.h.read(REVIEW_EVIDENCE/'selection.json')
    classes={c['id']:c for c in st['classes']};services={s['id']:s for s in st['services']}
    records=[];newclasses=[];master=openpyxl.load_workbook(x.MASTER)
    for item in old_review:
        old=item['source'];r=dict(lookup[(old['sheet'],old['row'])]);x.h.require(len(item['details'])==1,'Multiple service candidates')
        s=services[item['details'][0]['service']['id']];sourceclass=classes[s['class_id']]
        person=[p for p in st['people'] if p['mobile']==r['mobile']]
        snap=json.loads(s['service_snapshot'] or '{}')
        x.h.require(len(person)==1 and person[0]['id']==s['person_id'] and person[0]['name'].strip()==r['name']==s['student_name_snapshot']
            and snap.get('name',snap.get('productName'))==r['product'] and snap.get('skuName')==r['sku'],'Identity/course drift')
        x.h.require([x.norm(master[r['sheet']+'-学员导入填写'].cell(r['row'],col).value) for col in (1,2,4,5,6,7)]==
            [r['name'],r['mobile'],r['product'],r['sku'],r['class_name'],r['planner']],'Master/exception differs')
        r.update(service_id=s['id'],person_id=s['person_id'],old_class_id=s['class_id'],old_owner=s['owner_user_id'])
        if sourceclass['class_kind']=='EXCEPTION':
            x.h.require(s['teaching_closed']==b'\1','Exception class is not closed');r['outcome']='exclude'
        elif sourceclass['class_kind']=='UNSCHEDULED':
            x.h.require(s['teaching_closed']==b'\0' and s['teaching_delivery']==b'\1' and s['status']=='active'
                and s['acceptance_status']=='accepted' and s['delivery_stage']=='supervision','Unscheduled lifecycle differs')
            r['outcome']='transfer'
            users=[u for u in st['users'] if u['nickname']==r['planner'] and u['status']==0]
            x.h.require(len(users)==1,'Planner unavailable');r['owner_id']=users[0]['id']
            targets=[c for c in st['classes'] if c['class_name']==r['class_name'] and c['homeroom_user_id']==r['owner_id'] and c['status']=='SERVING' and c['class_kind']=='NORMAL']
            x.h.require(len(targets)<=1,'Target class ambiguous')
            if targets:r['class_id']=targets[0]['id']
            else:
                r['missing_class']=True
                similar=[c for c in st['classes'] if c['class_name']==r['class_name'] and c['class_kind']=='NORMAL' and c['status']=='SERVING']
                dates={(c['exam_schedule_id'],c['exam_schedule_snapshot']) for c in similar}
                x.h.require(len(dates)<=1,'Existing exam snapshots conflict')
                dateid,datesnap=next(iter(dates)) if dates else (None,None)
                depts=x.h.query('SELECT id,name FROM system_dept WHERE tenant_id=1 AND deleted=0 AND status=0 AND id=%s',(users[0]['dept_id'],))
                x.h.require(len(depts)==1,'Planner department unavailable')
                newclasses.append(dict(class_no='BJ'+datetime.now().strftime('%Y%m%d%H%M%S')+'0001',class_name=r['class_name'],
                    system_class=0,class_kind='NORMAL',homeroom_user_id=r['owner_id'],homeroom_user_name_snapshot=users[0]['nickname'],
                    dept_id=depts[0]['id'],dept_name_snapshot=depts[0]['name'],status='SERVING',version=0,
                    exam_schedule_id=dateid,exam_schedule_snapshot=datesnap))
        else:r['outcome']='hold'
        records.append(r)
    x.h.require(len(records)==22 and sum(r['outcome']=='transfer' for r in records)==19 and sum(r['outcome']=='exclude' for r in records)==2
        and len(newclasses)==1,'Reviewed scope changed')
    pids=','.join(map(str,sorted({r['person_id'] for r in records})));sids=','.join(map(str,sorted({r['service_id'] for r in records})))
    scope={x.P:(f'id IN ({pids})',[]),x.S:(f'person_id IN ({pids})',[]),x.G:(f'student_person_id IN ({pids})',[]),
        x.T:(f"biz_type='student_service' AND biz_id IN ({sids})",[]),x.R:(f'service_relation_id IN ({sids})',[]),
        x.CLAIM:(f'person_id IN ({pids})',[]),x.E:('creator=%s',[MARKER]),REQUEST:(f'service_relation_id IN ({sids})',[])}
    plan=dict(records=records,classes=newclasses,scope=scope,owners=sorted({r['old_owner'] for r in records}|{r['owner_id'] for r in records if r['outcome']=='transfer'}),
        hashes=hashes,script_sha256=x.sha(Path(__file__)))
    c=x.h.connection()
    try:
        with c.cursor() as q:
            q.execute('START TRANSACTION WITH CONSISTENT SNAPSHOT, READ ONLY');plan['before']=capture(q,plan)
    finally:c.rollback();c.close()
    x.h.require(not plan['before'][x.E] and not any(c['creator']==MARKER for c in plan['before'][x.C]),'Partial prior execution')
    x.h.require(not any(r['status']=='pending' and x.live(r) for r in plan['before'][REQUEST]),'Pending transfer needs resolution')
    eligible={r['user_id'] for r in plan['before']['planner_roles'] if x.live(r) and r['status']==0}
    x.h.require({r['owner_id'] for r in records if r['outcome']=='transfer'}<=eligible,'Target planner ineligible')
    plan['task_updates']=[]
    for task in plan['before'][x.T]:
        r=next((r for r in records if r['outcome']=='transfer' and r['service_id']==task['biz_id']),None)
        if r and x.live(task) and task['status']=='pending' and task['task_type'] in TASK_TYPES and task['assignee_id']!=r['owner_id']:
            plan['task_updates'].append(dict(id=task['id'],owner_id=r['owner_id']))
    x.h.require(all(x.sha(p)==hashes[p.name] for p in inputs),'Source changed')
    for p in inputs:
        dst=ROOT/p.name;x.h.require(not dst.exists(),'Backup exists');shutil.copy2(p,dst);dst.chmod(0o600)
    x.h.save(ROOT/'plan.json',plan)
    print(x.j(dict(transfers=19,closed_exclusions=2,held=1,new_classes=1,task_reassignments=len(plan['task_updates']),sha256=x.h.digest(plan))))


def audit(q,kind,object_id,key,refs,now):
    return x.insert(q,x.E,dict(event_type='system_data_repair',aggregate_type=kind,aggregate_id=object_id,
        operator_user_id=None,reason=x.REMARK,related_object_refs=x.j(refs),occurred_at=now,idempotency_key=key),now)


def verify(q,plan):
    current=capture(q,plan);before=plan['before'];events=current[x.E]
    x.h.require(len(events)==20,'Incomplete audit trail')
    created=[c for c in current[x.C] if c['creator']==MARKER];x.h.require(len(created)==1,'New class count differs')
    expectedclass=plan['classes'][0]
    for k,v in expectedclass.items():x.h.require(created[0][k]==(b'\0' if k=='system_class' else v),'Created class differs')
    tasks={t['id']:t for t in plan['task_updates']}
    targets={r['service_id']:r for r in plan['records'] if r['outcome']=='transfer'}
    for table in TABLES:
        if table==x.E:continue
        rows={r['id']:r for r in current[table]}
        x.h.require(set(rows)=={r['id'] for r in before[table]}|({created[0]['id']} if table==x.C else set()),'Unexpected row set: '+table)
        for old in before[table]:
            new=rows[old['id']];allowed=set()
            if table==x.S and old['id'] in targets:
                r=targets[old['id']];cid=created[0]['id'] if r.get('missing_class') else r['class_id']
                x.h.require(new['class_id']==cid and new['owner_user_id']==r['owner_id'] and new['version']==old['version']+1
                    and new['updater']==MARKER,'Assignment differs')
                allowed={'class_id','owner_user_id','version','updater','update_time'}
            if table==x.T and old['id'] in tasks:
                x.h.require(new['assignee_id']==tasks[old['id']]['owner_id'] and new['updater']==MARKER,'Task assignment differs')
                allowed={'assignee_id','updater','update_time'}
            x.h.require(all(x.h.encoded(new[k])==x.h.encoded(v) for k,v in old.items() if k not in allowed),'Protected fields differ: '+table)
    for key in ('users','planner_roles','departments'):x.h.require(x.h.encoded(current[key])==x.h.encoded(before[key]),'System configuration changed')
    for r in targets.values():
        event=[e for e in events if e['idempotency_key']==f'{MARKER}:service:{r["service_id"]}']
        x.h.require(len(event)==1 and event[0]['aggregate_id']==r['service_id'] and event[0]['operator_user_id'] is None
            and event[0]['reason']==x.REMARK and event[0]['from_status'] is None and event[0]['to_status'] is None,'Service audit differs')
        refs=json.loads(event[0]['related_object_refs']);x.h.require(refs['sourceRow']==r['row'] and refs['sourceSheet']==r['sheet']
            and refs['originalClassId']==r['old_class_id'] and refs['noNewService'],'Audit provenance differs')
    x.h.require(sum(e['idempotency_key']==MARKER+':class' for e in events)==1,'Class audit missing')
    hexes=x.rows(q,f'SELECT HEX(reason) h FROM {x.E} WHERE tenant_id=1 AND creator=%s',(MARKER,))
    x.h.require(all(v['h'].lower()==x.REMARK.encode().hex() for v in hexes),'Audit encoding differs')
    hexa=x.rows(q,f'SELECT HEX(class_name) h FROM {x.C} WHERE tenant_id=1 AND creator=%s',(MARKER,))[0]['h']
    x.h.require(hexa.lower()==expectedclass['class_name'].encode().hex(),'Class encoding differs')
    return current


def mutate(c,plan,fail_after=None):
    with c.cursor() as q:
        state=capture(q,plan,True)
        if state[x.E]:verify(q,plan);return 0
        x.h.require(x.h.encoded(state)==x.h.encoded(plan['before']),'Frozen preimage changed')
        now=x.rows(q,'SELECT NOW() now')[0]['now'];count=0
        def checkpoint():
            nonlocal count
            count+=1
            if count==fail_after:raise RuntimeError('Injected assignment failure')
        cid=x.insert(q,x.C,plan['classes'][0],now);checkpoint()
        audit(q,'delivery-class',cid,MARKER+':class',dict(sourceSha256=plan['hashes'][x.MASTER.name],historicalAssignment=True),now);checkpoint()
        before_services={r['id']:r for r in state[x.S]}
        for r in plan['records']:
            if r['outcome']!='transfer':continue
            old=before_services[r['service_id']];target=cid if r.get('missing_class') else r['class_id']
            q.execute(f'UPDATE {x.S} SET class_id=%s,owner_user_id=%s,version=version+1,updater=%s,update_time=%s '
                'WHERE tenant_id=1 AND id=%s AND version=%s AND deleted=0 AND teaching_delivery=1 AND teaching_closed=0',
                (target,r['owner_id'],MARKER,now,old['id'],old['version']))
            x.h.require(q.rowcount==1,'Transfer conflict');checkpoint()
            audit(q,'student-service',old['id'],f'{MARKER}:service:{old["id"]}',dict(sourceSheet=r['sheet'],sourceRow=r['row'],
                sourceSha256=plan['hashes'][x.MASTER.name],originalClassId=old['class_id'],targetClassId=target,
                originalOwnerId=old['owner_user_id'],targetOwnerId=r['owner_id'],noNewService=True,historyPreserved=True),now);checkpoint()
        for t in plan['task_updates']:
            q.execute(f'UPDATE {x.T} SET assignee_id=%s,updater=%s,update_time=%s WHERE tenant_id=1 AND id=%s AND deleted=0 AND status=%s',
                (t['owner_id'],MARKER,now,t['id'],'pending'));x.h.require(q.rowcount==1,'Pending task conflict');checkpoint()
        verify(q,plan);return count


def fixture(plan):
    c=x.h.connect('zsjos-mysql-1')
    try:
        with c.cursor() as q:
            q.execute('SET NAMES utf8mb4')
            for table in TABLES:
                ddl=x.rows(q,f'SHOW CREATE TABLE {table}')[0]['Create Table'];q.execute(ddl.replace('CREATE TABLE','CREATE TEMPORARY TABLE',1))
                cols=[r['COLUMN_NAME'] for r in x.rows(q,"SELECT COLUMN_NAME FROM information_schema.columns WHERE table_schema=DATABASE() AND table_name=%s AND EXTRA NOT LIKE '%%VIRTUAL GENERATED%%' AND EXTRA NOT LIKE '%%STORED GENERATED%%' ORDER BY ORDINAL_POSITION",(table,))]
                if plan['before'][table]:q.executemany(f'INSERT INTO {table} ('+','.join(cols)+') VALUES ('+','.join(['%s']*len(cols))+')',[tuple(r[k] for k in cols) for r in plan['before'][table]])
        c.commit()
        for point in (1,2,10,21,39):
            try:mutate(c,plan,point);raise AssertionError('Fault did not fire')
            except RuntimeError as e:x.h.require(str(e)=='Injected assignment failure','Unexpected fixture failure');c.rollback()
            with c.cursor() as q:x.h.require(x.h.encoded(capture(q,plan))==x.h.encoded(plan['before']),'Rollback differs')
            c.rollback()
        count=mutate(c,plan);c.commit();x.h.require(mutate(c,plan)==0,'Replay writes');c.rollback()
        with c.cursor() as q:
            q.execute(f'DELETE FROM {x.E} WHERE tenant_id=1 AND idempotency_key=%s',(MARKER+':class',))
            try:mutate(c,plan);raise AssertionError('Partial state accepted')
            except RuntimeError as e:x.h.require(str(e)=='Incomplete audit trail','Unexpected partial failure')
        c.rollback()
        with c.cursor() as q:verify(q,plan)
        x.h.save(ROOT/'fixture.json',dict(sha256=x.h.digest(plan),script_sha256=x.sha(Path(__file__)),writes=count,replay=0,fault_rollbacks=5))
        print(x.j(dict(fixture='PASS',writes=count,replay=0,fault_rollbacks=5,partial_rejected=True)))
    finally:c.rollback();c.close()


def apply(plan):
    ev=x.h.read(ROOT/'fixture.json');x.h.require(ev['sha256']==x.h.digest(plan) and ev['script_sha256']==x.sha(Path(__file__)),'Matching fixture required')
    x.h.require(all(x.sha(p)==plan['hashes'][p.name] for p in (x.MASTER,x.EXCEPTIONS,REVIEW,PENDING)),'Source changed')
    c=x.h.connect('zsjos-mysql-1')
    try:
        with c.cursor() as q:
            q.execute('SET NAMES utf8mb4');q.execute('SET SESSION innodb_lock_wait_timeout=5')
            x.h.require(x.rows(q,'SELECT GET_LOCK(%s,0) ok',(MARKER,))[0]['ok']==1,'Repair running')
        writes=mutate(c,plan);c.commit()
        with c.cursor() as q:after=verify(q,plan)
        x.h.save(ROOT/'after.json',after);print(x.j(dict(writes=writes,transfers=19,closed_preserved=2,new_classes=1,new_services=0)))
    finally:c.rollback();c.close()


def finish(plan):
    paths=(x.MASTER,x.EXCEPTIONS,REVIEW,PENDING,REPORT)
    if (ROOT/'files.json').exists():
        hashes=x.h.read(ROOT/'files.json');x.h.require(all(x.sha(p)==hashes[p.name] for p in paths),'Delivered file changed')
        print('File replay: no changes');return
    c=x.h.connection()
    try:
        with c.cursor() as q:
            q.execute('START TRANSACTION WITH CONSISTENT SNAPSHOT, READ ONLY');after=verify(q,plan)
    finally:c.rollback();c.close()
    x.h.save(ROOT/'independent-after.json',after)
    s=session()
    for r in plan['records']:
        if r['outcome']!='transfer':continue
        ctx=x.h.api(s,'GET',f'/zsjos/student/service/{r["service_id"]}/contact-context')
        x.h.require(ctx['acceptanceStatus']=='accepted' and ctx['deliveryStage']=='supervision','API stage changed')
    x.h.save(ROOT/'api.json',dict(contexts=19,supervision=True))
    x.h.require(all(x.sha(p)==plan['hashes'][p.name] for p in paths if p!=REPORT),'Workbook changed')
    master=openpyxl.load_workbook(x.MASTER);original=openpyxl.load_workbook(x.MASTER)
    exceptions=openpyxl.load_workbook(x.EXCEPTIONS);old_exc=openpyxl.load_workbook(x.EXCEPTIONS)
    review=openpyxl.load_workbook(REVIEW);old_review=openpyxl.load_workbook(REVIEW)
    done={(r['sheet'],r['row']):r for r in plan['records'] if r['outcome']!='hold'}
    excluded={(r['sheet'],r['row']) for r in plan['records'] if r['outcome']=='exclude'}
    report=openpyxl.Workbook();report.active.title='汇总'
    for row in [('处理项','数量','说明'),('已有服务按原表调整班级',19,'继续常规督学，保留原教学历史'),('异常结业保留系统并排除表格',2,'总表对应行清空；原始信息归档'),
                ('新增指定班主任同名班级',1,'复用已有同名班级考期配置'),('新增课程服务',0,'复用原服务，无重复导入'),('本次对比仍待确认',1,'普通班级班主任差异，未获处理决定'),('异常表剩余',20,'本次移出21条')]:report.active.append(row)
    full=report.create_sheet('完整原表与处理结果');headers=[c.value for c in original[plan['records'][0]['sheet']+'-学员导入填写'][1]]
    full.append(['总表工作表','总表行号','处理结果','备注']+headers)
    for i,r in enumerate(plan['records'],2):
        desc={'transfer':'已按原表调整现有服务班级/班主任','exclude':'保留系统异常结业，已清除导入表格','hold':'普通班级班主任差异待确认'}[r['outcome']]
        for col,v in enumerate([r['sheet'],r['row'],desc,x.REMARK if r['outcome']=='transfer' else '系统原状态和历史保留'],1):full.cell(i,col,v)
        source=original[r['sheet']+'-学员导入填写']
        for col in range(1,45):clone_cell(source.cell(r['row'],col),full.cell(i,col+4))
    for sh in report:x.format_sheet(sh)
    for sheet,row in excluded:
        for cell in master[sheet+'-学员导入填写'][row]:cell.value=None
    es=exceptions['异常与处理总表'];expected=[row for row in list(es.values)[1:] if (row[7],row[8]) not in done]
    for row in range(es.max_row,1,-1):
        if (es.cell(row,8).value,es.cell(row,9).value) in done:es.delete_rows(row)
    x.h.require(es.max_row==21,'Exception count differs');es.tables['ExceptionTable'].ref='A1:U21';es.auto_filter.ref='A1:U21'
    cs=exceptions['关联记录对比']
    for row in range(cs.max_row,1,-1):
        if (cs.cell(row,4).value,cs.cell(row,5).value) in done:cs.delete_rows(row)
    cs.tables['ComparisonTable'].ref=f'A1:K{cs.max_row}';cs.auto_filter.ref=cs.tables['ComparisonTable'].ref
    summary=exceptions['汇总']
    for row in range(2,summary.max_row+1):
        if summary.cell(row,1).value=='当前保留异常':summary.cell(row,2,20)
        if summary.cell(row,1).value=='当前异常构成':summary.cell(row,2,20);summary.cell(row,3,'1条普通班班主任差异、11条课程规格差异、8条多档案')
    summary.append(['2026-10-28考期班级处理',21,'19条按原表调整已有服务；2条异常结业保留系统并排除表格'])
    summary.tables['SummaryTable'].ref=f'A1:C{summary.max_row}';summary.auto_filter.ref=summary.tables['SummaryTable'].ref
    paired=review['逐行对比'];archived=review.create_sheet('已处理对照（处理前快照）')
    for col,c in enumerate(paired[1],1):clone_cell(c,archived.cell(1,col))
    archived.cell(1,22,'本次处理结果')
    for row in range(2,paired.max_row+1):
        key=(paired.cell(row,15).value,paired.cell(row,16).value)
        if key not in done:continue
        dest=archived.max_row+1
        for col,c in enumerate(paired[row],1):clone_cell(c,archived.cell(dest,col))
        archived.cell(dest,22,'按原表调整完成；保留教学历史' if done[key]['outcome']=='transfer' else '系统异常结业保持；表格已排除')
    for row in range(paired.max_row,1,-1):
        if (paired.cell(row,15).value,paired.cell(row,16).value) in done:paired.delete_rows(row)
    x.h.require(paired.max_row==3 and archived.max_row==43,'Paired partition differs')
    paired.tables['AssignmentPaired'].ref='A1:U3';paired.auto_filter.ref='A1:U3'
    x.format_sheet(archived);archived.freeze_panes='E2'
    rs=review['说明与汇总']
    for row in rs:
        if row[0].value=='待确认原表记录':row[1].value=1
        if row[0].value=='系统同课程同规格服务':row[1].value=1
        if row[0].value in ('班级不同','班级不同、班主任不同'):row[1].value=0
        if row[0].value=='本次操作':row[1].value='19条按原表调整；2条保留异常结业并排除表格；已处理对照为处理前快照，1条普通班主任差异待确认'
    rs.append(['处理记录','已处理21组移至已处理对照页，完整原表信息及原异常记录作为历史依据保留'])
    rs.auto_filter.ref=rs.dimensions
    books=((x.MASTER,master),(x.EXCEPTIONS,exceptions),(REVIEW,review),(REPORT,report))
    for path,book in books:
        temp=path.with_suffix('.tmp.xlsx');book.save(temp)
        with zipfile.ZipFile(temp) as z:x.h.require(z.testzip() is None,'Invalid workbook ZIP')
    saved={p:openpyxl.load_workbook(p.with_suffix('.tmp.xlsx')) for p,_ in books}
    for sh in original:
        for row in sh:
            for a in row:
                val=saved[x.MASTER][sh.title].cell(a.row,a.column).value
                x.h.require(val==(None if (sh.title.replace('-学员导入填写',''),a.row) in excluded else a.value),'Unexpected master edit')
    x.h.require(list(saved[x.EXCEPTIONS]['异常与处理总表'].values)[1:]==expected,'Retained exception rows differ')
    x.h.require(list(saved[x.EXCEPTIONS]['聊天处理依据'].values)==list(old_exc['聊天处理依据'].values),'Historical decisions differ')
    pending_expected=[r for r in list(old_review['逐行对比'].values)[1:] if (r[14],r[15]) not in done]
    x.h.require(list(saved[REVIEW]['逐行对比'].values)[1:]==pending_expected,'Remaining pair changed')
    for sh in ('原表完整信息','原异常记录'):x.h.require(list(saved[REVIEW][sh].values)==list(old_review[sh].values),'Review history changed')
    for row in list(saved[REPORT]['完整原表与处理结果'].rows)[1:]:
        src=original[row[0].value+'-学员导入填写']
        for col in range(1,45):
            a,b=src.cell(row[1].value,col),row[col+3]
            x.h.require((a.value,a.data_type,a.number_format)==(b.value,b.data_type,b.number_format),'Full source archive differs')
    x.h.require(all(x.sha(p)==plan['hashes'][p.name] for p in paths if p!=REPORT),'Source changed before replacement')
    for path,_ in books:path.with_suffix('.tmp.xlsx').replace(path)
    x.h.save(ROOT/'files.json',{p.name:x.sha(p) for p in paths})
    print('PASS:19 existing services assigned,2 closed exclusions,1comparison hold;20exceptions;2master rows cleared;original44fields archived')


def main():
    env=os.getenv('ZSJOS_AGENT_ENV')
    if env is None:env=Path('/etc/zsjos/agent-environment').read_text().strip()
    x.h.require(env=='test','Test only');os.umask(0o077);ROOT.mkdir(mode=0o700,parents=True,exist_ok=True)
    x.MARKER=MARKER;x.TABLES=TABLES
    parser=argparse.ArgumentParser();parser.add_argument('mode',choices=['prepare','fixture','apply','finish']);parser.add_argument('--sha256');a=parser.parse_args()
    if a.mode=='prepare':prepare();return
    plan=x.loads((ROOT/'plan.json').read_text());x.h.require(a.sha256==x.h.digest(plan),'Plan digest differs')
    if a.mode!='finish':x.h.require(plan['script_sha256']==x.sha(Path(__file__)),'Script changed')
    globals()[a.mode](plan)


if __name__=='__main__':main()
