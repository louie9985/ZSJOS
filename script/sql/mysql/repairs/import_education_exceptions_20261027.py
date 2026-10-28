#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""User-authorized tenant-1 historical exception import; never enable legacy SKUs.

Frozen inputs, private preimages, same-engine fixtures and guarded transactions.
No ordinary authorization rules or previous teaching services are rewritten.
See docs/operations/education-student-import-20261027.md for recovery boundaries.
"""
import argparse
from collections import Counter, defaultdict
from copy import copy
import datetime as dt
from decimal import Decimal
import hashlib
import json
import os
from pathlib import Path
import re
import shutil

import openpyxl

import import_education_students_20261027 as h
from advance_imported_students_20261027 import checklist, TYPES, REMARK
from qualify_education_leads_20261027 import loads

ROOT = Path('/opt/zsjos')
EVIDENCE = Path('/opt/zsjos-runtime/backups/education-exceptions-20261027')
MASTER = ROOT / 'output/教务学员导入总表.xlsx'
EXCEPTIONS = ROOT / 'output/教务学员_异常原因与已有处理汇总.xlsx'
REPORT = ROOT / 'output/教务学员_异常批次处理结果_20261027.xlsx'
MARKER = 'edu-exceptions-20261027'
P, S, C, K = 'zsjos_person', 'zsjos_service_relation', 'zsjos_delivery_class', 'zsjos_product_sku'
G, T, R, E = 'zsjos_collaboration_group', 'zsjos_business_task', 'zsjos_student_contact_record', 'zsjos_business_event'
CLAIM, COUNTER = 'zsjos_person_contact_claim', 'zsjos_person_no_daily_counter'
PROD, CAT, ATTR, VALUE = 'zsjos_product', 'zsjos_product_category', 'zsjos_product_attr', 'zsjos_product_attr_value'
CONFIG = 'zsjos_student_contact_config_version'
TABLES = (P, S, C, K, G, T, R, E, CLAIM, COUNTER, PROD, CAT, ATTR, VALUE, CONFIG)
CORRECTIONS = {441: (1, 1), 453: (4, 10), 454: (4, 10), 466: (4, 10)}
MASTER_HOLDS = {8: '总表同名不同联系方式待确认', 387: '总表同名不同联系方式待确认',
                467: '总表同名不同联系方式待确认', 452: '总表同课程不同规格待确认（同表112行）',
                134: '总表同联系方式不同姓名待确认（人社健康管理师166行）'}


def j(value):
    return json.dumps(value, ensure_ascii=False, separators=(',', ':'), default=lambda x: float(x) if isinstance(x, Decimal) else str(x))


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def rows(q, sql, args=()):
    q.execute(sql, args or None)
    return list(q.fetchall())


def norm(x):
    return str(x).strip() if x is not None else None


def live(row):
    return row['deleted'] in (0, b'\0', False)


def source_rows(path):
    w = openpyxl.load_workbook(path, data_only=True)
    return w, [dict(exception_row=i, name=norm(r[0]), mobile=norm(r[1]), state=r[2],
                    sheet=r[7], row=int(r[8]), product=r[9], sku=r[10], class_name=r[11],
                    planner=r[12], target=r[13], rules=r[17], issues=[])
               for i, r in enumerate(list(w['异常与处理总表'].values)[1:], 2) if r[0]]


def analyze():
    """Keep ambiguous identities/nearby specifications out of the authorized batch."""
    _, records = source_rows(EXCEPTIONS)
    before = h.snapshot()
    s = h.auth()
    options = h.api(s, 'GET', '/zsjos/student-entry/options')
    eligible_skus = {r['skuRef'] for r in options['courses']}
    eligible_classes = {r['id'] for r in options['classes']}
    master = openpyxl.load_workbook(MASTER, data_only=True)
    by_sheet = {s.title.split('-学员导入填写')[0]: s for s in master.worksheets[:8]}
    checks = {}
    for r in records:
        i = r['exception_row']
        raw = [norm(c.value) for c in by_sheet[r['sheet']][r['row']]][:7]
        if not any(raw):
            r['issues'].append('已从总表移除，已有服务待分班')
        else:
            expected = [r['name'], r['mobile'], None, r['product'], r['sku'], r['class_name'], r['planner']]
            if i in CORRECTIONS:
                raw[CORRECTIONS[i][0]] = expected[CORRECTIONS[i][0]]
            h.require(raw == expected, f'Unexpected master difference at exception row {i}')
        if i in MASTER_HOLDS:
            r['issues'].append(MASTER_HOLDS[i])
        if r['state'] == '仍需确认':
            r['issues'].append('原表仍需确认')
        if not r['mobile'] or not re.fullmatch(r'1[3-9]\d{9}', r['mobile']):
            r['issues'].append('联系方式格式待确认')
        products = [p for p in before['products'] if p['name'] == r['product'] and p['status'] == 0]
        h.require(len(products) == 1, 'Product missing/ambiguous')
        r['product_id'], r['product_ref'] = products[0]['id'], products[0]['product_ref']
        sku = [k for k in before['skus'] if k['spu_id'] == r['product_id'] and k['sku_name'] == r['sku']]
        if len(sku) > 1:
            r['issues'].append('规格不唯一')
        elif not sku:
            r['missing_sku'] = True
        else:
            r['sku_id'], r['skuRef'] = sku[0]['id'], sku[0]['sku_ref']
            if sku[0]['status'] != 0 or sku[0]['sku_ref'] not in eligible_skus:
                r['issues'].append('已有规格停用或不可选，需单独确认')
        users = [u for u in before['users'] if u['status'] == 0 and u['nickname'] == r['planner']]
        if len(users) != 1:
            r['issues'].append('规划师账号不唯一或不可用')
        else:
            r['owner_id'] = users[0]['id']
            classes = [c for c in before['classes'] if c['class_name'] == r['class_name']
                       and c['homeroom_user_id'] == r['owner_id'] and c['status'] == 'SERVING']
            if len(classes) > 1:
                r['issues'].append('班级不唯一')
            elif classes:
                r['class_id'] = classes[0]['id']
                if classes[0]['id'] not in eligible_classes:
                    r['issues'].append('班级不可入班')
            else:
                r['missing_class'] = True
        ids = {p['id'] for p in before['people'] if r['mobile'] in (p['mobile'], p['wechat_id'])}
        ids |= {c['person_id'] for c in before['claims'] if c['contact_value'] == r['mobile'] and c['person_id']}
        target = [p for p in before['people'] if p['person_no'] == r['target']] if r['target'] else []
        if r['target']:
            h.require(len(target) == 1, 'Confirmed target missing')
            ids.add(target[0]['id'])
        r['candidate_person_ids'] = sorted(ids)
        if len(ids) > 1:
            r['issues'].append('联系方式或指定档案涉及多档案')
        elif ids:
            r['person_id'] = next(iter(ids))
            person = next(p for p in before['people'] if p['id'] == r['person_id'])
            if person['name'].strip() != r['name']:
                r['issues'].append('档案姓名不一致')
            if r['mobile'] not in (person['mobile'], person['wechat_id']):
                r['issues'].append('指定档案联系方式不同')
        exact_identity = len(ids) == 1 and any(p['id'] in ids and p['name'].strip() == r['name']
                                              and r['mobile'] in (p['mobile'], p['wechat_id']) for p in before['people'])
        if not exact_identity and any(p['name'].strip() == r['name'] and p['id'] not in ids for p in before['people']):
            # Exact name/contact reuse was confirmed; a name alone is not identity.
            r['issues'].append('系统同名不同联系方式待确认')
        r['same_services'], r['other_services'] = [], []
        for rel in before['services']:
            if rel['person_id'] not in ids:
                continue
            snap = json.loads(rel['service_snapshot']) if rel['service_snapshot'] else {}
            if not snap:
                r['issues'].append('已有服务缺少课程快照')
                continue
            item = dict(id=rel['id'], product=snap.get('name') or snap.get('productName'),
                        sku=snap.get('skuName'), class_id=rel['class_id'], owner_id=rel['owner_user_id'], status=rel['status'])
            same = snap.get('productRef') == r['product_ref'] or item['product'] == r['product']
            r['same_services' if same else 'other_services'].append(item)
        if r['same_services']:
            r['issues'].append('已有同课程服务，不重复导入')
    groups = defaultdict(list)
    for r in records:
        groups[(r.get('person_id') or r['mobile'], r['product_id'])].append(r)
    for group in groups.values():
        if len(group) > 1:
            for r in group:
                r['issues'].append('异常表同档案或同号码同课程多行待确认')
    for r in records:
        if not r['issues']:
            if r['mobile'] not in checks:
                checks[r['mobile']] = h.api(s, 'POST', '/zsjos/student-entry/check', {'mobile': r['mobile'], 'wechatId': None})
            check = checks[r['mobile']]
            if check['status'] not in ('new', 'reuse') or check.get('personId') != r.get('person_id'):
                r['issues'].append('实时身份接口未通过')
            elif check['status'] == 'reuse' and check['name'].strip() != r['name']:
                r['issues'].append('实时姓名不一致')
    return records, before


def audit():
    records, _ = analyze()
    h.save(EVIDENCE / 'audit.json', records)
    ready = [r for r in records if not r['issues']]
    print(j({'total':len(records), 'ready':len(ready), 'new_classes':len({(r['class_name'],r['owner_id']) for r in ready if r.get('missing_class')}),
             'missing_sku_rows_ready':sum(bool(r.get('missing_sku')) for r in ready),
             'existing_same_product':sum(bool(r['same_services']) for r in records),
             'multi_course_ready':sum(bool(r['other_services']) for r in ready),
             'ready_refs':[(r['exception_row'],r['sheet'],r['row']) for r in ready],
             'issues':Counter(x for r in records for x in r['issues'])}))


def capture(q, plan, lock=False):
    state = {}
    for table in TABLES:
        condition, args = plan['scope'].get(table, ('1=1', []))
        state[table] = rows(q, f'SELECT * FROM {table} WHERE tenant_id=1 AND ({condition}) ORDER BY id'
                            + (' FOR UPDATE' if lock else ''), args)
    state['users'] = rows(q, 'SELECT id,nickname,status,dept_id,deleted FROM system_users WHERE tenant_id=1 '
                         'AND id IN (' + ','.join(map(str, plan['owners'] + [1])) + ') ORDER BY id')
    state['planner_roles'] = rows(q, "SELECT ur.user_id,ur.role_id,r.status,r.deleted FROM system_user_role ur "
                                  "JOIN system_role r ON r.id=ur.role_id AND r.tenant_id=ur.tenant_id "
                                  "WHERE ur.tenant_id=1 AND ur.deleted=0 AND r.code='study_planner' ORDER BY ur.user_id")
    state['departments'] = rows(q, 'SELECT id,name,status,deleted FROM system_dept WHERE tenant_id=1 ORDER BY id')
    return state


def course_snapshot(product, sku, state):
    categories = {c['id']: c for c in state[CAT] if live(c)}
    path, category_id = [], product['category_id']
    seen = set()
    while category_id:
        h.require(category_id not in seen and category_id in categories, 'Invalid category hierarchy')
        seen.add(category_id)
        cat = categories[category_id]
        path.insert(0, dict(id=cat['id'], name=cat['name']))
        category_id = cat['parent_id']
    h.require(path, 'Missing product category')
    attributes = {a['attr_key']: a for a in state[ATTR] if live(a) and a['spu_id'] == product['id']}
    specs = []
    for key, value in json.loads(sku['attr_values_json']).items():
        h.require(key in attributes, 'Missing product attribute')
        attr = attributes[key]
        labels = [v['label'] for v in state[VALUE] if live(v) and v['attr_id'] == attr['id'] and v['value'] == value]
        h.require(len(set(labels)) == 1, 'Missing or ambiguous attribute label')
        specs.append(dict(attrKey=key, attrName=attr['attr_name'], value=value, label=labels[0], labelMissing=False))
    return dict(productRef=product['product_ref'], name=product['name'], categoryId=path[-1]['id'],
                categoryName=path[-1]['name'], categoryPath=path, level1CategoryId=path[0]['id'],
                level1CategoryName=path[0]['name'], level2CategoryId=path[1]['id'] if len(path)>1 else None,
                level2CategoryName=path[1]['name'] if len(path)>1 else None,
                skuRef=sku['sku_ref'], skuName=sku['sku_name'], selectedAttrValuesJson=sku['attr_values_json'],
                price=sku['price'], spuUnknown=False, skuUnknown=False, specs=specs)


def sync_master():
    for path in (MASTER, EXCEPTIONS):
        dest = EVIDENCE / path.name
        if not dest.exists():
            shutil.copy2(path, dest)
            dest.chmod(0o600)
    baseline = openpyxl.load_workbook(EVIDENCE / MASTER.name)
    current = openpyxl.load_workbook(MASTER)
    exceptions = openpyxl.load_workbook(EVIDENCE / EXCEPTIONS.name, data_only=True)['异常与处理总表']
    edits = []
    for erow, (master_col, exception_col) in CORRECTIONS.items():
        raw = [c.value for c in exceptions[erow]]
        title, rowno = raw[7]+'-学员导入填写', int(raw[8])
        new = raw[exception_col]
        cell = current[title].cell(rowno, master_col+1)
        h.require(cell.value in (baseline[title].cell(rowno, master_col+1).value, new), 'Concurrent master edit')
        cell.value = new
        edits.append((title, rowno, master_col+1))
    for old_sheet, sheet in zip(baseline, current):
        h.require(old_sheet.title == sheet.title, 'Master sheet structure changed')
        for row in old_sheet:
            for cell in row:
                if (sheet.title, cell.row, cell.column) not in edits:
                    h.require(cell.value == sheet.cell(cell.row, cell.column).value, 'Unexpected concurrent master change')
    temp = MASTER.with_suffix('.tmp.xlsx')
    current.save(temp)
    openpyxl.load_workbook(temp).close()
    temp.replace(MASTER)
    h.save(EVIDENCE / 'master-sync.json', dict(cells=edits, before=sha(EVIDENCE/MASTER.name), after=sha(MASTER)))


def prepare(records=None, initial=None, *, sync=True, missing_sku_count=7, extra=None):
    h.require(not (EVIDENCE/'plan.json').exists(), 'Frozen plan already exists')
    if sync:
        sync_master()
    if records is None:
        records, initial = analyze()
    extra = extra or {}
    ready = [r for r in records if not r['issues']]
    h.require(ready, 'No eligible records')
    pids = sorted({r['person_id'] for r in ready if r.get('person_id')} |
                  {r['id'] for r in extra.get('name_updates', [])} |
                  {r['person_id'] for r in extra.get('service_name_updates', [])})
    mobiles = sorted({r['mobile'] for r in ready})
    names = sorted({r['name'] for r in ready})
    existing_services = [s['id'] for s in initial['services'] if s['person_id'] in pids]
    pid_sql = ','.join(map(str, pids)) or '-1'
    sid_sql = ','.join(map(str, existing_services)) or '-1'
    contact_marks = ','.join(['%s']*len(mobiles))
    name_marks = ','.join(['%s']*len(names))
    scope = {
        P: (f'id IN ({pid_sql}) OR mobile IN ({contact_marks}) OR wechat_id IN ({contact_marks}) OR name IN ({name_marks}) OR creator=%s', mobiles+mobiles+names+[MARKER]),
        S: (f'person_id IN ({pid_sql}) OR creator=%s', [MARKER]),
        CLAIM: (f'contact_value IN ({contact_marks}) OR creator=%s', mobiles+[MARKER]),
        G: (f'student_person_id IN ({pid_sql}) OR creator=%s', [MARKER]),
        T: (f"(biz_type='student_service' AND biz_id IN ({sid_sql})) OR creator=%s", [MARKER]),
        R: (f'service_relation_id IN ({sid_sql}) OR creator=%s', [MARKER]),
        E: ('creator=%s', [MARKER]),
        COUNTER: ('sequence_date=CURDATE()', []),
    }
    plan = dict(records=records, ready=ready, scope=scope, owners=sorted({r['owner_id'] for r in ready}),
                master_sha256=sha(MASTER), exception_sha256=sha(EXCEPTIONS), classes=[], skus=[], values=[])
    plan.update(extra)
    c = h.connection()
    try:
        with c.cursor() as q:
            q.execute('START TRANSACTION WITH CONSISTENT SNAPSHOT, READ ONLY')
            plan['before'] = capture(q, plan)
    finally:
        c.rollback(); c.close()
    state = plan['before']
    h.require(not state[E], 'Previous execution marker exists')
    h.require(all(not (r.get('creator') == MARKER) for t in TABLES for r in state[t]), 'Partial previous import')
    for r in ready:
        matched = {p['id'] for p in state[P] if live(p) and
                   (r['mobile']==p['mobile'] if plan.get('mobile_only') else r['mobile'] in (p['mobile'],p['wechat_id']))}
        claims = {a['person_id'] for a in state[CLAIM] if live(a) and a['contact_value']==r['mobile'] and a['person_id']}
        h.require(not claims-matched, 'Contact claim belongs to another identity')
        h.require(matched == ({r['person_id']} if r.get('person_id') else set()), 'Identity changed during preflight')
    configs = [f for f in state[CONFIG] if live(f) and f['status']=='published']
    h.require(len(configs)==1, 'Published contact configuration ambiguous')
    plan['config_id'], plan['checklist'] = configs[0]['id'], checklist(configs[0])
    eligible_owners = {r['user_id'] for r in state['planner_roles'] if live(r) and r['status']==0}
    h.require(set(plan['owners']) <= eligible_owners, 'Planner eligibility changed')
    users = {u['id']:u for u in state['users'] if live(u) and u['status']==0}
    h.require(set(plan['owners']+[1]) <= users.keys(), 'Executor/planner inactive')
    depts = {d['id']:d for d in state['departments'] if live(d) and d['status']==0}
    timestamp = dt.datetime.now().strftime('%Y%m%d%H%M%S')
    for index, (name, owner) in enumerate(sorted({(r['class_name'],r['owner_id']) for r in ready if r.get('missing_class')}),1):
        similar = [cl for cl in state[C] if live(cl) and cl['class_name']==name and cl['status']=='SERVING']
        dates = {(cl['exam_schedule_id'],cl['exam_schedule_snapshot']) for cl in similar}
        h.require(len(dates)<=1, 'Same-name class exam dates ambiguous')
        date_id, date_snapshot = next(iter(dates)) if dates else (None,None)
        user = users[owner]
        h.require(user['dept_id'] in depts, 'Planner department unavailable')
        plan['classes'].append(dict(class_no=f'BJ{timestamp}{index:04d}', class_name=name,
            system_class=0, class_kind='NORMAL', homeroom_user_id=owner,
            homeroom_user_name_snapshot=user['nickname'], dept_id=user['dept_id'],
            dept_name_snapshot=depts[user['dept_id']]['name'], status='SERVING', version=0,
            exam_schedule_id=date_id, exam_schedule_snapshot=date_snapshot))
    missing = sorted({(r['product_id'],r['sku']) for r in records if r.get('missing_sku')})
    h.require(len(missing)==missing_sku_count, 'Historical SKU scope changed')
    for index, (product_id, name) in enumerate(missing,1):
        product = next(p for p in state[PROD] if p['id']==product_id)
        attr = [a for a in state[ATTR] if live(a) and a['spu_id']==product_id and a['attr_key']=='sales_spec']
        h.require(len(attr)==1, 'Historical product attribute missing')
        value = name[len(product['name'])+1:] if name.startswith(product['name']+'｜') else name
        existing = [v for v in state[VALUE] if live(v) and v['attr_id']==attr[0]['id'] and v['value']==value]
        h.require(not existing, 'Historical attribute already exists; review required')
        plan['values'].append(dict(attr_id=attr[0]['id'], value=value, label=value, status=1, sort=999))
        attrs_json = j({'sales_spec': value})
        # Unit is copied from a real same-product specification with matching scope.
        suffix = name.rsplit('｜',1)[-1]
        samples = [k for k in state[K] if live(k) and k['spu_id']==product_id and k['sku_name'].endswith('｜'+suffix)]
        if not samples:
            samples = [k for k in state[K] if live(k) and k['spu_id']==product_id]
        units = {k['price_unit'] for k in samples}
        h.require(len(units)==1, 'Historical price unit ambiguous')
        plan['skus'].append(dict(spu_id=product_id, sku_ref=f'edu_hist_20261027_{index}', sku_name=name,
            attr_values_json=attrs_json, attr_values_hash=hashlib.sha256(attrs_json.encode()).hexdigest(),
            price=Decimal('0.00'), retail_price=Decimal('0.00'), min_deal_price=Decimal('0.00'),
            min_deal_type=None, min_deal_rate=None, exam_fee=Decimal('0.00'), price_unit=next(iter(units)),
            pricing_note='历史学员导入专用，不用于后续销售', status=1, sort=999, remark=REMARK))
    virtual = dict(state)
    virtual[VALUE] = state[VALUE]+[dict(v,deleted=0) for v in plan['values']]
    for r in ready:
        product = next(p for p in state[PROD] if p['id']==r['product_id'])
        sku = next(k for k in (plan['skus'] if r.get('missing_sku') else state[K])
                   if k['spu_id']==r['product_id'] and k['sku_name']==r['sku'] and (not 'deleted' in k or live(k)))
        r['skuRef'] = sku['sku_ref']
        r['snapshot'] = course_snapshot(product, sku, virtual)
        r['key'] = f'{MARKER}:row:{r["exception_row"]}'
    h.save(EVIDENCE/'plan.json',plan)
    print(j(dict(ready=len(ready), classes=len(plan['classes']), disabled_zero_skus=len(plan['skus']),
                 historical_services=sum(bool(r.get('missing_sku')) for r in ready), sha256=h.digest(plan))))


def insert(q, table, values, now):
    fields = dict(values, tenant_id=1, creator=MARKER, updater=MARKER, create_time=now, update_time=now)
    q.execute(f'INSERT INTO {table} ('+','.join(fields)+') VALUES ('+','.join(['%s']*len(fields))+')', tuple(fields.values()))
    h.require(q.rowcount==1, 'Insert count differs')
    return q.lastrowid


def event(q, kind, aggregate_id, key, refs, now):
    return insert(q, E, dict(event_type='system_data_repair', aggregate_type=kind,
        aggregate_id=aggregate_id, operator_user_id=None, reason=REMARK, occurred_at=now,
        to_status='supervision' if kind=='student-service' and not refs.get('nameCorrection') else None,
        related_object_refs=j(dict(refs, executor=MARKER, executorUserId=1)), idempotency_key=key), now)


def verify(q, plan):
    after = capture(q, plan)
    before = plan['before']
    own = {t:[r for r in after[t] if r.get('creator')==MARKER] for t in TABLES}
    n = len(plan['ready'])
    h.require(len(own[S])==len(own[G])==n and len(own[T])==len(own[R])==2*n, 'Incomplete service/task/record graph')
    h.require(len(own[C])==len(plan['classes']) and len(own[K])==len(plan['skus'])
              and len(own[VALUE])==len(plan['values']), 'Incomplete class/SKU definitions')
    name_updates={r['id']:r for r in plan.get('name_updates', [])}
    service_name_updates={r['id']:r for r in plan.get('service_name_updates', [])}
    h.require(len(own[E])==n+len(plan['classes'])+len(plan['skus'])+len(name_updates)+len(service_name_updates), 'Incomplete audit trail')
    # Existing rows are immutable except native identity promotion and daily numbering.
    promoted = {r['person_id'] for r in plan['ready'] if r.get('person_id')}
    event_map = {e['idempotency_key']:e for e in own[E]}
    service_map = {s['id']:s for s in own[S]}
    new_count = len({r['mobile'] for r in plan['ready'] if not r.get('person_id')})
    h.require(len(own[P])==new_count, 'Unexpected new identity count')
    for table in TABLES:
        current = {r['id']:r for r in after[table]}
        for old in before[table]:
            h.require(old['id'] in current, 'Protected row missing: '+table)
            found = current[old['id']]
            allowed = set()
            promotion=table==P and old['id'] in promoted and old['identity_status']!='student'
            rename=table==P and old['id'] in name_updates
            if promotion or rename:
                allowed = {'version','updater','update_time'}
                h.require(found['version']==old['version']+int(promotion)+int(rename) and found['updater']==MARKER, 'Identity version differs')
                if promotion:
                    allowed.add('identity_status')
                    h.require(found['identity_status']=='student', 'Identity transition differs')
                if rename:
                    allowed.add('name')
                    h.require(found['name']==name_updates[old['id']]['name'], 'Corrected person name differs')
            if table==S and old['id'] in service_name_updates:
                allowed={'student_name_snapshot','version','updater','update_time'}
                h.require(found['student_name_snapshot']==service_name_updates[old['id']]['name']
                          and found['version']==old['version']+1 and found['updater']==MARKER, 'Service name correction differs')
            if table==COUNTER:
                allowed={'current_value','update_time'}
                expected = (old['current_value']+new_count-1)%9999+1
                h.require(found['current_value']==expected, 'Daily counter differs')
            h.require(all(h.encoded(found[k])==h.encoded(v) for k,v in old.items() if k not in allowed), 'Protected fields changed: '+table)
        if table not in (P,COUNTER):
            h.require({r['id'] for r in after[table] if r.get('creator')!=MARKER}=={r['id'] for r in before[table]}, 'Unexpected external rows: '+table)
    for key in ('users','planner_roles','departments'):
        h.require(h.encoded(after[key])==h.encoded(before[key]), 'Protected system configuration changed')
    for table,changes,column in ((P,name_updates,'name'),(S,service_name_updates,'student_name_snapshot')):
        for rid,change in changes.items():
            audit=event_map.get(f'{MARKER}:name:{table}:{rid}')
            h.require(audit and audit['aggregate_id']==rid and audit['operator_user_id'] is None
                      and audit['reason']==REMARK and json.loads(audit['related_object_refs'])['nameCorrection'], 'Name correction audit missing')
            hexa=rows(q,f'SELECT HEX({column}) AS h FROM {table} WHERE tenant_id=1 AND id=%s',(rid,))[0]['h']
            h.require(hexa.lower()==change['name'].encode().hex(), 'Corrected name encoding differs')
    for sku in own[K]:
        expected = next(k for k in plan['skus'] if k['sku_ref']==sku['sku_ref'])
        h.require(all((json.loads(sku[k])==json.loads(v)) if k=='attr_values_json' else sku[k]==v for k,v in expected.items()), 'Historical SKU differs')
        h.require(sku['status']==1 and sku['price']==0, 'Historical SKU must remain disabled/zero')
        hexa=rows(q,f'SELECT HEX(sku_name) AS h FROM {K} WHERE tenant_id=1 AND id=%s',(sku['id'],))[0]['h']
        h.require(hexa.lower()==sku['sku_name'].encode().hex(), 'SKU encoding differs')
    for value in own[VALUE]:
        expected=next(v for v in plan['values'] if v['attr_id']==value['attr_id'] and v['value']==value['value'])
        h.require(all(value[k]==v for k,v in expected.items()), 'Disabled attribute differs')
    for cl in own[C]:
        expected=next(c for c in plan['classes'] if c['class_no']==cl['class_no'])
        h.require(all(cl[k]==(b'\0' if k=='system_class' else v) for k,v in expected.items()), 'Class differs')
        hexa=rows(q,f'SELECT HEX(class_name) AS h FROM {C} WHERE tenant_id=1 AND id=%s',(cl['id'],))[0]['h']
        h.require(hexa.lower()==cl['class_name'].encode().hex(), 'Class encoding differs')
    persons={p['id']:p for p in after[P]}
    classes={c['id']:c for c in after[C]}
    tasks={t['id']:t for t in own[T]}
    contacts={r['id']:r for r in own[R]}
    result=[]
    for source in plan['ready']:
        audit=event_map[source['key']]
        rel=service_map[audit['aggregate_id']]
        refs=json.loads(audit['related_object_refs'])
        person=persons[rel['person_id']]
        cl=classes[rel['class_id']]
        h.require(audit['operator_user_id'] is None and audit['reason']==REMARK and refs['historicalImport']
                  and refs['sourceRow']==source['row'] and refs['sourceSheet']==source['sheet']
                  and refs['masterSha256']==plan['master_sha256'], 'Service provenance differs')
        h.require(person['name'].strip()==source['name'] and source['mobile'] in (person['mobile'],person['wechat_id'])
                  and person['identity_status']=='student' and live(person), 'Imported identity differs')
        h.require(not source.get('person_id') or rel['person_id']==source['person_id'], 'Reused identity differs')
        claim=[a for a in after[CLAIM] if a['contact_value']==source['mobile'] and live(a)]
        h.require(len(claim)==1 and claim[0]['person_id']==person['id'] and claim[0]['reservation_key'] is None, 'Identity claim missing')
        h.require(rel['student_name_snapshot']==source['name'] and rel['student_name_source']=='independent_entry'
                  and rel['status']=='active' and rel['acceptance_status']=='accepted' and rel['delivery_stage']=='supervision'
                  and rel['accepted_by_user_id']==1 and rel['owner_user_id']==source['owner_id']
                  and rel['teaching_delivery']==b'\1' and rel['teaching_closed']==b'\0' and live(rel)
                  and rel['order_id'] is None and rel['order_item_id'] is None and rel['registration_case_id'] is None,
                  'Service state differs')
        h.require(cl['class_name']==source['class_name'] and cl['homeroom_user_id']==source['owner_id'], 'Service destination differs')
        h.require(json.loads(rel['service_snapshot'])==json.loads(j(source['snapshot']))
                  and json.loads(rel['first_contact_progress_json'])==plan['checklist'], 'Service snapshot differs')
        groups=[g for g in own[G] if g['id']==rel['collaboration_group_id'] and g['source_service_relation_id']==rel['id']
                and g['student_person_id']==person['id'] and g['status']=='active']
        h.require(len(groups)==1, 'Collaboration link missing')
        for kind,title,action in TYPES:
            ids=refs[kind]; task=tasks[ids['taskId']]; contact=contacts[ids['recordId']]
            h.require(task['biz_type']=='student_service' and task['biz_id']==rel['id'] and task['task_type']==kind
                      and task['status']=='completed' and task['assignee_id']==source['owner_id']
                      and task['summary_snapshot']==REMARK and task['title_snapshot']==title and task['action_code']==action
                      and task['completed_at']==rel['accepted_at'] and task['due_at'] is None and task['remind_at'] is None
                      and json.loads(task['payload'])['generationSource']=='historical_import', 'Completed task differs')
            h.require(contact['service_relation_id']==rel['id'] and contact['task_id']==task['id']
                      and contact['contact_type']==kind and contact['successful']==b'\1' and contact['remark']==REMARK
                      and contact['operator_user_id']==1 and contact['submitted_at']==rel['accepted_at']
                      and contact['next_contact_at'] is None and json.loads(contact['attachment_file_ids_json'])==[],
                      'Completion record differs')
            if kind==TYPES[0][0]:
                h.require(json.loads(contact['checklist_snapshot_json'])==plan['checklist']
                          and json.loads(contact['checklist_result_json'])==[x['key'] for x in plan['checklist']['items']], 'Checklist differs')
        # Audit and all completion facts record this operation, not invented earlier dates.
        h.require(rel['accepted_at']==audit['occurred_at']==rel['activated_at'], 'Completion timestamps differ')
        result.append(dict(exception_row=source['exception_row'],sheet=source['sheet'],row=source['row'],
                           person_id=person['id'],service_id=rel['id'],class_id=cl['id'],skuRef=source['skuRef']))
    hexes=rows(q,f'SELECT HEX(remark) AS h FROM {R} WHERE tenant_id=1 AND creator=%s',(MARKER,))
    h.require(len(hexes)==2*n and all(x['h'].lower()==REMARK.encode().hex() for x in hexes),'Remark encoding differs')
    names_hex=rows(q,f'SELECT id,HEX(student_name_snapshot) AS h FROM {S} WHERE tenant_id=1 AND creator=%s',(MARKER,))
    h.require(all(x['h'].lower()==service_map[x['id']]['student_name_snapshot'].encode().hex() for x in names_hex),'Name encoding differs')
    return after,result


def check_existing_courses(source, current):
    # A user-confirmed second enrollment may share a product but never a SKU.
    # Freeze the exact existing service IDs so confirmation cannot cover later drift.
    allowed = set(source.get('confirmed_same_product_ids', []))
    seen = set()
    for rel in current:
        snap = json.loads(rel['service_snapshot']) if rel['service_snapshot'] else {}
        h.require(snap, 'Existing service lacks course snapshot')
        if snap.get('productRef') == source['product_ref'] or snap.get('name', snap.get('productName')) == source['product']:
            h.require(rel['id'] in allowed and snap.get('skuRef') != source['skuRef']
                      and snap.get('skuName') != source['sku'], 'Same-product service appeared')
            seen.add(rel['id'])
    h.require(seen == allowed, 'Confirmed existing course scope changed')


def mutate(c, plan, fail_after=None):
    with c.cursor() as q:
        state=capture(q,plan,lock=True)
        if state[E]:
            verify(q,plan)
            return 0
        h.require(h.encoded(state)==h.encoded(plan['before']), 'Frozen preimage changed; refuse overwrite')
        now=rows(q,'SELECT NOW() AS now')[0]['now']
        count=0
        def checkpoint():
            nonlocal count
            count+=1
            if count==fail_after:
                raise RuntimeError('Injected exception-import failure')
        for table,changes,column in ((P,plan.get('name_updates',[]),'name'),(S,plan.get('service_name_updates',[]),'student_name_snapshot')):
            for change in changes:
                old=next(r for r in state[table] if r['id']==change['id'])
                q.execute(f'UPDATE {table} SET {column}=%s,version=version+1,updater=%s,update_time=%s '
                          'WHERE tenant_id=1 AND id=%s AND version=%s AND deleted=0',
                          (change['name'],MARKER,now,old['id'],old['version']))
                h.require(q.rowcount==1,'Name correction concurrent change'); checkpoint()
                event(q,'person' if table==P else 'student-service',old['id'],f'{MARKER}:name:{table}:{old["id"]}',
                      dict(nameCorrection=True,sourceSheet=change['sheet'],sourceRow=change['row'],sourceSha256=plan['master_sha256']),now); checkpoint()
                # Local working state carries the new version for a later student promotion.
                old[column]=change['name']; old['version']+=1
        class_ids={(c['class_name'],c['homeroom_user_id']):c['id'] for c in state[C] if live(c) and c['status']=='SERVING'}
        for fields in plan['classes']:
            key=(fields['class_name'],fields['homeroom_user_id'])
            h.require(key not in class_ids, 'Class already exists')
            cid=insert(q,C,fields,now); checkpoint()
            class_ids[key]=cid
            event(q,'delivery-class',cid,MARKER+':class:'+fields['class_no'],dict(sourceSha256=plan['exception_sha256']),now); checkpoint()
        for fields in plan['values']:
            insert(q,VALUE,fields,now); checkpoint()
        for fields in plan['skus']:
            kid=insert(q,K,fields,now); checkpoint()
            event(q,'product-sku',kid,MARKER+':sku:'+fields['sku_ref'],dict(disabled=True,price=0,sourceSha256=plan['exception_sha256']),now); checkpoint()
        created_people={}
        promoted=set()
        for source in plan['ready']:
            pid=source.get('person_id') or created_people.get(source['mobile'])
            if pid is None:
                q.execute(f"INSERT INTO {COUNTER} (sequence_date,current_value,creator,updater,tenant_id) VALUES (%s,1,%s,%s,1) "
                          'ON DUPLICATE KEY UPDATE current_value=IF(current_value>=9999,1,current_value+1),update_time=NOW()', (now.date(),MARKER,MARKER))
                sequence=rows(q,f'SELECT current_value FROM {COUNTER} WHERE tenant_id=1 AND sequence_date=%s FOR UPDATE',(now.date(),))[0]['current_value']
                checkpoint()
                pid=insert(q,P,dict(person_no='XY'+now.strftime('%Y%m%d%H%M%S')+f'{sequence:04d}',name=source['name'],
                    mobile=source['mobile'],wechat_id=None,identity_status='student',first_seen_at=now,last_seen_at=now,version=0),now)
                checkpoint(); created_people[source['mobile']]=pid
            elif pid not in promoted:
                person=next((p for p in state[P] if p['id']==pid),None)
                if person and person['identity_status']!='student':
                    q.execute(f"UPDATE {P} SET identity_status='student',version=version+1,updater=%s,update_time=%s "
                              'WHERE tenant_id=1 AND id=%s AND version=%s AND deleted=0',(MARKER,now,pid,person['version']))
                    h.require(q.rowcount==1,'Identity promotion conflict'); checkpoint()
                promoted.add(pid)
            claims=rows(q,f'SELECT * FROM {CLAIM} WHERE tenant_id=1 AND contact_value=%s FOR UPDATE',(source['mobile'],))
            if claims:
                h.require(len(claims)==1 and live(claims[0]) and claims[0]['person_id']==pid and claims[0]['reservation_key'] is None, 'Contact claim conflict')
            else:
                insert(q,CLAIM,dict(contact_value=source['mobile'],person_id=pid,reservation_key=None),now); checkpoint()
            current=rows(q,f'SELECT id,service_snapshot FROM {S} WHERE tenant_id=1 AND deleted=0 AND person_id=%s',(pid,))
            check_existing_courses(source,current)
            rid=insert(q,S,dict(person_id=pid,student_name_snapshot=source['name'],student_name_source='independent_entry',
                teaching_delivery=1,media_ready=0,class_id=class_ids[(source['class_name'],source['owner_id'])],
                status='active',owner_user_id=source['owner_id'],acceptance_status='accepted',accepted_by_user_id=1,
                accepted_at=now,director_stage='precheck',service_snapshot=j(source['snapshot']),delivery_stage='supervision',
                activated_at=now,version=0,teaching_closed=0,first_contact_progress_json=j(plan['checklist'])),now); checkpoint()
            gid=insert(q,G,dict(source_service_relation_id=rid,student_person_id=pid,status='active',version=0),now); checkpoint()
            q.execute(f'UPDATE {S} SET collaboration_group_id=%s WHERE tenant_id=1 AND id=%s AND creator=%s',(gid,rid,MARKER))
            h.require(q.rowcount==1,'Group link failed'); checkpoint()
            refs=dict(historicalImport=True,sourceSheet=source['sheet'],sourceRow=source['row'],exceptionOriginalRow=source['exception_row'],
                      masterSha256=plan['master_sha256'],exceptionSha256=plan['exception_sha256'],personId=pid,
                      reusedPerson=bool(source.get('person_id')),disabledHistoricalSku=bool(source.get('missing_sku')))
            for kind,title,action in TYPES:
                tid=insert(q,T,dict(task_type=kind,biz_type='student_service',biz_id=rid,status='completed',assignee_type='user',
                    assignee_id=source['owner_id'],title_snapshot=title,summary_snapshot=REMARK,action_code=action,
                    completed_at=now,version=1,payload=j(dict(serviceRelationId=rid,configVersionId=plan['config_id'],
                    generationSource='historical_import',executor=MARKER)),idempotency_key=source['key']+':'+kind),now); checkpoint()
                first=kind==TYPES[0][0]
                contact_id=insert(q,R,dict(service_relation_id=rid,task_id=tid,contact_type=kind,successful=1,remark=REMARK,
                    attachment_file_ids_json='[]',checklist_result_json=j([x['key'] for x in plan['checklist']['items']] if first else []),
                    checklist_snapshot_json=j(plan['checklist']) if first else None,study_plan_completed_at_first_contact=0 if first else None,
                    delivery_data_json='{}',operator_user_id=1,submitted_at=now,next_contact_at=None,
                    idempotency_key=source['key']+':record:'+kind,request_fingerprint=h.digest([MARKER,source['exception_row'],kind])),now); checkpoint()
                refs[kind]=dict(taskId=tid,recordId=contact_id)
            event(q,'student-service',rid,source['key'],refs,now); checkpoint()
        verify(q,plan)
        return count


def fixture(plan):
    c=h.connect('zsjos-mysql-1')
    try:
        with c.cursor() as q:
            q.execute('SET NAMES utf8mb4')
            for table in TABLES:
                ddl=rows(q,f'SHOW CREATE TABLE {table}')[0]['Create Table']
                q.execute(ddl.replace('CREATE TABLE','CREATE TEMPORARY TABLE',1))
                cols=[r['COLUMN_NAME'] for r in rows(q,"SELECT COLUMN_NAME FROM information_schema.columns WHERE table_schema=DATABASE() "
                    "AND table_name=%s AND EXTRA NOT LIKE '%%VIRTUAL GENERATED%%' AND EXTRA NOT LIKE '%%STORED GENERATED%%' ORDER BY ORDINAL_POSITION",(table,))]
                if plan['before'][table]:
                    q.executemany(f'INSERT INTO {table} ('+','.join(cols)+') VALUES ('+','.join(['%s']*len(cols))+')',
                                  [tuple(r[k] for k in cols) for r in plan['before'][table]])
        c.commit()
        for point in plan.get('fixture_failure_points',(1,30,51,100,1000)):
            try:
                mutate(c,plan,point)
                raise AssertionError('Fault injection did not fire')
            except RuntimeError as e:
                h.require(str(e)=='Injected exception-import failure','Unexpected fixture failure: '+str(e))
                c.rollback()
            with c.cursor() as q:
                h.require(h.encoded(capture(q,plan))==h.encoded(plan['before']),'Failed transaction left effects')
            c.rollback()
        count=mutate(c,plan); c.commit()
        h.require(mutate(c,plan)==0,'Replay wrote data'); c.rollback()
        with c.cursor() as q:
            q.execute(f'DELETE FROM {E} WHERE tenant_id=1 AND idempotency_key=%s',(plan['ready'][0]['key'],))
            try:
                mutate(c,plan)
                raise AssertionError('Partial state accepted')
            except RuntimeError as e:
                h.require(str(e)=='Incomplete audit trail','Unexpected partial-state failure')
        c.rollback()
        with c.cursor() as q:
            verify(q,plan)
        h.save(EVIDENCE/'fixture.json',dict(plan_sha256=h.digest(plan),script_sha256=sha(Path(__file__)),writes=count,
                                          failure_points=5,repeat=0,partial_rejected=True,recovery=True))
        print(j(dict(fixture='PASS',writes=count,fault_rollbacks=5,replay=0,partial_state_rejected=True)))
    finally:
        c.rollback(); c.close()


def apply(plan):
    evidence=loads((EVIDENCE/'fixture.json').read_text())
    h.require(evidence['plan_sha256']==h.digest(plan) and evidence['script_sha256']==sha(Path(__file__)),'Matching fixture required')
    h.require(sha(MASTER)==plan['master_sha256'] and sha(EXCEPTIONS)==plan['exception_sha256'],'Input workbook changed')
    c=h.connect('zsjos-mysql-1')
    try:
        with c.cursor() as q:
            q.execute('SET NAMES utf8mb4'); q.execute('SET SESSION innodb_lock_wait_timeout=5')
            h.require(rows(q,'SELECT GET_LOCK(%s,0) AS ok',(MARKER,))[0]['ok']==1,'Import already running')
        count=mutate(c,plan); c.commit()
        with c.cursor() as q:
            after,result=verify(q,plan)
        h.save(EVIDENCE/'after.json',after); h.save(EVIDENCE/'result.json',result)
        print(j(dict(writes=count,services=len(result),people=len({r['person_id'] for r in result}),classes=len(plan['classes']),disabled_skus=len(plan['skus']))))
    finally:
        c.rollback(); c.close()


def same_course_decision(source):
    same=source['same_services']
    if not same:
        return '其他问题待确认'
    exact=[e for e in same if e['sku']==source['sku']]
    if not exact:
        return '同课程规格不同，单独确认'
    if any(e['class_id']==source.get('class_id') and e['owner_id']==source.get('owner_id') for e in exact):
        return '已有同规格同班级同班主任，不重复导入'
    return '已有同规格，班级或班主任不同，不重复导入'


def format_sheet(sheet):
    from openpyxl.styles import Font, PatternFill, Alignment
    sheet.freeze_panes='A2'
    sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:
        cell.font=Font(bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='264478')
    for row in sheet:
        for cell in row:
            cell.alignment=Alignment(vertical='top',wrap_text=True)
    for col in sheet.columns:
        letter=col[0].column_letter
        sheet.column_dimensions[letter].width=min(65,max(14,len(str(col[0].value or ''))*2+3))


def finish_files(plan):
    """Only verified newly inserted services leave the exception list; keep all holds."""
    c=h.connection()
    try:
        with c.cursor() as q:
            q.execute('START TRANSACTION WITH CONSISTENT SNAPSHOT, READ ONLY')
            after,result=verify(q,plan)
    finally:
        c.rollback(); c.close()
    h.require(sha(MASTER)==plan['master_sha256'],'Master changed after import')
    done={r['exception_row'] for r in result}
    remaining=[r for r in plan['records'] if r['exception_row'] not in done]
    h.require(len(done)==len(plan['ready']) and len(done)+len(remaining)==494,'Exception partition differs')
    if (EVIDENCE/'files-result.json').exists():
        old=loads((EVIDENCE/'files-result.json').read_text())
        h.require(sha(EXCEPTIONS)==old['exceptions_sha256'] and sha(REPORT)==old['report_sha256'],'Delivered files changed; refuse overwrite')
        print('Workbook replay: no changes')
        return
    h.require(sha(EXCEPTIONS)==plan['exception_sha256'],'Exception workbook changed; refuse overwrite')
    original=openpyxl.load_workbook(EVIDENCE/EXCEPTIONS.name)
    target=openpyxl.load_workbook(EXCEPTIONS)
    sheet=target['异常与处理总表']
    for r in remaining:
        row=r['exception_row']
        sheet.cell(row,3, same_course_decision(r) if r['same_services'] else '待逐项确认')
        sheet.cell(row,4,'；'.join(r['issues']))
        action='核对档案或报名关系后处理'
        if r['same_services']:
            action=same_course_decision(r)
        sheet.cell(row,6,action)
        previous=sheet.cell(row,7).value
        sheet.cell(row,7,(str(previous)+'\n' if previous else '')+'2026-10-27：本批实时复核，未新增；具体对照见异常批次处理结果。')
    for row in sorted(done,reverse=True):
        sheet.delete_rows(row)
    sheet.tables['ExceptionTable'].ref=f'A1:U{len(remaining)+1}'
    sheet.auto_filter.ref=sheet.tables['ExceptionTable'].ref
    refs={(r['sheet'],r['row']) for r in result}
    comparison=target['关联记录对比']
    removed_comparison=0
    for row in range(comparison.max_row,1,-1):
        if (comparison.cell(row,4).value,comparison.cell(row,5).value) in refs:
            comparison.delete_rows(row); removed_comparison+=1
    comparison.tables['ComparisonTable'].ref=f'A1:K{comparison.max_row}'
    comparison.auto_filter.ref=comparison.tables['ComparisonTable'].ref
    summary=target['汇总']
    summary.delete_rows(1,summary.max_row)
    summary.append(['项目','数量','说明'])
    summary_rows=[
        ('处理前异常课程记录',494,'以原始异常工作表为准，备份保留'),
        ('本批成功导入并移出',len(result),'全部进入常规督学，首联及学习计划已完成'),
        ('当前保留待处理记录',len(remaining),'仅删除本批经数据库核验成功的新导入记录'),
        ('已有同课程服务',sum(bool(r['same_services']) for r in remaining),'未重复新增；身份、规格及归属问题保留'),
        ('已有同课程但无相同规格',sum(bool(r['same_services']) and not any(e['sku']==r['sku'] for e in r['same_services']) for r in remaining),'需单独确认规格关系'),
        ('其他保留记录',sum(not r['same_services'] for r in remaining),'身份、联系方式、表内重复、规划师等问题'),
        ('本批新增班级',len(plan['classes']),'按表内班主任创建同名独立班级，不改旧班'),
        ('本批新增历史规格',len(plan['skus']),'全部0元、停用；属性选项同时停用'),
        ('总表同步修正',4,'已确认的1处联系方式及3处课程规格'),
        ('关联对比移出记录',removed_comparison,'仅移出对应本批成功导入的来源记录'),
        ('行号规则',None,'原表行号仍指总表实际行；异常表删行后当前行号变化'),
        ('处理结果文件',None,REPORT.name),
        ('历史依据',None,'聊天处理依据保留原文，仅作历史依据；最新状态以本次汇总为准'),
    ]
    for row in summary_rows: summary.append(row)
    summary.tables['SummaryTable'].ref=f'A1:C{summary.max_row}'
    format_sheet(summary)
    report=openpyxl.Workbook()
    report.active.title='汇总'
    report.active.append(['项目','数量','说明'])
    for row in summary_rows: report.active.append(row)
    imported=report.create_sheet('本批已导入')
    imported.append(['总表工作表','总表行号','处理前异常表行号','课程','规格','班级','身份处理','当前阶段','完成事项','备注'])
    for r in plan['ready']:
        imported.append([r['sheet'],r['row'],r['exception_row'],r['product'],r['sku'],r['class_name'],
                         '复用档案' if r.get('person_id') else '新档案','常规督学','首联、学习计划均完成',REMARK])
    compare=report.create_sheet('已有课程差异')
    compare.append(['总表工作表','总表行号','处理前异常表行号','表内课程','表内规格','已有课程','已有规格',
                    '表内班级','已有班级','班主任一致','处理结论','其他待确认事项'])
    classes={c['id']:c for c in after[C]}
    for r in remaining:
        if r['same_services']:
            compare.append([r['sheet'],r['row'],r['exception_row'],r['product'],r['sku'],
                '\n'.join(dict.fromkeys(e['product'] or '快照缺失' for e in r['same_services'])),
                '\n'.join(dict.fromkeys(e['sku'] or '快照缺失' for e in r['same_services'])),r['class_name'],
                '\n'.join(dict.fromkeys(classes.get(e['class_id'],{}).get('class_name') or '未分班' for e in r['same_services'])),
                '全部一致' if all(e['owner_id']==r.get('owner_id') for e in r['same_services']) else '存在差异',
                same_course_decision(r),'；'.join(x for x in r['issues'] if x!='已有同课程服务，不重复导入')])
    held=report.create_sheet('剩余待处理')
    held.append(['总表工作表','总表行号','处理前异常表行号','课程','规格','本次处理','待确认事项'])
    for r in remaining:
        held.append([r['sheet'],r['row'],r['exception_row'],r['product'],r['sku'],'未新增，保留异常表','；'.join(r['issues'])])
    classes_sheet=report.create_sheet('新增班级')
    classes_sheet.append(['班级名称','班主任依据（总表）','考期'])
    for cl in plan['classes']:
        source=next(r for r in plan['ready'] if r['class_name']==cl['class_name'] and r['owner_id']==cl['homeroom_user_id'])
        classes_sheet.append([cl['class_name'],f'{source["sheet"]} 第{source["row"]}行规划师',cl['exam_schedule_snapshot'] or '空'])
    sku_sheet=report.create_sheet('历史停用规格')
    sku_sheet.append(['课程规格','价格','状态','用途'])
    for sku in plan['skus']:sku_sheet.append([sku['sku_name'],0,'停用','历史导入专用，禁止后续销售使用'])
    for ws in report:format_sheet(ws)
    report_tmp=REPORT.with_suffix('.tmp.xlsx'); report.save(report_tmp)
    exceptions_tmp=EXCEPTIONS.with_suffix('.tmp.xlsx'); target.save(exceptions_tmp)
    checked=openpyxl.load_workbook(exceptions_tmp,data_only=True)
    h.require(checked['异常与处理总表'].max_row==len(remaining)+1,'Remaining Excel count differs')
    expected_keys=[(r['sheet'],r['row']) for r in remaining]
    h.require([(r[7],int(r[8])) for r in list(checked['异常与处理总表'].values)[1:]]==expected_keys,'Remaining source references differ')
    protected_columns=[1,2,5,8,9,10,11,12,13,14,15,16,17,18,19,20,21]
    for newrow,r in enumerate(remaining,2):
        for col in protected_columns:
            h.require(checked['异常与处理总表'].cell(newrow,col).value==original['异常与处理总表'].cell(r['exception_row'],col).value,'Remaining source data changed')
    h.require(list(checked['聊天处理依据'].values)==list(original['聊天处理依据'].values),'Historical decisions changed')
    report_check=openpyxl.load_workbook(report_tmp,data_only=True)
    h.require(report_check['本批已导入'].max_row==len(result)+1 and report_check['剩余待处理'].max_row==len(remaining)+1,'Report partition differs')
    h.require(report_check['已有课程差异'].max_row==135,'Existing service report count differs')
    report_tmp.replace(REPORT)
    exceptions_tmp.replace(EXCEPTIONS)
    h.save(EVIDENCE/'files-result.json',dict(exceptions_sha256=sha(EXCEPTIONS),report_sha256=sha(REPORT),
        master_sha256=sha(MASTER),removed=len(done),remaining=len(remaining),comparison_removed=removed_comparison))
    print(j(dict(workbooks='PASS',removed=len(done),remaining=len(remaining),comparison_removed=removed_comparison)))


def main():
    env = os.getenv('ZSJOS_AGENT_ENV')
    env = env if env is not None else Path('/etc/zsjos/agent-environment').read_text().strip()
    h.require(env == 'test', 'Test-only execution')
    os.umask(0o077)
    EVIDENCE.mkdir(mode=0o700, parents=True, exist_ok=True)
    parser = argparse.ArgumentParser()
    parser.add_argument('mode', choices=['audit','prepare','fixture','apply','verify','finish-files'])
    parser.add_argument('--sha256')
    args = parser.parse_args()
    if args.mode in ('audit','prepare'):
        globals()[args.mode]()
        return
    plan=loads((EVIDENCE/'plan.json').read_text())
    h.require(args.sha256==h.digest(plan),'Exact frozen plan hash required')
    if args.mode=='finish-files':
        finish_files(plan)
    elif args.mode in ('fixture','apply'):
        globals()[args.mode](plan)
    else:
        c=h.connection()
        try:
            with c.cursor() as q:
                q.execute('START TRANSACTION WITH CONSISTENT SNAPSHOT, READ ONLY')
                after,result=verify(q,plan)
            h.save(EVIDENCE/'independent-after.json',after)
            print(j(dict(verified=len(result),completed_tasks=2*len(result),exact_remarks=2*len(result))))
        finally:
            c.rollback(); c.close()


if __name__ == '__main__':
    main()
