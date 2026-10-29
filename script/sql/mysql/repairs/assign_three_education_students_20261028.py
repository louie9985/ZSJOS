#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Assign three user-confirmed historical services; no new service or stage reset."""
import argparse
import json
import os
from pathlib import Path
import shutil
import zipfile

import openpyxl
import import_education_exceptions_20261027 as x
from import_two_confirmed_students_20261028 import session
from archive_matching_education_students_20261028 import clone_cell, PENDING
from export_education_assignment_comparison_20261028 import OUTPUT as REVIEW

ROOT=Path('/opt/zsjos-runtime/backups/education-three-assignment-20261028')
MARKER='edu-three-assignment-20261028'
REPORT=x.ROOT/'output/教务学员_三条分班处理结果_20261028.xlsx'
TARGETS={('中医营养配餐',17),('中医营养配餐',18),('中医营养配餐',22)}
REQUEST='zsjos_class_transfer_request'
TABLES=tuple(t for t in x.TABLES if t!=x.COUNTER)+(REQUEST,)
INPUTS=(x.MASTER,x.EXCEPTIONS,PENDING,REVIEW)


def capture(q,plan,lock=False):return x.capture(q,plan,lock)


def prepare():
    x.h.require(not (ROOT/'plan.json').exists(),'Frozen plan exists')
    _,allrows=x.source_rows(x.EXCEPTIONS);records=[r for r in allrows if (r['sheet'],r['row']) in TARGETS]
    x.h.require(len(allrows)==17 and len(records)==3,'Source scope changed')
    hashes={p.name:x.sha(p) for p in INPUTS};st=x.h.snapshot();classes={c['id']:c for c in st['classes']}
    master=openpyxl.load_workbook(x.MASTER)
    for r in records:
        claims=[c for c in st['claims'] if c['contact_value']==r['mobile']]
        x.h.require(len(claims)==1 and claims[0]['person_id'] and claims[0]['reservation_key'] is None,'Ambiguous claim')
        ps=[p for p in st['people'] if p['id']==claims[0]['person_id']]
        x.h.require(len(ps)==1 and ps[0]['mobile']==r['mobile'] and ps[0]['name']==r['name'],'Claim identity differs')
        owner=[u for u in st['users'] if u['nickname']==r['planner'] and u['status']==0]
        x.h.require(len(owner)==1,'Owner unavailable')
        target=[c for c in st['classes'] if c['class_name']==r['class_name'] and c['homeroom_user_id']==owner[0]['id'] and c['status']=='SERVING' and c['class_kind']=='NORMAL']
        x.h.require(len(target)==1,'Target class unavailable')
        products={p['product_ref'] for p in st['products'] if p['name']==r['product']}
        services=[]
        for s in st['services']:
            snap=json.loads(s['service_snapshot'] or '{}')
            if s['person_id']==ps[0]['id'] and snap.get('productRef') in products and snap.get('name',snap.get('productName'))==r['product'] and snap.get('skuName')==r['sku']:services.append(s)
        x.h.require(len(services)==1,'Course service ambiguous');s=services[0]
        x.h.require(classes[s['class_id']]['class_kind']=='UNSCHEDULED' and s['owner_user_id']==owner[0]['id'] and s['student_name_snapshot']==r['name']
            and s['status']=='active' and s['teaching_delivery']==b'\1' and s['teaching_closed']==b'\0' and s['acceptance_status']=='accepted' and s['delivery_stage']=='supervision','Service lifecycle/assignment differs')
        x.h.require(not any(c.value is not None for c in master[r['sheet']+'-学员导入填写'][r['row']]),'Previously empty master row changed')
        r.update(person_id=ps[0]['id'],service_id=s['id'],claim_id=claims[0]['id'],owner_id=owner[0]['id'],class_id=target[0]['id'],old_class_id=s['class_id'])
    pids=','.join(map(str,sorted({r['person_id'] for r in records})));sids=','.join(map(str,sorted({r['service_id'] for r in records})))
    mobiles=[r['mobile'] for r in records];marks=','.join(['%s']*len(mobiles))
    scope={x.P:(f'id IN ({pids}) OR mobile IN ({marks}) OR wechat_id IN ({marks})',mobiles+mobiles),x.S:(f'person_id IN ({pids})',[]),
        x.G:(f'student_person_id IN ({pids})',[]),x.T:(f"biz_type='student_service' AND biz_id IN ({sids})",[]),x.R:(f'service_relation_id IN ({sids})',[]),
        x.CLAIM:(f'contact_value IN ({marks})',mobiles),x.E:('creator=%s',[MARKER]),REQUEST:(f'service_relation_id IN ({sids})',[])}
    plan=dict(records=records,scope=scope,owners=sorted({r['owner_id'] for r in records}),hashes=hashes,script_sha256=x.sha(Path(__file__)))
    c=x.h.connection()
    try:
        with c.cursor() as q:q.execute('START TRANSACTION WITH CONSISTENT SNAPSHOT, READ ONLY');plan['before']=capture(q,plan)
    finally:c.rollback();c.close()
    eligible={r['user_id'] for r in plan['before']['planner_roles'] if x.live(r) and r['status']==0}
    x.h.require(set(plan['owners'])<=eligible and not plan['before'][x.E],'Planner or prior execution differs')
    x.h.require(not any(x.live(r) and r['status']=='pending' for r in plan['before'][REQUEST]),'Pending class transfer')
    x.h.require(all(x.sha(p)==hashes[p.name] for p in INPUTS),'Source changed')
    for p in (x.MASTER,x.EXCEPTIONS):
        dst=ROOT/p.name;x.h.require(not dst.exists(),'Backup exists');shutil.copy2(p,dst);dst.chmod(0o600)
    x.h.save(ROOT/'plan.json',plan);print(x.j(dict(transfers=3,new_services=0,owner_changes=0,sha256=x.h.digest(plan))))


def verify(q,plan):
    current=capture(q,plan);targets={r['service_id']:r for r in plan['records']};events=current[x.E]
    x.h.require(len(events)==3,'Incomplete audit trail')
    for table in TABLES:
        if table==x.E:continue
        before=plan['before'][table];now={r['id']:r for r in current[table]}
        x.h.require(set(now)=={r['id'] for r in before},'Unexpected row set: '+table)
        for old in before:
            actual=now[old['id']];allowed=set()
            if table==x.S and old['id'] in targets:
                r=targets[old['id']];allowed={'class_id','version','updater','update_time'}
                x.h.require(actual['class_id']==r['class_id'] and actual['version']==old['version']+1 and actual['updater']==MARKER,'Assignment differs')
            x.h.require(all(x.h.encoded(actual[k])==x.h.encoded(v) for k,v in old.items() if k not in allowed),'Protected fields changed: '+table)
    for key in ('users','planner_roles','departments'):x.h.require(x.h.encoded(current[key])==x.h.encoded(plan['before'][key]),'System configuration changed')
    for r in plan['records']:
        matches=[e for e in events if e['idempotency_key']==f'{MARKER}:service:{r["service_id"]}']
        x.h.require(len(matches)==1,'Audit key missing');e=matches[0];refs=json.loads(e['related_object_refs'])
        x.h.require(e['aggregate_id']==r['service_id'] and e['operator_user_id'] is None and e['reason']==x.REMARK and e['from_status'] is None and e['to_status'] is None
            and refs['sourceSheet']==r['sheet'] and refs['sourceRow']==r['row'] and refs['targetClassId']==r['class_id'] and refs['originalClassId']==r['old_class_id'],'Audit provenance differs')
    hexes=x.rows(q,f'SELECT HEX(reason) h FROM {x.E} WHERE tenant_id=1 AND creator=%s',(MARKER,))
    x.h.require(all(h['h'].lower()==x.REMARK.encode().hex() for h in hexes),'Audit encoding differs');return current


def mutate(c,plan,fail_after=None):
    with c.cursor() as q:
        before=capture(q,plan,True)
        if before[x.E]:verify(q,plan);return 0
        x.h.require(x.h.encoded(before)==x.h.encoded(plan['before']),'Frozen preimage changed')
        now=x.rows(q,'SELECT NOW() now')[0]['now'];count=0
        def point():
            nonlocal count
            count+=1
            if count==fail_after:raise RuntimeError('Injected assignment failure')
        for r in plan['records']:
            old=next(s for s in before[x.S] if s['id']==r['service_id'])
            q.execute(f'UPDATE {x.S} SET class_id=%s,version=version+1,updater=%s,update_time=%s WHERE tenant_id=1 AND id=%s AND version=%s AND class_id=%s AND deleted=0 AND teaching_closed=0 AND teaching_delivery=1',
                (r['class_id'],MARKER,now,r['service_id'],old['version'],r['old_class_id']));x.h.require(q.rowcount==1,'Service conflict');point()
            x.insert(q,x.E,dict(event_type='system_data_repair',aggregate_type='student-service',aggregate_id=r['service_id'],operator_user_id=None,reason=x.REMARK,
                occurred_at=now,idempotency_key=f'{MARKER}:service:{r["service_id"]}',related_object_refs=x.j(dict(sourceSheet=r['sheet'],sourceRow=r['row'],
                sourceSha256=plan['hashes'][x.EXCEPTIONS.name],originalClassId=r['old_class_id'],targetClassId=r['class_id'],claimId=r['claim_id'],noNewService=True,historyPreserved=True))),now);point()
        verify(q,plan);return count


def fixture(plan):
    c=x.h.connect('zsjos-mysql-1')
    try:
        with c.cursor() as q:
            q.execute('SET NAMES utf8mb4')
            for table in TABLES:
                ddl=x.rows(q,f'SHOW CREATE TABLE {table}')[0]['Create Table'];q.execute(ddl.replace('CREATE TABLE','CREATE TEMPORARY TABLE',1))
                cols=[r['COLUMN_NAME'] for r in x.rows(q,"SELECT COLUMN_NAME FROM information_schema.columns WHERE table_schema=DATABASE() AND table_name=%s AND EXTRA NOT LIKE '%%VIRTUAL GENERATED%%' AND EXTRA NOT LIKE '%%STORED GENERATED%%' ORDER BY ORDINAL_POSITION",(table,))]
                if plan['before'][table]:q.executemany(f'INSERT INTO {table} ('+','.join(cols)+') VALUES ('+','.join(['%s']*len(cols))+')',[tuple(r[k] for k in cols) for r in plan['before'][table]])
        c.commit()
        for point in (1,2,3,4,6):
            try:mutate(c,plan,point);raise AssertionError('Fault did not fire')
            except RuntimeError as e:x.h.require(str(e)=='Injected assignment failure','Unexpected fixture failure');c.rollback()
            with c.cursor() as q:x.h.require(x.h.encoded(capture(q,plan))==x.h.encoded(plan['before']),'Rollback differs')
            c.rollback()
        count=mutate(c,plan);c.commit();x.h.require(mutate(c,plan)==0,'Replay writes');c.rollback()
        with c.cursor() as q:
            q.execute(f'DELETE FROM {x.E} WHERE tenant_id=1 AND idempotency_key=%s',(f'{MARKER}:service:{plan["records"][0]["service_id"]}',))
            try:mutate(c,plan);raise AssertionError('Partial state accepted')
            except RuntimeError as e:x.h.require(str(e)=='Incomplete audit trail','Unexpected partial failure')
        c.rollback()
        with c.cursor() as q:verify(q,plan)
        x.h.save(ROOT/'fixture.json',dict(sha256=x.h.digest(plan),script_sha256=x.sha(Path(__file__)),writes=count,replay=0,fault_rollbacks=5))
        print(x.j(dict(fixture='PASS',writes=count,replay=0,fault_rollbacks=5,partial_rejected=True)))
    finally:c.rollback();c.close()


def apply(plan):
    ev=x.h.read(ROOT/'fixture.json');x.h.require(ev['sha256']==x.h.digest(plan) and ev['script_sha256']==x.sha(Path(__file__)),'Matching fixture required')
    x.h.require(all(x.sha(p)==plan['hashes'][p.name] for p in (x.MASTER,x.EXCEPTIONS,REVIEW,PENDING)),'Source changed')
    c=x.h.connect('zsjos-mysql-1')
    try:
        with c.cursor() as q:
            q.execute('SET NAMES utf8mb4');q.execute('SET SESSION innodb_lock_wait_timeout=5')
            x.h.require(x.rows(q,'SELECT GET_LOCK(%s,0) ok',(MARKER,))[0]['ok']==1,'Repair running')
        writes=mutate(c,plan);c.commit()
        with c.cursor() as q:after=verify(q,plan)
        x.h.save(ROOT/'after.json',after);print(x.j(dict(writes=writes,transfers=3,new_classes=0,new_services=0)))
    finally:c.rollback();c.close()


def finish(plan):
    if (ROOT/'files.json').exists():
        hashes=x.h.read(ROOT/'files.json');x.h.require(all(x.sha(p)==hashes[p.name] for p in (*INPUTS,REPORT)),'Delivered files changed')
        print('File replay: no changes');return
    c=x.h.connection()
    try:
        with c.cursor() as q:q.execute('START TRANSACTION WITH CONSISTENT SNAPSHOT, READ ONLY');after=verify(q,plan)
    finally:c.rollback();c.close()
    x.h.save(ROOT/'independent-after.json',after)
    s=session()
    for r in plan['records']:
        ctx=x.h.api(s,'GET',f'/zsjos/student/service/{r["service_id"]}/contact-context')
        x.h.require(ctx['acceptanceStatus']=='accepted' and ctx['deliveryStage']=='supervision','API lifecycle differs')
    x.h.save(ROOT/'api.json',dict(contexts=3,passed=True))
    x.h.require(all(x.sha(p)==plan['hashes'][p.name] for p in INPUTS),'Source changed')
    exc=openpyxl.load_workbook(x.EXCEPTIONS);original=openpyxl.load_workbook(x.EXCEPTIONS)
    out=openpyxl.Workbook();out.active.title='处理结果'
    out.active.append(['总表工作表','总表行号','原班级','调整后班级','处理结果','备注'])
    for r in plan['records']:out.active.append([r['sheet'],r['row'],'未确定考期',r['class_name'],'已有服务分班完成，常规督学及历史保留',x.REMARK])
    out.active.append(['汇总',3,None,None,'异常表17→14；无新增服务、班级或档案',None])
    out.active.append(['原始资料说明',None,None,None,'3条总表原行此前已清空；原异常记录全21列保存在另页，未补造原表字段',None])
    archive=out.create_sheet('原异常完整记录');archive.append([c.value for c in original['异常与处理总表'][1]])
    for index,r in enumerate(plan['records'],2):
        for col in range(1,22):clone_cell(original['异常与处理总表'].cell(r['exception_row'],col),archive.cell(index,col))
    for sh in out:x.format_sheet(sh)
    es=exc['异常与处理总表'];expected=[r for r in list(es.values)[1:] if (r[7],r[8]) not in TARGETS]
    for row in range(es.max_row,1,-1):
        if (es.cell(row,8).value,es.cell(row,9).value) in TARGETS:es.delete_rows(row)
    x.h.require(es.max_row==15,'Remaining exception count differs');es.tables['ExceptionTable'].ref='A1:U15';es.auto_filter.ref='A1:U15'
    cs=exc['关联记录对比'];removed=0
    for row in range(cs.max_row,1,-1):
        if (cs.cell(row,4).value,cs.cell(row,5).value) in TARGETS:cs.delete_rows(row);removed+=1
    cs.tables['ComparisonTable'].ref=f'A1:K{cs.max_row}';cs.auto_filter.ref=cs.tables['ComparisonTable'].ref
    summary=exc['汇总']
    for row in range(2,summary.max_row+1):
        if summary.cell(row,1).value=='当前保留异常':summary.cell(row,2,14)
        if summary.cell(row,1).value=='当前异常构成':summary.cell(row,2,14);summary.cell(row,3,'2条无目标课程（涉及多档案）、11条规格差异、1条班主任差异')
    summary.append(['2026-10-28三条中医营养配餐分班',3,'已有服务按原表分班完成，已从异常表移出；原始异常信息另行归档'])
    summary.tables['SummaryTable'].ref=f'A1:C{summary.max_row}';summary.auto_filter.ref=summary.tables['SummaryTable'].ref
    for p,book in ((x.EXCEPTIONS,exc),(REPORT,out)):
        temp=p.with_suffix('.tmp.xlsx');book.save(temp)
        with zipfile.ZipFile(temp) as z:x.h.require(z.testzip() is None,'Invalid workbook ZIP')
    saved=openpyxl.load_workbook(x.EXCEPTIONS.with_suffix('.tmp.xlsx'));savedout=openpyxl.load_workbook(REPORT.with_suffix('.tmp.xlsx'))
    x.h.require(list(saved['异常与处理总表'].values)[1:]==expected,'Retained exception cells changed')
    x.h.require(list(saved['聊天处理依据'].values)==list(original['聊天处理依据'].values),'History changed')
    x.h.require(list(savedout['原异常完整记录'].values)[1:]==[r for r in list(original['异常与处理总表'].values)[1:] if (r[7],r[8]) in TARGETS],'Archive differs')
    x.h.require(all(x.sha(p)==plan['hashes'][p.name] for p in INPUTS),'Source changed before replacement')
    REPORT.with_suffix('.tmp.xlsx').replace(REPORT);x.EXCEPTIONS.with_suffix('.tmp.xlsx').replace(x.EXCEPTIONS)
    x.h.save(ROOT/'files.json',{p.name:x.sha(p) for p in (*INPUTS,REPORT)})
    print(x.j(dict(assigned=3,remaining_exceptions=14,comparison_removed=removed,new_services=0)))


def main():
    env=os.getenv('ZSJOS_AGENT_ENV')
    if env is None:env=Path('/etc/zsjos/agent-environment').read_text().strip()
    x.h.require(env=='test','Test only');os.umask(0o077);ROOT.mkdir(mode=0o700,parents=True,exist_ok=True)
    x.MARKER=MARKER;x.TABLES=TABLES
    parser=argparse.ArgumentParser();parser.add_argument('mode',choices=['prepare','fixture','apply','finish']);parser.add_argument('--sha256');a=parser.parse_args()
    if a.mode=='prepare':prepare();return
    plan=x.loads((ROOT/'plan.json').read_text());x.h.require(a.sha256==x.h.digest(plan),'Plan digest differs')
    if a.mode!='finish':x.h.require(plan['script_sha256']==x.sha(Path(__file__)),'Script changed')
    globals()[a.mode](plan)


if __name__=='__main__':main()
