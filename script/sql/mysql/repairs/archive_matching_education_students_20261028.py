#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Archive fully matching existing services; only local review workbooks change.

Personal data stays in user-requested full-row workbooks and private backups.
Database connections are read-only. Existing services are never recreated/reset.
"""
from copy import copy
import json
import os
from pathlib import Path
import shutil
import zipfile

import openpyxl
import import_education_exceptions_20261027 as x

ROOT=Path('/opt/zsjos-runtime/backups/education-matched-archive-20261028')
ARCHIVE=x.ROOT/'output/教务学员_已导入确认记录_20261028.xlsx'
PENDING=x.ROOT/'output/教务学员_多条报名待确认_20261027.xlsx'


def match_records(raw_rows,state):
    found=[]
    for rowno,raw in enumerate(raw_rows,2):
        name,mobile=x.norm(raw[0]),x.norm(raw[1])
        people=[p for p in state['people'] if p['mobile']==mobile]
        if len(people)!=1 or people[0]['name'].strip()!=name:
            continue
        users=[u for u in state['users'] if u['nickname']==raw[12] and u['status']==0]
        if len(users)!=1:
            continue
        class_ids={c['id'] for c in state['classes'] if c['class_name']==raw[11]
                   and c['homeroom_user_id']==users[0]['id'] and c['status']=='SERVING'}
        product_refs={p['product_ref'] for p in state['products'] if p['name']==raw[9]}
        services=[]
        for rel in state['services']:
            if rel['person_id']!=people[0]['id'] or rel['student_name_snapshot']!=name:
                continue
            snap=json.loads(rel['service_snapshot'] or '{}')
            if (snap.get('name',snap.get('productName'))==raw[9] and snap.get('productRef') in product_refs
                and snap.get('skuName')==raw[10] and rel['class_id'] in class_ids
                and rel['owner_user_id']==users[0]['id'] and rel['status']=='active'
                and rel['teaching_closed'] in (0,b'\0') and rel['teaching_delivery'] in (1,b'\1')):
                services.append(rel['id'])
        if services:
            found.append(dict(exception_row=rowno,sheet=raw[7],row=int(raw[8]),service_ids=sorted(services)))
    return found


def clone_cell(source,dest):
    dest.value=source.value;dest.data_type=source.data_type
    dest.font=copy(source.font);dest.fill=copy(source.fill);dest.border=copy(source.border)
    dest.alignment=copy(source.alignment);dest.protection=copy(source.protection)
    dest.number_format=source.number_format
    if source.comment:dest.comment=copy(source.comment)
    if source.hyperlink:dest.hyperlink=copy(source.hyperlink)


def run():
    environment=os.getenv('ZSJOS_AGENT_ENV')
    if environment is None:environment=Path('/etc/zsjos/agent-environment').read_text().strip()
    x.h.require(environment=='test','Unexpected environment')
    os.umask(0o077);ROOT.mkdir(mode=0o700,parents=True,exist_ok=True)
    if (ROOT/'delivery.json').exists():
        delivery=x.h.read(ROOT/'delivery.json')
        x.h.require(all(x.sha(p)==delivery['hashes'][p.name] for p in (x.MASTER,x.EXCEPTIONS,PENDING,ARCHIVE)),
                    'Delivered workbook changed; fresh review required')
        print('Replay: no changes');return
    hashes={p.name:x.sha(p) for p in (x.MASTER,x.EXCEPTIONS,PENDING)}
    for p in (x.MASTER,x.EXCEPTIONS,PENDING):
        backup=ROOT/p.name
        if backup.exists():x.h.require(x.sha(backup)==hashes[p.name],'Existing backup differs')
        else:shutil.copy2(p,backup);backup.chmod(0o600)
    current=openpyxl.load_workbook(x.EXCEPTIONS)
    original=openpyxl.load_workbook(x.EXCEPTIONS)
    raw_rows=list(current['异常与处理总表'].values)[1:]
    x.h.require(len(raw_rows)==154,'Exception scope changed')
    state=x.h.snapshot();selected=match_records(raw_rows,state)
    x.h.require(len(selected)==87,'Exact-match scope changed; do not guess')
    keys={(r['sheet'],r['row']) for r in selected}
    x.h.require(len(keys)==87,'Duplicate source references')
    master=openpyxl.load_workbook(x.MASTER)
    pending=openpyxl.load_workbook(PENDING,data_only=True)['待确认报名']
    pending_keys={(r[1],r[2]) for r in list(pending.values)[1:]}
    x.h.require(not keys & pending_keys,'Pending decision must remain')
    archive=openpyxl.Workbook();archive.active.title='汇总'
    archive.active.append(['项目','数量','说明'])
    for row in [('确认为已导入',87,'唯一手机号档案；姓名、课程规格、班级和班主任全部一致'),
                ('异常表剩余',67,'包含10条已单独列出的待确认报名及57条其他异常'),
                ('本次新增系统服务',0,'仅整理表格，系统数据未修改'),
                ('字段说明',44,'已导入完整信息保留总表全部44个原始字段；原异常记录另页保留')]:archive.active.append(row)
    full=archive.create_sheet('已导入完整信息')
    first=master[selected[0]['sheet']+'-学员导入填写']
    headers=[c.value for c in first[1]]
    x.h.require(len(headers)==44,'Master fields changed')
    full.append(['总表工作表','总表行号','归档前异常表行号','处理结果']+headers)
    history=archive.create_sheet('原异常完整记录')
    history.append([c.value for c in original['异常与处理总表'][1]])
    for out_row,r in enumerate(selected,2):
        source=master[r['sheet']+'-学员导入填写']
        raw=raw_rows[r['exception_row']-2]
        x.h.require([c.value for c in source[1]]==headers,'Master headings differ')
        # Excel stores master phone numbers numerically and exception phones as text;
        # compare normalized identity values, then preserve each source's exact cells.
        x.h.require([x.norm(source.cell(r['row'],col).value) for col in (1,2,4,5,6,7)]==[x.norm(raw[col]) for col in (0,1,9,10,11,12)],
                    'Master/exception source mismatch')
        for col,value in enumerate([r['sheet'],r['row'],r['exception_row'],'已导入（系统已有完全一致服务）'],1):full.cell(out_row,col,value)
        for col in range(1,45):clone_cell(source.cell(r['row'],col),full.cell(out_row,col+4))
        for col in range(1,22):clone_cell(original['异常与处理总表'].cell(r['exception_row'],col),history.cell(out_row,col))
    for ws in archive:x.format_sheet(ws)
    full.freeze_panes='G2'
    sheet=current['异常与处理总表']
    for r in sorted(selected,key=lambda r:r['exception_row'],reverse=True):sheet.delete_rows(r['exception_row'])
    sheet.tables['ExceptionTable'].ref=f'A1:U{sheet.max_row}';sheet.auto_filter.ref=sheet.tables['ExceptionTable'].ref
    comparison=current['关联记录对比'];comparison_removed=0
    for row in range(comparison.max_row,1,-1):
        if (comparison.cell(row,4).value,comparison.cell(row,5).value) in keys:
            comparison.delete_rows(row);comparison_removed+=1
    comparison.tables['ComparisonTable'].ref=f'A1:K{comparison.max_row}';comparison.auto_filter.ref=comparison.tables['ComparisonTable'].ref
    summary=current['汇总']
    for row in range(2,summary.max_row+1):
        if summary.cell(row,1).value=='当前保留异常':summary.cell(row,2,67)
    summary.append(['2026-10-28完全一致已导入归档',87,'已从异常主表移出；系统服务未修改'])
    summary.append(['已导入归档文件',None,ARCHIVE.name])
    summary.append(['当前异常构成',67,'10条待确认报名＋57条其他异常；独立待确认表另有1条不在此表'])
    summary.tables['SummaryTable'].ref=f'A1:C{summary.max_row}';summary.auto_filter.ref=summary.tables['SummaryTable'].ref
    expected=[r for r in raw_rows if (r[7],int(r[8])) not in keys]
    x.h.require(len(expected)==67,'Exception partition differs')
    for path,book in ((x.EXCEPTIONS,current),(ARCHIVE,archive)):
        temp=path.with_suffix('.tmp.xlsx');book.save(temp)
        with zipfile.ZipFile(temp) as z:x.h.require(z.testzip() is None,'Invalid ZIP')
    saved=openpyxl.load_workbook(x.EXCEPTIONS.with_suffix('.tmp.xlsx'))
    saved_archive=openpyxl.load_workbook(ARCHIVE.with_suffix('.tmp.xlsx'))
    x.h.require(list(saved['异常与处理总表'].values)[1:]==expected,'Retained exception data changed')
    x.h.require(list(saved['聊天处理依据'].values)==list(original['聊天处理依据'].values),'Historical decisions changed')
    for out_row,r in enumerate(selected,2):
        source=master[r['sheet']+'-学员导入填写']
        for col in range(1,45):
            a,b=source.cell(r['row'],col),saved_archive['已导入完整信息'].cell(out_row,col+4)
            x.h.require(a.value==b.value and a.data_type==b.data_type and a.number_format==b.number_format,'Archive source field differs')
        x.h.require(tuple(c.value for c in saved_archive['原异常完整记录'][out_row])==raw_rows[r['exception_row']-2],
                    'Archive exception row differs')
    x.h.require(match_records(raw_rows,x.h.snapshot())==selected,'Live matching changed before file replacement')
    x.h.require(all(x.sha(p)==hashes[p.name] for p in (x.MASTER,x.EXCEPTIONS,PENDING)),'Input changed concurrently')
    x.h.save(ROOT/'matches.json',selected)
    x.h.save(ROOT/'before.json',state)
    ARCHIVE.with_suffix('.tmp.xlsx').replace(ARCHIVE)
    x.EXCEPTIONS.with_suffix('.tmp.xlsx').replace(x.EXCEPTIONS)
    x.h.save(ROOT/'delivery.json',dict(archived=87,remaining=67,comparison_removed=comparison_removed,database_writes=0,
        hashes={p.name:x.sha(p) for p in (x.MASTER,x.EXCEPTIONS,PENDING,ARCHIVE)}))
    print(x.j(dict(archived=87,remaining=67,comparison_removed=comparison_removed,database_writes=0,full_source_columns=44)))


if __name__=='__main__':run()
