#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Add two confirmed distinct courses to claimed identities; preserve old services."""
import argparse
import json
import os
from pathlib import Path
import shutil
import zipfile

import openpyxl
import import_education_exceptions_20261027 as x
from import_two_confirmed_students_20261028 import session
from archive_matching_education_students_20261028 import clone_cell, PENDING
from export_education_assignment_comparison_20261028 import OUTPUT as REVIEW

ROOT=Path('/opt/zsjos-runtime/backups/education-two-claimed-courses-20261028')
MARKER='edu-two-claimed-courses-20261028'
REPORT=x.ROOT/'output/教务学员_两条新增课程处理结果_20261028.xlsx'
TARGETS={('中医营养配餐',3),('中医营养配餐',20)}
INPUTS=(x.MASTER,x.EXCEPTIONS,PENDING,REVIEW)


def prepare():
    x.h.require(not (ROOT/'plan.json').exists(),'Frozen plan exists')
    _,allrows=x.source_rows(x.EXCEPTIONS);selected=[r for r in allrows if (r['sheet'],r['row']) in TARGETS]
    x.h.require(len(allrows)==14 and len(selected)==2,'Source scope changed')
    hashes={p.name:x.sha(p) for p in INPUTS};st=x.h.snapshot();s=session();options=x.h.api(s,'GET','/zsjos/student-entry/options')
    master=openpyxl.load_workbook(x.MASTER);existing=[]
    for r in selected:
        src=master[r['sheet']+'-学员导入填写']
        x.h.require([x.norm(src.cell(r['row'],col).value) for col in (1,2,4,5,6,7)]==[r['name'],r['mobile'],r['product'],r['sku'],r['class_name'],r['planner']],'Master/source mismatch')
        claims=[c for c in st['claims'] if c['contact_value']==r['mobile']]
        x.h.require(len(claims)==1 and claims[0]['person_id'] and claims[0]['reservation_key'] is None,'Claim ambiguous')
        ps=[p for p in st['people'] if p['id']==claims[0]['person_id']]
        x.h.require(len(ps)==1 and ps[0]['name']==r['name'] and ps[0]['mobile']==r['mobile'] and ps[0]['identity_status']=='student','Confirmed person differs')
        r['person_id']=r['confirmed_claim_person_id']=ps[0]['id']
        r['confirmed_contact_candidate_ids']=sorted(p['id'] for p in st['people'] if r['mobile'] in (p['mobile'],p['wechat_id']))
        products=[p for p in st['products'] if p['name']==r['product'] and p['status']==0];x.h.require(len(products)==1,'Product unavailable')
        r['product_id'],r['product_ref']=products[0]['id'],products[0]['product_ref']
        skus=[k for k in st['skus'] if k['spu_id']==r['product_id'] and k['sku_name']==r['sku'] and k['status']==0]
        x.h.require(len(skus)==1 and skus[0]['sku_ref'] in {k['skuRef'] for k in options['courses']},'SKU unavailable')
        r['sku_id'],r['skuRef']=skus[0]['id'],skus[0]['sku_ref']
        owners=[u for u in st['users'] if u['nickname']==r['planner'] and u['status']==0];x.h.require(len(owners)==1,'Planner unavailable');r['owner_id']=owners[0]['id']
        classes=[c for c in st['classes'] if c['class_name']==r['class_name'] and c['homeroom_user_id']==r['owner_id'] and c['status']=='SERVING' and c['class_kind']=='NORMAL']
        x.h.require(len(classes)==1 and classes[0]['id'] in {c['id'] for c in options['classes']},'Target class unavailable');r['class_id']=classes[0]['id']
        old=[v for v in st['services'] if v['person_id']==r['person_id']];x.check_existing_courses(r,old)
        x.h.require(all(json.loads(v['service_snapshot']).get('name')=='大健康轻创网红博主' for v in old),'Existing service scope changed')
        existing.extend(old);r['same_services']=[];r['other_services']=[v['id'] for v in old];r['issues']=[]
    x.h.require(len(existing)==3 and sum(v['status']=='cancelled' for v in existing)==1,'Prior media service scope changed')
    x.h.require(all(x.sha(p)==hashes[p.name] for p in INPUTS),'Input changed')
    for p in (x.MASTER,x.EXCEPTIONS):
        dst=ROOT/p.name;x.h.require(not dst.exists(),'Backup exists');shutil.copy2(p,dst);dst.chmod(0o600)
    x.prepare(selected,st,sync=False,missing_sku_count=0,extra=dict(mobile_only=False,fixture_failure_points=[1,3,7,12,16],
        wrapper_sha256=x.sha(Path(__file__)),source_count=14,protected_old_services=existing,input_hashes=hashes))


def finish(plan):
    if (ROOT/'files.json').exists():
        hashes=x.h.read(ROOT/'files.json');x.h.require(all(x.sha(p)==hashes[p.name] for p in (*INPUTS,REPORT)),'Delivered files changed')
        print('File replay: no changes');return
    c=x.h.connection()
    try:
        with c.cursor() as q:
            q.execute('START TRANSACTION WITH CONSISTENT SNAPSHOT, READ ONLY');after,result=x.verify(q,plan)
            for old in plan['protected_old_services']:
                actual=x.rows(q,f'SELECT * FROM {x.S} WHERE tenant_id=1 AND id=%s',(old['id'],))
                x.h.require(len(actual)==1 and x.h.encoded(actual[0])==x.h.encoded(old),'Prior service changed')
    finally:c.rollback();c.close()
    x.h.save(ROOT/'independent-after.json',after)
    s=session()
    for row in result:
        rid=row['service_id'];ctx=x.h.api(s,'GET',f'/zsjos/student/service/{rid}/contact-context');stages={v['code']:v for v in ctx['deliveryStages']}
        x.h.require(ctx['acceptanceStatus']=='accepted' and ctx['deliveryStage']=='supervision' and ctx['currentTask'] is None
            and stages['first_contact']['status']==stages['study_plan']['status']=='done' and all(v['completed'] for v in ctx['firstContactProgress']['items']),'API completion differs')
        hist=x.h.api(s,'GET',f'/zsjos/student/service/{rid}/contact-records?pageNo=1&pageSize=100')
        x.h.require(hist['total']==2 and {v['contactType'] for v in hist['list']}=={v[0] for v in x.TYPES}
            and all(v['successful'] and v['remark']==x.REMARK for v in hist['list']),'API records differ')
    x.h.save(ROOT/'api.json',dict(contexts=2,histories=2,passed=True))
    x.h.require(all(x.sha(p)==plan['input_hashes'][p.name] for p in INPUTS),'Input changed')
    exc=openpyxl.load_workbook(x.EXCEPTIONS);original=openpyxl.load_workbook(x.EXCEPTIONS);master=openpyxl.load_workbook(x.MASTER)
    out=openpyxl.Workbook();out.active.title='汇总'
    for row in [('处理项','数量','说明'),('复用档案新增中医营养配餐服务',2,'完成首联、学习计划，常规督学；指定10.27备注'),
                ('已有网红博主服务保持',3,'包括2条有效和1条取消服务，原字段及历史不变'),('异常表剩余',12,'11条规格差异、1条班主任差异'),('新增档案或班级',0,'复用当前档案和班级')]:out.active.append(row)
    full=out.create_sheet('原表完整信息');headers=[c.value for c in master['中医营养配餐-学员导入填写'][1]]
    x.h.require(len(headers)==44,'Master field count differs');full.append(['总表工作表','总表行号','处理结果','备注']+headers)
    archive=out.create_sheet('原异常完整记录');archive.append([c.value for c in original['异常与处理总表'][1]])
    for i,r in enumerate(plan['ready'],2):
        for col,v in enumerate([r['sheet'],r['row'],'已新增课程，原服务保留',x.REMARK],1):full.cell(i,col,v)
        for col in range(1,45):clone_cell(master[r['sheet']+'-学员导入填写'].cell(r['row'],col),full.cell(i,col+4))
        for col in range(1,22):clone_cell(original['异常与处理总表'].cell(r['exception_row'],col),archive.cell(i,col))
    for sh in out:x.format_sheet(sh)
    es=exc['异常与处理总表'];expected=[r for r in list(es.values)[1:] if (r[7],r[8]) not in TARGETS]
    for row in range(es.max_row,1,-1):
        if (es.cell(row,8).value,es.cell(row,9).value) in TARGETS:es.delete_rows(row)
    x.h.require(es.max_row==13,'Exception partition differs');es.tables['ExceptionTable'].ref='A1:U13';es.auto_filter.ref='A1:U13'
    cs=exc['关联记录对比'];removed=0
    for row in range(cs.max_row,1,-1):
        if (cs.cell(row,4).value,cs.cell(row,5).value) in TARGETS:cs.delete_rows(row);removed+=1
    cs.tables['ComparisonTable'].ref=f'A1:K{cs.max_row}';cs.auto_filter.ref=cs.tables['ComparisonTable'].ref
    summary=exc['汇总']
    for row in range(2,summary.max_row+1):
        if summary.cell(row,1).value=='当前保留异常':summary.cell(row,2,12)
        if summary.cell(row,1).value=='当前异常构成':summary.cell(row,2,12);summary.cell(row,3,'11条规格差异、1条班主任差异')
    summary.append(['2026-10-28两条独立课程新增',2,'复用号码占用档案；原网红博主服务完全保留；已移出异常表'])
    summary.tables['SummaryTable'].ref=f'A1:C{summary.max_row}';summary.auto_filter.ref=summary.tables['SummaryTable'].ref
    for p,book in ((x.EXCEPTIONS,exc),(REPORT,out)):
        temp=p.with_suffix('.tmp.xlsx');book.save(temp)
        with zipfile.ZipFile(temp) as z:x.h.require(z.testzip() is None,'Invalid workbook ZIP')
    saved=openpyxl.load_workbook(x.EXCEPTIONS.with_suffix('.tmp.xlsx'));savedout=openpyxl.load_workbook(REPORT.with_suffix('.tmp.xlsx'))
    x.h.require(list(saved['异常与处理总表'].values)[1:]==expected,'Retained rows differ')
    x.h.require(list(saved['聊天处理依据'].values)==list(original['聊天处理依据'].values),'History changed')
    x.h.require(list(savedout['原异常完整记录'].values)[1:]==[r for r in list(original['异常与处理总表'].values)[1:] if (r[7],r[8]) in TARGETS],'Archived exception differs')
    for row in list(savedout['原表完整信息'].rows)[1:]:
        for col in range(1,45):
            a=master[row[0].value+'-学员导入填写'].cell(row[1].value,col);b=row[col+3]
            x.h.require((a.value,a.data_type,a.number_format)==(b.value,b.data_type,b.number_format),'Original44fields differ')
    x.h.require(all(x.sha(p)==plan['input_hashes'][p.name] for p in INPUTS),'Concurrent input change')
    REPORT.with_suffix('.tmp.xlsx').replace(REPORT);x.EXCEPTIONS.with_suffix('.tmp.xlsx').replace(x.EXCEPTIONS)
    x.h.save(ROOT/'files.json',{p.name:x.sha(p) for p in (*INPUTS,REPORT)})
    print(x.j(dict(imported=2,old_services_unchanged=3,remaining_exceptions=12,comparison_removed=removed)))


def main():
    env=os.getenv('ZSJOS_AGENT_ENV')
    if env is None:env=Path('/etc/zsjos/agent-environment').read_text().strip()
    x.h.require(env=='test','Test only');os.umask(0o077);ROOT.mkdir(mode=0o700,parents=True,exist_ok=True);x.EVIDENCE=ROOT;x.MARKER=MARKER
    parser=argparse.ArgumentParser();parser.add_argument('mode',choices=['prepare','fixture','apply','finish']);parser.add_argument('--sha256');a=parser.parse_args()
    if a.mode=='prepare':prepare();return
    plan=x.loads((ROOT/'plan.json').read_text());x.h.require(a.sha256==x.h.digest(plan),'Plan digest differs')
    if a.mode!='finish':x.h.require(plan['wrapper_sha256']==x.sha(Path(__file__)),'Mutation wrapper changed')
    if a.mode=='finish':finish(plan)
    else:getattr(x,a.mode)(plan)


if __name__=='__main__':main()
