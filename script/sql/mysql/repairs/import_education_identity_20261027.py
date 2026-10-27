#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Follow-up under explicit mobile/name decisions; retain ambiguous registrations.

Reuses the audited historical import transaction. No contact reassignment, identity
merging, existing teaching-stage reset, schema changes or notification dispatch.
"""
import argparse
from collections import Counter, defaultdict
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import zipfile

import openpyxl
import import_education_exceptions_20261027 as x

MARKER='edu-identity-20261027'
ROOT=Path('/opt/zsjos-runtime/backups/education-identity-20261027')
REPORT=x.ROOT/'output/教务学员_身份确认处理结果_20261027.xlsx'
PENDING=x.ROOT/'output/教务学员_多条报名待确认_20261027.xlsx'
GROUPS=[ [('学会公共营养师',35),('学会公共营养师',36)],
         [('学会公共营养师',37),('学会公共营养师',53),('人社公共营养师',40)],
         [('人社健康管理师',27),('人社健康管理师',128)],
         [('心理咨询师',19),('心理咨询师',48)],
         [('执业中西药师',39),('执业中西药师',95)],
         [('执业中西药师',40),('执业中西药师',112)] ]
CONFIRMED_PLANNER=('学会公共营养师',120)


def clean_name(value):
    return re.sub(r'[（(]\s*代理(?:学员)?\s*[）)]','',str(value)).strip()


def configure():
    env=os.getenv('ZSJOS_AGENT_ENV')
    if env is None:env=Path('/etc/zsjos/agent-environment').read_text().strip()
    x.h.require(env=='test','Test-only authorized execution')
    os.umask(0o077);ROOT.mkdir(mode=0o700,parents=True,exist_ok=True)
    x.EVIDENCE=ROOT;x.MARKER=MARKER


def audit():
    _,records=x.source_rows(x.EXCEPTIONS)
    x.h.require(len(records)==241,'Initial exception scope changed')
    state=x.h.snapshot();session=x.h.auth()
    options=x.h.api(session,'GET','/zsjos/student-entry/options')
    eligible={c['id'] for c in options['classes']}
    master=openpyxl.load_workbook(x.MASTER,data_only=True)
    names_by_mobile=defaultdict(set)
    for sheet in master.worksheets[:8]:
        for raw in list(sheet.values)[1:]:
            if raw[0] and raw[1]:names_by_mobile[x.norm(raw[1])].add(clean_name(raw[0]))
    holds={key for group in GROUPS for key in group}
    name_changes={};service_changes={};checks={}
    for r in records:
        r['issues']=[];r['original_name']=r['name'];r['name']=clean_name(r['name'])
        key=r['sheet'],r['row']
        if key==CONFIRMED_PLANNER:r['planner']='周老师【梁颖】'
        raw=list(master[r['sheet']+'-学员导入填写'].values)[r['row']-1]
        r['removed_from_master']=not any(raw)
        if any(raw):
            actual=[clean_name(raw[0]),x.norm(raw[1]),x.norm(raw[3]),x.norm(raw[4]),x.norm(raw[5]),
                    r['planner'] if key==CONFIRMED_PLANNER else x.norm(raw[6])]
            x.h.require(actual==[r['name'],r['mobile'],r['product'],r['sku'],r['class_name'],r['planner']],
                        'Unexpected master/source difference')
        else:r['issues'].append('已从总表移除，已有服务待分班')
        products=[p for p in state['products'] if p['name']==r['product'] and p['status']==0]
        x.h.require(len(products)==1,'Product ambiguous')
        r['product_id'],r['product_ref']=products[0]['id'],products[0]['product_ref']
        skus=[k for k in state['skus'] if k['spu_id']==r['product_id'] and k['sku_name']==r['sku']]
        x.h.require(len(skus)==1,'SKU ambiguous')
        sku=skus[0];r['skuRef'],r['sku_id']=sku['sku_ref'],sku['id']
        r['historical_sku']=sku['status']==1
        if sku['status']==1:
            x.h.require(sku['price']==0 and sku['creator']=='edu-exceptions-20261027','Disabled SKU not previously authorized')
        people=[p for p in state['people'] if p['mobile']==r['mobile']]
        r['mobile_match_count']=len(people)
        claims={a['person_id'] for a in state['claims'] if a['contact_value']==r['mobile'] and a['person_id']}
        wechat=[p for p in state['people'] if p['wechat_id']==r['mobile']]
        r['wechat_match_count']=len(wechat)
        r['candidate_person_ids']=[p['id'] for p in people]
        if len(people)>1:r['issues'].append('系统手机号字段重复，无法唯一选择档案')
        elif len(people)==1:
            r['person_id']=people[0]['id']
            if claims-{people[0]['id']}:r['issues'].append('手机号唯一占用记录指向其他档案')
        elif wechat or claims:r['issues'].append('号码仅匹配微信或占用记录，未匹配系统手机号')
        if len(names_by_mobile[r['mobile']])>1:r['issues'].append('总表同手机号有不同姓名，需确认')
        if not re.fullmatch(r'1[3-9]\d{9}',r['mobile']):r['issues'].append('手机号格式异常')
        users=[u for u in state['users'] if u['status']==0 and u['nickname']==r['planner']]
        x.h.require(len(users)==1,'Planner unresolved after explicit confirmation')
        r['owner_id']=users[0]['id']
        classes=[cl for cl in state['classes'] if cl['class_name']==r['class_name']
                 and cl['homeroom_user_id']==r['owner_id'] and cl['status']=='SERVING']
        if len(classes)>1:r['issues'].append('班级不唯一')
        elif classes:
            r['class_id']=classes[0]['id']
            x.h.require(r['class_id'] in eligible,'Class unavailable')
        else:r['missing_class']=True
        r['same_services']=[];r['other_services']=[]
        for rel in state['services']:
            if rel['person_id'] not in r['candidate_person_ids']:continue
            snap=json.loads(rel['service_snapshot'] or '{}')
            item=dict(id=rel['id'],person_id=rel['person_id'],product=snap.get('name',snap.get('productName')),
                      sku=snap.get('skuName'),class_id=rel['class_id'],owner_id=rel['owner_user_id'],
                      status=rel['status'],name=rel['student_name_snapshot'])
            same=snap.get('productRef')==r['product_ref'] or item['product']==r['product']
            r['same_services' if same else 'other_services'].append(item)
        if r['same_services']:r['issues'].append('已有同课程服务，不重复导入')
        if key in holds:r['issues'].append('多条报名保留待确认')
        exact=[e for e in r['same_services'] if e['sku']==r['sku'] and e['class_id']==r.get('class_id') and e['owner_id']==r['owner_id']]
        r['duplicate_category']='无同课程服务'
        if len(people)>1:r['duplicate_category']='手机号匹配多个档案，不能判为同一学员'
        elif exact:
            r['duplicate_category']='姓名、手机号、课程规格、班级、班主任完全相同' if people[0]['name'].strip()==r['name'] and any(e['name']==r['name'] for e in exact) else '课程规格班级班主任相同，姓名需按表修正'
        elif r['same_services']:
            r['duplicate_category']='规格相同，班级或班主任不同' if any(e['sku']==r['sku'] for e in r['same_services']) else '同课程规格不同'
        if len(people)==1 and not r['removed_from_master'] and len(names_by_mobile[r['mobile']])<=1 and not claims-{people[0]['id']}:
            person=people[0]
            if person['name']!=r['name']:
                change=dict(id=person['id'],name=r['name'],sheet=r['sheet'],row=r['row'])
                x.h.require(person['id'] not in name_changes or name_changes[person['id']]['name']==r['name'],'Conflicting name corrections')
                name_changes[person['id']]=change
            for e in r['same_services']:
                if e['sku']==r['sku'] and e['name']!=r['name']:
                    service_changes[e['id']]=dict(id=e['id'],person_id=person['id'],name=r['name'],sheet=r['sheet'],row=r['row'])
        if not r['issues']:
            if r['mobile'] not in checks:checks[r['mobile']]=x.h.api(session,'POST','/zsjos/student-entry/check',dict(mobile=r['mobile'],wechatId=None))
            check=checks[r['mobile']]
            # The user's mobile-only decision resolves cross-field ambiguity, but never
            # overrides a different contact claim or multiple mobile matches.
            if check['status']=='conflict':
                x.h.require(len(people)==1 and wechat and not claims-{people[0]['id']},'Unresolved API identity conflict')
                r['mobile_priority_override']=True
            else:x.h.require(check['status'] in ('new','reuse') and check.get('personId')==r.get('person_id'),'Identity API changed')
    return records,state,list(name_changes.values()),list(service_changes.values())


def sync(records):
    master=openpyxl.load_workbook(x.MASTER);exceptions=openpyxl.load_workbook(x.EXCEPTIONS)
    edits=[]
    for r in records:
        erow=r['exception_row'];title=r['sheet']+'-学员导入填写'
        if r['name']!=r['original_name']:
            x.h.require(not r['removed_from_master'],'Cannot restore excluded row')
            master[title].cell(r['row'],1,r['name']);exceptions['异常与处理总表'].cell(erow,1,r['name'])
            edits.append(dict(sheet=title,row=r['row'],column=1,exception_row=erow,exception_column=1))
        if (r['sheet'],r['row'])==CONFIRMED_PLANNER:
            master[title].cell(r['row'],7,r['planner']);exceptions['异常与处理总表'].cell(erow,13,r['planner'])
            edits.append(dict(sheet=title,row=r['row'],column=7,exception_row=erow,exception_column=13))
    for path,w in ((x.MASTER,master),(x.EXCEPTIONS,exceptions)):
        backup=openpyxl.load_workbook(ROOT/path.name)
        allowed={(e['sheet'],e['row'],e['column']) if path==x.MASTER else ('异常与处理总表',e['exception_row'],e['exception_column']) for e in edits}
        for old in backup:
            for row in old:
                for c in row:
                    if (old.title,c.row,c.column) not in allowed:
                        x.h.require(c.value==w[old.title].cell(c.row,c.column).value,'Unrelated source edit')
        temp=path.with_suffix('.tmp.xlsx');w.save(temp);openpyxl.load_workbook(temp).close();temp.replace(path)
    x.h.save(ROOT/'source-edits.json',edits)


def prepare():
    x.h.require(not (ROOT/'plan.json').exists(),'Frozen plan exists')
    for path in (x.MASTER,x.EXCEPTIONS):
        dest=ROOT/path.name
        x.h.require(not dest.exists(),'Previous backup exists; review before resuming')
        shutil.copy2(path,dest);dest.chmod(0o600)
    records,state,names,services=audit()
    sync(records)
    x.prepare(records,state,sync=False,missing_sku_count=0,extra=dict(mobile_only=True,name_updates=names,
        service_name_updates=services,fixture_failure_points=[1,20,80,250,500],
        wrapper_sha256=x.sha(Path(__file__)),source_count=241))
    print(x.j(dict(ready=sum(not r['issues'] for r in records),name_corrections=len(names),service_display_name_corrections=len(services),
                   initial_comparison=Counter(r['duplicate_category'] for r in records))))


def verify_api(plan,result):
    session=x.h.auth();selected={}
    for r in result:selected.setdefault(r['class_id'],r)
    for r in selected.values():
        context=x.h.api(session,'GET',f'/zsjos/student/service/{r["service_id"]}/contact-context')
        stages={v['code']:v for v in context['deliveryStages']}
        x.h.require(context['deliveryStage']=='supervision' and context['currentTask'] is None
                    and stages['first_contact']['status']==stages['study_plan']['status']=='done','API stage differs')
        history=x.h.api(session,'GET',f'/zsjos/student/service/{r["service_id"]}/contact-records?pageNo=1&pageSize=100')
        x.h.require(history['total']==2 and all(v['successful'] and v['remark']==x.REMARK for v in history['list']),'API completion differs')
    catalog=x.h.api(session,'GET','/zsjos/student-entry/options')
    x.h.require(not {r['skuRef'] for r in plan['ready'] if r['historical_sku']} & {k['skuRef'] for k in catalog['courses']},'Disabled SKU exposed')
    x.h.save(ROOT/'api.json',dict(samples=len(selected),contexts=True,histories=True,disabled_excluded=True))


def finish(plan):
    c=x.h.connection()
    try:
        with c.cursor() as q:
            q.execute('START TRANSACTION WITH CONSISTENT SNAPSHOT, READ ONLY')
            after,result=x.verify(q,plan)
    finally:c.rollback();c.close()
    x.h.save(ROOT/'independent-after.json',after)
    verify_api(plan,result)
    done={r['exception_row'] for r in result};remain=[r for r in plan['records'] if r['exception_row'] not in done]
    x.h.require(len(done)+len(remain)==241,'Partition differs')
    x.h.require(x.sha(x.MASTER)==plan['master_sha256'],'Master changed')
    if (ROOT/'files.json').exists():
        old=x.h.read(ROOT/'files.json')
        x.h.require(x.sha(x.EXCEPTIONS)==old['exceptions'] and x.sha(REPORT)==old['report'] and x.sha(PENDING)==old['pending'],'Delivered files changed')
        print('Files verified; replay zero changes');return
    x.h.require(x.sha(x.EXCEPTIONS)==plan['exception_sha256'],'Exception file changed')
    w=openpyxl.load_workbook(x.EXCEPTIONS);baseline=openpyxl.load_workbook(x.EXCEPTIONS)
    sheet=w['异常与处理总表'];lookup={(r['sheet'],r['row']):r for r in plan['records']}
    for r in remain:
        row=r['exception_row'];issues=[v for v in r['issues'] if v!='已有同课程服务，不重复导入']
        sheet.cell(row,3,'已有课程，保留核对' if r['same_services'] else '待确认')
        sheet.cell(row,4,'；'.join(r['issues']))
        sheet.cell(row,6,'；'.join(issues) if issues else '已有服务不新增；核对班级和班主任是否调整')
        old=sheet.cell(row,7).value or ''
        sheet.cell(row,7,str(old)+'\n2026-10-27：按手机号字段和最新姓名规则复核；本轮未新增，详见身份确认处理结果。')
    for row in sorted(done,reverse=True):sheet.delete_rows(row)
    sheet.tables['ExceptionTable'].ref=f'A1:U{sheet.max_row}';sheet.auto_filter.ref=sheet.tables['ExceptionTable'].ref
    compare=w['关联记录对比'];refs={(r['sheet'],r['row']) for r in result}
    for row in range(compare.max_row,1,-1):
        if (compare.cell(row,4).value,compare.cell(row,5).value) in refs:compare.delete_rows(row)
    compare.tables['ComparisonTable'].ref=f'A1:K{compare.max_row}';compare.auto_filter.ref=compare.tables['ComparisonTable'].ref
    summary=w['汇总'];summary.delete_rows(1,summary.max_row);summary.append(['项目','数量','说明'])
    summary_rows=[('上轮剩余异常',241,'本轮处理起点'),('本轮导入成功',len(result),'常规督学；完成首联和学习计划'),
        ('当前保留异常',len(remain),'已有服务及仍需确认记录保留'),('人员姓名修正',len(plan['name_updates']),'以表格姓名为准，去除代理标记'),
        ('已有服务显示姓名修正',len(plan['service_name_updates']),'仅匹配手机号和相同课程规格的服务，其他状态未改'),
        ('多条报名待确认',13,'6组；12条在当前异常表，另1条是总表执业中西药师112行'),
        ('已确认规划师修正',1,'学会公共营养师120行按确认账号归属'),('本轮结果文件',None,REPORT.name),('独立待确认文件',None,PENDING.name)]
    for row in summary_rows:summary.append(row)
    summary.tables['SummaryTable'].ref=f'A1:C{summary.max_row}';x.format_sheet(summary)
    out=openpyxl.Workbook();out.active.title='汇总';out.active.append(['项目','数量','说明'])
    for row in summary_rows:out.active.append(row)
    imported=out.create_sheet('本轮导入');imported.append(['总表工作表','总表行号','课程','规格','班级','阶段','备注'])
    for r in plan['ready']:imported.append([r['sheet'],r['row'],r['product'],r['sku'],r['class_name'],'常规督学',x.REMARK])
    held=out.create_sheet('保留异常');held.append(['总表工作表','总表行号','课程','规格','手机号档案数量','复核分类','保留原因'])
    for r in remain:held.append([r['sheet'],r['row'],r['product'],r['sku'],r['mobile_match_count'],r['duplicate_category'],'；'.join(r['issues'])])
    detail=out.create_sheet('已有课程逐条对照');detail.append(['总表工作表','总表行号','表内规格','系统规格','表内班级','系统班级','班主任一致','学生姓名按表一致','复核结论'])
    dbservices={s['id']:s for s in after[x.S]};dbpeople={p['id']:p for p in after[x.P]};dbclasses={c['id']:c for c in after[x.C]}
    exact_count=0
    for r in plan['records']:
        found_exact=False
        for e in r['same_services']:
            rel=dbservices.get(e['id']) or next(v for v in x.h.query('SELECT * FROM zsjos_service_relation WHERE tenant_id=1 AND id=%s',(e['id'],)))
            person=dbpeople.get(rel['person_id']) or x.h.query('SELECT * FROM zsjos_person WHERE tenant_id=1 AND id=%s',(rel['person_id'],))[0]
            name_ok=rel['student_name_snapshot']==r['name'] and person['name'].strip()==r['name']
            eq=r['mobile_match_count']==1 and e['sku']==r['sku'] and e['class_id']==r.get('class_id') and e['owner_id']==r['owner_id'] and name_ok
            found_exact|=eq
            detail.append([r['sheet'],r['row'],r['sku'],e['sku'],r['class_name'],dbclasses.get(e['class_id'],{}).get('class_name','未分班'),
                '是' if e['owner_id']==r['owner_id'] else '否','是' if name_ok else '否','完全重复' if eq else r['duplicate_category']])
        exact_count+=int(found_exact)
    pending=openpyxl.Workbook();pending.active.title='待确认报名'
    ps=pending.active
    ps.append(['组号','总表工作表','总表行号','课程','规格','班级','班主任依据（总表）','差异说明','您的处理意见'])
    master=openpyxl.load_workbook(x.MASTER,data_only=True)
    notes=['同规格同班级，班主任不同；选保留行或明确两条独立服务',
           '学会两行班主任不同，另有人社规格；确认保留行及学会、人社是否均保留',
           '同规格同班级，班主任不同；选保留行或明确两条独立服务',
           '学会三级VIP4980与心协/学会无级别，考期班及班主任不同',
           '同规格同班级，班主任不同；选保留行或明确两条独立服务',
           'VIP全程全科与基础精讲全科，同班级同班主任；选保留行或明确两条独立服务']
    aliases={'小江老师【钟茂源】':'小江老师','小敏老师【汪雅琪1】':'小敏老师','小束老师【汪雅琪2】':'小束老师','小惠老师【张雨娴】':'小惠老师'}
    for i,group in enumerate(GROUPS,1):
        for short,rowno in group:
            raw=list(master[short+'-学员导入填写'].values)[rowno-1]
            ps.append([i,short,rowno,raw[3],raw[4],raw[5],aliases.get(raw[6],f'{short}第{rowno}行规划师'),notes[i-1],None])
    for book in (out,pending):
        for ws in book:x.format_sheet(ws)
    outputs=[(x.EXCEPTIONS,w),(REPORT,out),(PENDING,pending)]
    for path,book in outputs:
        temp=path.with_suffix('.tmp.xlsx');book.save(temp)
        with zipfile.ZipFile(temp) as z:x.h.require(z.testzip() is None,'Invalid workbook ZIP')
    saved=openpyxl.load_workbook(x.EXCEPTIONS.with_suffix('.tmp.xlsx'),data_only=True)['异常与处理总表']
    x.h.require(saved.max_row==len(remain)+1,'Exception count differs')
    protected=[1,2,5,8,9,10,11,12,13,14,15,16,17,18,19,20,21]
    for i,r in enumerate(remain,2):
        for col in protected:x.h.require(saved.cell(i,col).value==baseline['异常与处理总表'].cell(r['exception_row'],col).value,'Protected cell changed')
    x.h.require(ps.max_row==14,'Pending group count differs')
    for path,_ in outputs:path.with_suffix('.tmp.xlsx').replace(path)
    x.h.save(ROOT/'files.json',dict(exceptions=x.sha(x.EXCEPTIONS),report=x.sha(REPORT),pending=x.sha(PENDING),
                                   imported=len(result),remaining=len(remain),fully_duplicate_rows=exact_count))
    print(x.j(dict(imported=len(result),remaining=len(remain),fully_duplicate_rows=exact_count,pending_groups=6,pending_rows=13)))


def main():
    configure();parser=argparse.ArgumentParser();parser.add_argument('mode',choices=['audit','prepare','fixture','apply','finish']);parser.add_argument('--sha256');args=parser.parse_args()
    if args.mode=='audit':
        records,_,names,services=audit()
        print(x.j(dict(ready=sum(not r['issues'] for r in records),names=len(names),service_names=len(services),
            classes=len({(r['class_name'],r['owner_id']) for r in records if not r['issues'] and r.get('missing_class')}),
            categories=Counter(r['duplicate_category'] for r in records),issues=Counter(v for r in records for v in r['issues']))));return
    if args.mode=='prepare':prepare();return
    plan=x.loads((ROOT/'plan.json').read_text())
    x.h.require(args.sha256==x.h.digest(plan),'Exact plan hash required')
    x.h.require(plan['wrapper_sha256']==x.sha(Path(__file__)),'Wrapper changed after planning')
    if args.mode=='finish':finish(plan)
    else:getattr(x,args.mode)(plan)


if __name__=='__main__':main()
