#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Read-only paired comparison of the twelve remaining historical enrollments."""
from datetime import datetime
import json
import os
from pathlib import Path
import zipfile

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
import import_education_exceptions_20261027 as x
from archive_matching_education_students_20261028 import clone_cell, PENDING
from export_education_assignment_comparison_20261028 import OUTPUT as PREVIOUS

ROOT=Path('/opt/zsjos-runtime/backups/education-remaining-comparison-20261028')
OUTPUT=x.ROOT/'output/教务学员_12条规格及班主任差异对比_20261028.xlsx'
NOTES={('学会公共营养师',166):'班型不同：VIP高端班 / 超值精品班；系统班级未确定考期',
       ('学会公共营养师',179):'班型不同：超值精品班 / VIP高端班；系统手机号为空，已确认复用唯一微信档案',
       ('学会公共营养师',178):'仅班主任不同；按用户要求暂时搁置',
       ('人社健康管理师',154):'等级与班型不同：三级超值精品 / 二级VIP',
       ('心理咨询师',68):'原表机构/规格名称与系统不同，未明确班型；系统初级VIP3980，且未确定考期',
       ('执业中西药师',59):'班型不同：无忧通关 / 基础精讲，均全科',
       ('执业中西药师',70):'范围不同：全科 / 单科，均基础精讲',
       ('执业中西药师',73):'班型与范围不同：基础精讲单科 / 金牌押题全科'}
for _row in (10,28,32,78):NOTES[('执业中西药师',_row)]='班型不同：基础精讲 / 金牌押题，均全科'


def run():
    env=os.getenv('ZSJOS_AGENT_ENV')
    if env is None:env=Path('/etc/zsjos/agent-environment').read_text().strip()
    x.h.require(env=='test','Test-only export');os.umask(0o077);ROOT.mkdir(mode=0o700,parents=True,exist_ok=True)
    inputs=(x.MASTER,x.EXCEPTIONS,PENDING,PREVIOUS);hashes={p.name:x.sha(p) for p in inputs}
    if (ROOT/'delivery.json').exists():
        old=x.h.read(ROOT/'delivery.json');x.h.require(old['source_hashes']==hashes and old['output_sha256']==x.sha(OUTPUT),'Changed input/output; preserve user review')
        print('Replay: no changes');return
    x.h.require(not OUTPUT.exists(),'Do not overwrite existing review')
    _,records=x.source_rows(x.EXCEPTIONS);x.h.require(len(records)==12 and {(r['sheet'],r['row']) for r in records}==set(NOTES),'Remaining scope changed')
    # Keep the deferred owner decision last so the active specification review is contiguous.
    records.sort(key=lambda r:((r['sheet'],r['row'])==('学会公共营养师',178),r['exception_row']))
    state=x.h.snapshot();classes={c['id']:c for c in state['classes']};users={u['id']:u for u in state['users']}
    master=openpyxl.load_workbook(x.MASTER);exceptions=openpyxl.load_workbook(x.EXCEPTIONS)
    selected=[]
    for r in records:
        claims=[c for c in state['claims'] if c['contact_value']==r['mobile']]
        x.h.require(len(claims)==1 and claims[0]['person_id'] and claims[0]['reservation_key'] is None,'Claim differs')
        ps=[p for p in state['people'] if p['id']==claims[0]['person_id']]
        x.h.require(len(ps)==1 and ps[0]['name'].strip()==r['name'] and r['mobile'] in (ps[0]['mobile'],ps[0]['wechat_id']),'Identity differs')
        products={p['product_ref'] for p in state['products'] if p['name']==r['product']}
        services=[]
        for s in state['services']:
            snap=json.loads(s['service_snapshot'] or '{}')
            if s['person_id']==ps[0]['id'] and snap.get('productRef') in products and snap.get('name',snap.get('productName'))==r['product']:services.append((s,snap))
        x.h.require(len(services)==1,'Same-product comparison ambiguous');s,snap=services[0];c=classes[s['class_id']]
        x.h.require(c['homeroom_user_id'] in users and s['owner_user_id'] in users,'Assignment user unavailable')
        held=(r['sheet'],r['row'])==('学会公共营养师',178)
        x.h.require((snap['skuName']==r['sku'])==held,'SKU classification drift')
        if held:x.h.require(c['class_name']==r['class_name'] and users[c['homeroom_user_id']]['nickname']!=r['planner'],'Held difference changed')
        selected.append(dict(source=r,person=ps[0],service=s,snapshot=snap,cl=c,homeroom=users[c['homeroom_user_id']],owner=users[s['owner_user_id']],held=held))
    book=openpyxl.Workbook();ws=book.active;ws.title='逐行对比'
    ws.append(['序号','数据来源','姓名','手机号','微信号','课程','课程规格','班级','班主任（原表规划师）','服务负责人',
        '服务显示姓名','差异说明','您的处理意见','补充说明','总表工作表','总表行号','处理状态',
        '服务状态（原值）','接收状态（原值）','交付阶段（原值）','教学是否关闭','服务创建时间'])
    info=book.create_sheet('说明与汇总')
    for row in [('项目','说明'),('核对时间',datetime.now().strftime('%Y-%m-%d %H:%M:%S 北京时间')),('记录数','12组、24数据行：11组规格差异，1组班主任差异'),
                ('排列','每组先原表数据、后系统数据；相同字段上下对齐；黄色单元格为差异'),('填写','在原表数据行绿色“您的处理意见”和“补充说明”栏填写；原有意见不作预选'),
                ('班主任口径','系统班主任为当前班级实际班主任；服务负责人独立显示'),('暂缓项','最后一组：学会公共营养师178行班主任差异，按用户要求暂时搁置'),
                ('完整资料','原表完整信息保留全部44个字段，原异常记录保留全部21列；无对应系统字段留空'),
                ('数据范围','系统课程行仅展示本次同产品待核对服务；不代表学员所有其他课程'),('本次操作','只读系统并生成对照，不修改系统或任何源表，不执行处理意见')]:info.append(row)
    full=book.create_sheet('原表完整信息');headers=[c.value for c in master[records[0]['sheet']+'-学员导入填写'][1]]
    x.h.require(len(headers)==44,'Master fields differ');full.append(['序号','总表工作表','总表行号','差异说明']+headers)
    history=book.create_sheet('原异常记录');history.append(['序号']+[c.value for c in exceptions['异常与处理总表'][1]])
    for i,a in enumerate(selected,1):
        r=a['source'];p=a['person'];s=a['service'];snap=a['snapshot'];c=a['cl'];note=NOTES[(r['sheet'],r['row'])];status='暂缓：班主任差异' if a['held'] else '待确认：课程规格'
        src=master[r['sheet']+'-学员导入填写']
        x.h.require([c.value for c in src[1]]==headers,'Master headings differ')
        x.h.require([x.norm(src.cell(r['row'],col).value) for col in (1,2,4,5,6,7)]==[r['name'],r['mobile'],r['product'],r['sku'],r['class_name'],r['planner']],'Source mismatch')
        ws.append([i,'原表数据',r['name'],r['mobile'],src.cell(r['row'],3).value,r['product'],r['sku'],r['class_name'],r['planner'],None,None,note,None,None,r['sheet'],r['row'],status,None,None,None,None,None])
        ws.append([i,'系统数据',p['name'],p['mobile'],p['wechat_id'],snap.get('name',snap.get('productName')),snap['skuName'],c['class_name'],a['homeroom']['nickname'],a['owner']['nickname'],s['student_name_snapshot'],note,None,None,r['sheet'],r['row'],status,s['status'],s['acceptance_status'],s['delivery_stage'],'是' if s['teaching_closed'] in (1,b'\1') else '否',s['create_time']])
        for col,val in enumerate([i,r['sheet'],r['row'],note],1):full.cell(i+1,col,val)
        for col in range(1,45):clone_cell(src.cell(r['row'],col),full.cell(i+1,col+4))
        history.cell(i+1,1,i)
        for col in range(1,22):clone_cell(exceptions['异常与处理总表'].cell(r['exception_row'],col),history.cell(i+1,col+1))
    for sh in book:x.format_sheet(sh);sh.sheet_view.zoomScale=85
    for sh,name in ((ws,'RemainingPaired'),(full,'RemainingFull'),(history,'RemainingOriginal')):
        table=Table(displayName=name,ref=sh.dimensions);table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=False);sh.add_table(table)
    widths=[8,13,17,18,20,20,53,42,29,29,17,50,32,32,22,12,24,19,19,19,17,24]
    for col,width in enumerate(widths,1):ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width=width
    ws.freeze_panes='E2';full.freeze_panes='G2';history.freeze_panes='D2';info.column_dimensions['A'].width=20;info.column_dimensions['B'].width=110
    line=Side(style='thin',color='8397B0')
    for top in range(2,26,2):
        for row in (top,top+1):
            ws.row_dimensions[row].height=70
            for cell in ws[row]:
                cell.fill=PatternFill('solid',fgColor='F2F6FC' if row==top else 'FFFFFF');cell.alignment=Alignment(vertical='center',wrap_text=True)
                cell.border=Border(bottom=line if row==top+1 else Side())
            ws.cell(row,2).font=Font(bold=True,color='264478' if row==top else '548235')
            for col in (4,5):ws.cell(row,col).number_format='@'
        for col in range(3,10):
            if ws.cell(top,col).value!=ws.cell(top+1,col).value:
                for row in (top,top+1):ws.cell(row,col).fill=PatternFill('solid',fgColor='FFF2CC');ws.cell(row,col).font=Font(color='9C0006')
        for col in (13,14):ws.cell(top,col).fill=PatternFill('solid',fgColor='E2F0D9')
    temp=OUTPUT.with_suffix('.tmp.xlsx');book.save(temp)
    with zipfile.ZipFile(temp) as z:x.h.require(z.testzip() is None,'Invalid ZIP')
    saved=openpyxl.load_workbook(temp)
    x.h.require(saved['逐行对比'].max_row==25 and list(saved['逐行对比'].values)==list(ws.values),'Saved pairs differ')
    for i,a in enumerate(selected,2):
        r=a['source'];src=master[r['sheet']+'-学员导入填写']
        for col in range(1,45):
            v,b=src.cell(r['row'],col),saved['原表完整信息'].cell(i,col+4)
            x.h.require((v.value,v.data_type,v.number_format)==(b.value,b.data_type,b.number_format),'Source44fields differ')
        x.h.require(tuple(c.value for c in saved['原异常记录'][i])[1:]==tuple(c.value for c in exceptions['异常与处理总表'][r['exception_row']]),'Original exception differs')
    x.h.require(all(x.sha(p)==hashes[p.name] for p in inputs),'Input changed during export')
    x.h.save(ROOT/'selection.json',selected);temp.replace(OUTPUT)
    x.h.save(ROOT/'delivery.json',dict(source_hashes=hashes,output_sha256=x.sha(OUTPUT),pairs=12,data_rows=24,sku_differences=11,teacher_hold=1,database_writes=0))
    print('PASS:12pairs/24rows;11SKU differences+1teacher hold;44original fields;source files unchanged;database writes0')


if __name__=='__main__':run()
