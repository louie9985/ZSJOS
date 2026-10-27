#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Reconcile the 22-row follow-up workbook; SQL connections are strictly read-only.

Exclusion annotations concern import rows. They do not complete refunds or delete
existing teaching services. Source and master remain unchanged; see operation log.
"""
from collections import Counter
import os
from pathlib import Path
import re
import shutil
import zipfile

import openpyxl

import import_education_exceptions_20261027 as x

SOURCE = x.ROOT / 'output/异常学员(1)1111.xlsx'
REPORT = x.ROOT / 'output/异常学员(1)1111_核对更新结果.xlsx'
EVIDENCE = Path('/opt/zsjos-runtime/backups/education-1111-review-20261027')
NOTE = '2026-10-27：已核对异常学员(1)1111；身份问题未新增明确结论，继续保留，不重复导入。'


def run():
    env = os.getenv('ZSJOS_AGENT_ENV')
    env = env if env is not None else Path('/etc/zsjos/agent-environment').read_text().strip()
    x.h.require(env == 'test', 'Test-only review')
    os.umask(0o077)
    EVIDENCE.mkdir(mode=0o700, parents=True, exist_ok=True)
    hashes = {p.name:x.sha(p) for p in (SOURCE, x.MASTER, x.EXCEPTIONS)}
    if (EVIDENCE/'delivery.json').exists():
        old = x.h.read(EVIDENCE/'delivery.json')
        x.h.require(old['source_sha256']==hashes[SOURCE.name] and old['master_sha256']==hashes[x.MASTER.name]
                    and old['exceptions_sha256']==hashes[x.EXCEPTIONS.name] and old['report_sha256']==x.sha(REPORT),
                    'Delivered inputs/outputs changed; fresh review required')
        print('Replay: no file changes')
        return
    for path in (SOURCE, x.MASTER, x.EXCEPTIONS):
        dest = EVIDENCE/path.name
        if dest.exists():
            x.h.require(x.sha(dest)==hashes[path.name], 'Existing backup differs; refuse overwrite')
        else:
            shutil.copy2(path,dest); dest.chmod(0o600)
    master = openpyxl.load_workbook(x.MASTER,data_only=True)
    exceptions = openpyxl.load_workbook(x.EXCEPTIONS)
    original = openpyxl.load_workbook(x.EXCEPTIONS)
    sheet = exceptions['异常与处理总表']
    exception_rows = {(r[7],int(r[8])):i for i,r in enumerate(list(sheet.values)[1:],2)}
    source = openpyxl.load_workbook(SOURCE,data_only=True)
    data = list(source.active.values)[1:]
    x.h.require(len(data)==22, 'Source count changed')
    before = x.h.snapshot()
    x.h.save(EVIDENCE/'before.json',before)
    session = x.h.auth()
    results, matched_ids, touched = [],set(),set()
    for input_row, raw in enumerate(data,2):
        match = re.fullmatch(r'(.+)第(\d+)行',raw[7])
        x.h.require(match is not None,'Invalid master reference')
        short,rowno = match[1],int(match[2])
        key = short,rowno
        expected = list(master[short+'-学员导入填写'].values)[rowno-1]
        mobile = x.norm(raw[2])
        people = [p for p in before['people'] if mobile in (p['mobile'],p['wechat_id'])]
        claims = {c['person_id'] for c in before['claims'] if c['contact_value']==mobile and c['person_id']}
        person_ids = {p['id'] for p in people}|claims
        services = [s for s in before['services'] if s['person_id'] in person_ids]
        result = dict(input_row=input_row,sheet=short,row=rowno,sku=re.split('[【（]',raw[3])[0],
                      class_name=raw[4],service_ids=[],status=None,detail=None)
        if '删除' in raw[3]:
            x.h.require(not any(expected) and key not in exception_rows,'Exclusion not yet reflected; review required')
            result.update(status='已排除导入',detail='总表对应行已清空，当前异常表无此记录；保持排除。')
            if '退费审批中' in raw[3]:
                x.h.require(services,'Refund-related historical service missing')
                result['detail'] += '系统仍有历史二级服务；退费审批中说明不等于已退费，不关闭或删除该服务。'
                matched_ids.update(s['id'] for s in services)
        else:
            x.h.require([x.norm(v) for v in (raw[1],raw[2],raw[3],raw[4],raw[5])]
                        ==[x.norm(v) for v in (expected[0],expected[1],expected[4],expected[5],expected[6])],
                        'New source/master correction requires explicit mapping')
            if key in exception_rows:
                x.h.require(input_row in (2,3,4,6),'Unexpected additional unresolved row')
                result.update(status='待确认身份',detail=('同一手机号对应不同姓名，尚无明确修正结论。' if input_row==6
                              else '同名记录使用不同号码，尚未明确是否同一人或号码有误。'))
                erow = exception_rows[key]
                old = sheet.cell(erow,7).value
                sheet.cell(erow,7,(str(old)+'\n' if old else '')+NOTE)
                touched.add((sheet.title,erow,7))
            else:
                x.h.require(len(people)==1 and person_ids=={people[0]['id']} and people[0]['name'].strip()==x.norm(raw[1]),
                            'Completed identity mismatch')
                classes = {c['id']:c for c in before['classes']}
                owners = [u for u in before['users'] if u['nickname']==x.norm(raw[5]) and u['status']==0]
                x.h.require(len(owners)==1,'Owner unavailable')
                exact = [s for s in services if x.json.loads(s['service_snapshot'] or '{}').get('skuName')==raw[3]
                         and classes.get(s['class_id'],{}).get('class_name')==raw[4] and s['owner_user_id']==owners[0]['id']]
                x.h.require(len(exact)==1,'Completed service specification/destination mismatch')
                rel = exact[0]; rid=rel['id']; matched_ids.add(rid)
                x.h.require(rel['acceptance_status']=='accepted' and rel['delivery_stage']=='supervision','Completed state differs')
                context=x.h.api(session,'GET',f'/zsjos/student/service/{rid}/contact-context')
                stages={a['code']:a for a in context['deliveryStages']}
                x.h.require(context['deliveryStage']=='supervision' and stages['first_contact']['status']==stages['study_plan']['status']=='done'
                            and all(a['completed'] for a in context['firstContactProgress']['items']), 'API completion differs')
                records=x.h.api(session,'GET',f'/zsjos/student/service/{rid}/contact-records?pageNo=1&pageSize=100')
                for kind,_,_ in x.TYPES:
                    x.h.require(any(r['contactType']==kind and r['successful'] and r['remark']==x.REMARK for r in records['list']),
                                'Missing exact completion remark')
                result.update(status='已导入并完成',detail='规格、班级、班主任一致；常规督学，首联与学习计划已完成，无需重复导入。',service_ids=[rid])
                if input_row==23:
                    sku=[s for s in before['skus'] if s['sku_ref']==x.json.loads(rel['service_snapshot'])['skuRef']]
                    x.h.require(len(sku)==1 and sku[0]['status']==1 and sku[0]['price']==0,'Historical zero/disabled specification differs')
                    result['detail']+='该历史规格已按后续确认建为0元且停用，旧“不新增配置”说明已被后续决定替代。'
        results.append(result)
    counts=Counter(r['status'] for r in results)
    x.h.require(counts=={'已导入并完成':11,'已排除导入':7,'待确认身份':4},'Unexpected result partition')
    summary=exceptions['汇总']
    start=summary.max_row+1
    for row in [('异常学员(1)1111复核',22,'已完成11条、已排除7条、身份待确认4条；本轮新增导入0条'),
                ('本轮复核结果文件',None,REPORT.name)]: summary.append(row)
    summary.tables['SummaryTable'].ref=f'A1:C{summary.max_row}'
    summary.auto_filter.ref=summary.tables['SummaryTable'].ref
    report=openpyxl.Workbook(); report.active.title='汇总'
    report.active.append(['项目','数量','说明'])
    for label,count in counts.items(): report.active.append([label,count,'逐条核对总表及当前测试库'])
    report.active.append(['本轮新增导入',0,'已完成不重复；身份待确认不自动放行'])
    report.active.append(['当前异常表剩余',sheet.max_row-1,'本轮仅更新4条核对说明'])
    report.active.append(['退费相关历史服务',1,'对应学会公共营养师185、188行；导入排除已落实，系统历史服务保留'])
    for title,subset in [('逐条核对',results),('待确认4条',[r for r in results if r['status']=='待确认身份'])]:
        ws=report.create_sheet(title)
        ws.append(['1111文件行号','总表工作表','总表行号','课程规格','班级','当前结果','核对说明'])
        for r in subset: ws.append([r[k] for k in ('input_row','sheet','row','sku','class_name','status','detail')])
    for ws in report:x.format_sheet(ws)
    for old_sheet in original:
        new_sheet=exceptions[old_sheet.title]
        for row in old_sheet:
            for cell in row:
                if (old_sheet.title,cell.row,cell.column) not in touched:
                    x.h.require(cell.value==new_sheet.cell(cell.row,cell.column).value,'Unrelated cell changed')
    x.h.require(sheet.max_row==242 and len(touched)==4,'Exception count/annotation scope differs')
    after_services=x.h.query('SELECT * FROM zsjos_service_relation WHERE tenant_id=1 AND deleted=0 AND id IN ('
                            +','.join(map(str,sorted(matched_ids)))+') ORDER BY id')
    x.h.require(x.h.encoded(after_services)==x.h.encoded([s for s in before['services'] if s['id'] in matched_ids]),'Observed service changed during review')
    x.h.require(all(x.sha(p)==hashes[p.name] for p in (SOURCE,x.MASTER,x.EXCEPTIONS)),'Workbook changed concurrently')
    etmp=x.EXCEPTIONS.with_suffix('.tmp.xlsx'); rtmp=REPORT.with_suffix('.tmp.xlsx')
    exceptions.save(etmp); report.save(rtmp)
    for path in (etmp,rtmp):
        with zipfile.ZipFile(path) as z:x.h.require(z.testzip() is None,'Invalid workbook archive')
    checked=openpyxl.load_workbook(etmp,data_only=True)
    x.h.require(checked['异常与处理总表'].max_row==242,'Saved exception count differs')
    rc=openpyxl.load_workbook(rtmp,data_only=True)
    x.h.require(rc['逐条核对'].max_row==23 and rc['待确认4条'].max_row==5,'Saved report count differs')
    rtmp.replace(REPORT); etmp.replace(x.EXCEPTIONS)
    x.h.save(EVIDENCE/'results.json',results)
    x.h.save(EVIDENCE/'delivery.json',dict(source_sha256=x.sha(SOURCE),master_sha256=x.sha(x.MASTER),
        exceptions_sha256=x.sha(x.EXCEPTIONS),report_sha256=x.sha(REPORT),counts=dict(counts),database_writes=0,
        api_contexts=11,api_histories=11,exception_rows=241))
    print(x.j(dict(counts=counts,new_imports=0,database_writes=0,exception_rows=241,annotated_rows=4)))


if __name__=='__main__':
    run()
