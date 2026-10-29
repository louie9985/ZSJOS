#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Accept the user's final system-precedence decision; Excel changes only."""
from datetime import datetime
import json
import os
from pathlib import Path
import shutil
import zipfile

import openpyxl
import import_education_exceptions_20261027 as x
import export_remaining_education_comparison_20261028 as review_export
from archive_matching_education_students_20261028 import clone_cell, PENDING
from export_education_assignment_comparison_20261028 import OUTPUT as PREVIOUS

ROOT=Path('/opt/zsjos-runtime/backups/education-final-system-20261028')
REVIEW=review_export.OUTPUT
OUTPUT=x.ROOT/'output/教务学员_最后12条确认处理结果_20261028.xlsx'
INPUTS=(x.MASTER,x.EXCEPTIONS,REVIEW,PREVIOUS,PENDING)


def resolve(selection,state):
    people={p['id']:p for p in state['people']};services={s['id']:s for s in state['services']}
    classes={c['id']:c for c in state['classes']};users={u['id']:u for u in state['users']};result=[]
    for entry in selection:
        r=entry['source'];s=services[entry['service']['id']];p=people[s['person_id']];c=classes[s['class_id']];u=users[c['homeroom_user_id']]
        snap=json.loads(s['service_snapshot'] or '{}');claims=[a for a in state['claims'] if a['contact_value']==r['mobile']]
        x.h.require(len(claims)==1 and claims[0]['person_id']==p['id'] and claims[0]['reservation_key'] is None,'Claim identity changed')
        x.h.require(p['name'].strip()==r['name']==s['student_name_snapshot'] and r['mobile'] in (p['mobile'],p['wechat_id']),'Identity changed')
        x.h.require(s['person_id']==entry['person']['id'] and snap==entry['snapshot'] and s['class_id']==entry['cl']['id']
            and c['class_name']==entry['cl']['class_name'] and u['nickname']==entry['homeroom']['nickname']
            and s['owner_user_id']==c['homeroom_user_id']==entry['owner']['id'],'Reviewed assignment or SKU changed')
        # System precedence also preserves a reviewed cancelled service; it must
        # never be silently reopened or represented as an active enrollment.
        x.h.require(all(s[k]==entry['service'][k] for k in ('status','teaching_closed','teaching_delivery','acceptance_status','delivery_stage')),
                    'Reviewed lifecycle changed')
        result.append(dict(source=r,service=s,person=p,cl=c,owner=u,claim=claims[0],snapshot=snap))
    return result


def run():
    env=os.getenv('ZSJOS_AGENT_ENV')
    if env is None:env=Path('/etc/zsjos/agent-environment').read_text().strip()
    x.h.require(env=='test','Test-only execution');os.umask(0o077);ROOT.mkdir(mode=0o700,parents=True,exist_ok=True)
    if (ROOT/'delivery.json').exists():
        old=x.h.read(ROOT/'delivery.json');x.h.require(all(x.sha(p)==old['hashes'][p.name] for p in (*INPUTS,OUTPUT)),'Delivered files changed')
        print('Replay: no changes');return
    x.h.require(not OUTPUT.exists(),'Existing report must not be overwritten')
    hashes={p.name:x.sha(p) for p in INPUTS};selection=x.h.read(review_export.ROOT/'selection.json')
    state=x.h.snapshot();resolved=resolve(selection,state);keys={(a['source']['sheet'],a['source']['row']) for a in resolved}
    _,records=x.source_rows(x.EXCEPTIONS);x.h.require(len(records)==12 and {(r['sheet'],r['row']) for r in records}==keys,'Exception scope differs')
    master=openpyxl.load_workbook(x.MASTER);old_master=openpyxl.load_workbook(x.MASTER)
    exc=openpyxl.load_workbook(x.EXCEPTIONS);old_exc=openpyxl.load_workbook(x.EXCEPTIONS)
    review=openpyxl.load_workbook(REVIEW);old_review=openpyxl.load_workbook(REVIEW)
    prior=openpyxl.load_workbook(PREVIOUS);old_prior=openpyxl.load_workbook(PREVIOUS)
    for path in INPUTS:
        backup=ROOT/path.name
        if backup.exists():x.h.require(x.sha(backup)==hashes[path.name],'Backup differs from original source')
        else:shutil.copy2(path,backup);backup.chmod(0o600)
    edits=[]
    for a in resolved:
        r=a['source'];sh=master[r['sheet']+'-学员导入填写']
        x.h.require([x.norm(sh.cell(r['row'],col).value) for col in (1,2,4,5,6,7)]==[r['name'],r['mobile'],r['product'],r['sku'],r['class_name'],r['planner']],'Master source drift')
        values={4:a['snapshot'].get('name',a['snapshot'].get('productName')),5:a['snapshot']['skuName'],6:a['cl']['class_name'],7:a['owner']['nickname']}
        for col,val in values.items():
            if sh.cell(r['row'],col).value!=val:
                edits.append(dict(sheet=sh.title,row=r['row'],column=col,before=sh.cell(r['row'],col).value,after=val));sh.cell(r['row'],col,val)
    x.h.require(len(edits)==14 and sum(e['column']==5 for e in edits)==11 and sum(e['column']==6 for e in edits)==2 and sum(e['column']==7 for e in edits)==1,'Master edit scope changed')
    es=exc['异常与处理总表'];es.delete_rows(2,es.max_row-1);es.tables['ExceptionTable'].ref='A1:U1';es.auto_filter.ref='A1:U1'
    cs=exc['关联记录对比'];removed=0
    for row in range(cs.max_row,1,-1):
        if (cs.cell(row,4).value,cs.cell(row,5).value) in keys:cs.delete_rows(row);removed+=1
    cs.tables['ComparisonTable'].ref=f'A1:K{cs.max_row}';cs.auto_filter.ref=cs.tables['ComparisonTable'].ref
    summary=exc['汇总']
    for row in range(2,summary.max_row+1):
        if summary.cell(row,1).value=='当前保留异常':summary.cell(row,2,0)
        if summary.cell(row,1).value=='当前异常构成':summary.cell(row,2,0);summary.cell(row,3,'本批异常全部处理完成；最终12条按系统现有数据确认')
    summary.append(['2026-10-28最后12条按系统确认',12,'11条规格差异及1条班主任差异均以系统为准；总表14处同步，异常清零'])
    summary.tables['SummaryTable'].ref=f'A1:C{summary.max_row}';summary.auto_filter.ref=summary.tables['SummaryTable'].ref
    timestamp=datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    statuses={(a['source']['sheet'],a['source']['row']):a['service']['status'] for a in resolved}
    ws=review['逐行对比'];x.h.require(ws.max_row==25,'Current review count differs')
    for row in range(2,ws.max_row+1):
        x.h.require((ws.cell(row,15).value,ws.cell(row,16).value) in keys,'Review source differs')
        cancelled=statuses[(ws.cell(row,15).value,ws.cell(row,16).value)]=='cancelled'
        ws.cell(row,17,'已处理：系统已取消，保持取消' if cancelled else '已确认：以系统为准')
        if ws.cell(row,2).value=='原表数据':
            old=ws.cell(row,13).value
            x.h.require(old in (None,'','以系统为准'),'Conflicting saved user decision')
            ws.cell(row,13,'以系统为准')
    info=review['说明与汇总']
    for row in info:
        if row[0].value=='暂缓项':row[1].value='原暂缓班主任差异已随本次全部以系统为准的决定确认'
        if row[0].value=='本次操作':row[1].value='12条已确认以系统为准；同步总表、清理异常，系统服务未修改'
        if row[0].value=='填写':row[1].value='已记录用户最终决定；逐行原表/系统为核对时快照，现行总表已同步系统值'
    info.append(['最终确认时间',timestamp]);info.auto_filter.ref=info.dimensions
    paired=prior['逐行对比'];archive=prior['已处理对照（处理前快照）']
    x.h.require(paired.max_row==3,'Previous review pending scope differs')
    for row in range(2,4):
        x.h.require((paired.cell(row,15).value,paired.cell(row,16).value)==('学会公共营养师',178),'Previous review pending item differs')
        dest=archive.max_row+1
        for col,cell in enumerate(paired[row],1):clone_cell(cell,archive.cell(dest,col))
        archive.cell(dest,22,'用户已确认以系统班主任为准；总表已同步，异常已清理')
        if archive.cell(dest,2).value=='原表数据':archive.cell(dest,13,'以系统为准')
    paired.delete_rows(2,2);paired.tables['AssignmentPaired'].ref='A1:U1';paired.auto_filter.ref='A1:U1';archive.auto_filter.ref=archive.dimensions
    for row in prior['说明与汇总']:
        if row[0].value in ('待确认原表记录','系统同课程同规格服务','班主任不同'):row[1].value=0
        if row[0].value=='本次操作':row[1].value='最后1条班主任差异已确认以系统为准；全部22组处理完成，快照归档保留'
    out=openpyxl.Workbook();out.active.title='汇总'
    for row in [('项目','数量','说明'),('最终确认记录',12,'均以系统现有课程规格、班级、班主任为准'),('总表同步单元格',14,'11处规格、2处班级、1处班主任'),
                ('异常表剩余',0,'12条异常已移出；既有3条其他来源关联对比历史保留'),('系统已取消服务',1,'执业中西药师70行保持取消，不作为有效在读'),('数据库写入',0,'原系统服务和状态完全保留')]:out.active.append(row)
    full=out.create_sheet('总表同步前后完整信息');headers=[c.value for c in old_master[resolved[0]['source']['sheet']+'-学员导入填写'][1]]
    x.h.require(len(headers)==44,'Original fields differ');full.append(['序号','数据版本','总表工作表','总表行号','确认结果']+headers)
    for i,a in enumerate(resolved,1):
        r=a['source']
        for version,source in (('同步前原表',old_master),('同步后总表',master)):
            row=full.max_row+1
            decision='以系统为准，已处理；系统已取消，保持取消' if a['service']['status']=='cancelled' else '以系统为准，已处理'
            for col,val in enumerate([i,version,r['sheet'],r['row'],decision],1):full.cell(row,col,val)
            sh=source[r['sheet']+'-学员导入填写']
            for col in range(1,45):clone_cell(sh.cell(r['row'],col),full.cell(row,col+5))
    history=out.create_sheet('原异常完整记录')
    for row in old_exc['异常与处理总表']:
        for cell in row:clone_cell(cell,history.cell(cell.row,cell.column))
    compare=out.create_sheet('已确认逐行对比')
    for row in ws:
        for cell in row:clone_cell(cell,compare.cell(cell.row,cell.column))
    for sh in out:x.format_sheet(sh)
    full.freeze_panes='H2';compare.freeze_panes='E2'
    books=((x.MASTER,master),(x.EXCEPTIONS,exc),(REVIEW,review),(PREVIOUS,prior),(OUTPUT,out))
    for p,book in books:
        temp=p.with_suffix('.tmp.xlsx');book.save(temp)
        with zipfile.ZipFile(temp) as z:x.h.require(z.testzip() is None,'Invalid workbook ZIP')
    saved={p:openpyxl.load_workbook(p.with_suffix('.tmp.xlsx')) for p,_ in books}
    allowed={(e['sheet'],e['row'],e['column']):e['after'] for e in edits}
    for sh in old_master:
        for row in sh:
            for cell in row:
                new=saved[x.MASTER][sh.title].cell(cell.row,cell.column)
                x.h.require(new.value==allowed.get((sh.title,cell.row,cell.column),cell.value),'Unexpected master cell change')
    x.h.require(saved[x.EXCEPTIONS]['异常与处理总表'].max_row==1,'Exception not empty')
    expected_comparison=[r for r in list(old_exc['关联记录对比'].values)[1:] if (r[3],r[4]) not in keys]
    x.h.require(list(saved[x.EXCEPTIONS]['关联记录对比'].values)[1:]==expected_comparison,'Unrelated comparison history changed')
    x.h.require(list(saved[x.EXCEPTIONS]['聊天处理依据'].values)==list(old_exc['聊天处理依据'].values),'Historical decisions changed')
    for p,old in ((REVIEW,old_review),(PREVIOUS,old_prior)):
        for title in ('原表完整信息','原异常记录'):x.h.require(list(saved[p][title].values)==list(old[title].values),'Historical source changed')
    for row in list(saved[OUTPUT]['总表同步前后完整信息'].rows)[1:]:
        source=old_master if row[1].value=='同步前原表' else saved[x.MASTER]
        for col in range(1,45):
            a=source[row[2].value+'-学员导入填写'].cell(row[3].value,col);b=row[col+4]
            x.h.require((a.value,a.data_type,a.number_format)==(b.value,b.data_type,b.number_format),'Full source archive differs')
    x.h.require(list(saved[OUTPUT]['原异常完整记录'].values)==list(old_exc['异常与处理总表'].values),'Exception archive differs')
    x.h.require(saved[PREVIOUS]['逐行对比'].max_row==1 and saved[PREVIOUS]['已处理对照（处理前快照）'].max_row==45,'Previous review completion differs')
    x.h.require(x.h.encoded(resolve(selection,x.h.snapshot()))==x.h.encoded(resolved),'Scoped live records changed')
    x.h.require(all(x.sha(p)==hashes[p.name] for p in INPUTS),'Files changed concurrently')
    x.h.save(ROOT/'resolved.json',resolved);x.h.save(ROOT/'master-edits.json',edits)
    # The standalone archive is installed before any source cleanup, keeping all originals recoverable.
    for p in (OUTPUT,x.MASTER,x.EXCEPTIONS,REVIEW,PREVIOUS):p.with_suffix('.tmp.xlsx').replace(p)
    x.h.save(ROOT/'delivery.json',dict(confirmed=12,master_cells=14,remaining_exceptions=0,database_writes=0,hashes={p.name:x.sha(p) for p in (*INPUTS,OUTPUT)}))
    print('PASS:12system decisions;14master cells synced;0exceptions;both review queues resolved;database writes0')


if __name__=='__main__':run()
