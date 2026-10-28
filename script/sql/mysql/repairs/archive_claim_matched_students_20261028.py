#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Archive three explicitly authorized exact matches without changing the database."""
import json
import os
from pathlib import Path
import shutil
import zipfile

import openpyxl
import import_education_exceptions_20261027 as x
from archive_matching_education_students_20261028 import clone_cell, PENDING
from export_education_assignment_comparison_20261028 import OUTPUT as REVIEW

ROOT=Path('/opt/zsjos-runtime/backups/education-claim-matched-20261028')
OUTPUT=x.ROOT/'output/教务学员_三条已有服务清理记录_20261028.xlsx'
TARGETS={('注册营养师',175),('卫健委体重管理',12),('中医营养配餐',30)}
EMPTY_MASTER={('卫健委体重管理',12),('中医营养配餐',30)}


def resolve(records,state):
    result=[]
    for r in records:
        claims=[c for c in state['claims'] if c['contact_value']==r['mobile']]
        x.h.require(len(claims)==1 and claims[0]['person_id'] and claims[0]['reservation_key'] is None,'Claim unavailable or ambiguous')
        people=[p for p in state['people'] if p['id']==claims[0]['person_id']]
        x.h.require(len(people)==1 and people[0]['name'].strip()==r['name'] and people[0]['mobile']==r['mobile'],'Claim identity differs')
        owners=[u for u in state['users'] if u['nickname']==r['planner'] and u['status']==0]
        x.h.require(len(owners)==1,'Planner unavailable')
        classes={c['id'] for c in state['classes'] if c['class_name']==r['class_name'] and c['homeroom_user_id']==owners[0]['id'] and c['status']=='SERVING'}
        products={p['product_ref'] for p in state['products'] if p['name']==r['product']}
        services=[]
        for s in state['services']:
            snap=json.loads(s['service_snapshot'] or '{}')
            if (s['person_id']==people[0]['id'] and s['student_name_snapshot']==r['name'] and s['class_id'] in classes
                and s['owner_user_id']==owners[0]['id'] and s['status']=='active' and s['teaching_closed'] in (0,b'\0')
                and s['teaching_delivery'] in (1,b'\1') and snap.get('name',snap.get('productName'))==r['product']
                and snap.get('productRef') in products and snap.get('skuName')==r['sku']):services.append(s['id'])
        x.h.require(len(services)==1,'Exact service not unique')
        result.append(dict(sheet=r['sheet'],row=r['row'],person_id=people[0]['id'],service_id=services[0],claim_id=claims[0]['id']))
    return result


def run():
    env=os.getenv('ZSJOS_AGENT_ENV')
    if env is None:env=Path('/etc/zsjos/agent-environment').read_text().strip()
    x.h.require(env=='test','Test only');os.umask(0o077);ROOT.mkdir(mode=0o700,parents=True,exist_ok=True)
    inputs=(x.MASTER,x.EXCEPTIONS,PENDING,REVIEW)
    if (ROOT/'delivery.json').exists():
        meta=x.h.read(ROOT/'delivery.json');x.h.require(all(x.sha(p)==meta['hashes'][p.name] for p in (*inputs,OUTPUT)),'Delivered files changed')
        print('Replay: no changes');return
    x.h.require(not OUTPUT.exists(),'Existing archive must not be replaced')
    hashes={p.name:x.sha(p) for p in inputs}
    original,records=x.source_rows(x.EXCEPTIONS)
    chosen=[r for r in records if (r['sheet'],r['row']) in TARGETS]
    x.h.require(len(records)==20 and len(chosen)==3,'Source scope changed')
    state=x.h.snapshot();matches=resolve(chosen,state)
    for p in (x.MASTER,x.EXCEPTIONS):
        backup=ROOT/p.name
        if backup.exists():x.h.require(x.sha(backup)==hashes[p.name],'Backup differs from unchanged source')
        else:shutil.copy2(p,backup);backup.chmod(0o600)
    master=openpyxl.load_workbook(x.MASTER);exc=openpyxl.load_workbook(x.EXCEPTIONS)
    out=openpyxl.Workbook();out.active.title='汇总'
    for row in [('项目','数量','说明'),('已有一致服务归档',3,'占用记录所指档案的姓名、手机号、课程规格、班级和班主任一致'),
                ('异常表剩余',17,'3条需分班、2条无目标课程、11条规格差异、1条班主任差异'),('数据库修改',0,'未合并或删除档案、未修改已有服务'),
                ('总表原行已为空',2,'现有历史备份同样为空，不补造原始字段；3条原异常记录全部21列完整保存在另页')]:out.active.append(row)
    full=out.create_sheet('原表完整信息');headers=[c.value for c in master[chosen[0]['sheet']+'-学员导入填写'][1]]
    x.h.require(len(headers)==44,'Master columns differ');full.append(['总表工作表','总表行号','处理结果']+headers)
    history=out.create_sheet('原异常记录');history.append([c.value for c in original['异常与处理总表'][1]])
    for idx,r in enumerate(chosen,2):
        sh=master[r['sheet']+'-学员导入填写']
        x.h.require([c.value for c in sh[1]]==headers,'Master headers differ')
        empty=(r['sheet'],r['row']) in EMPTY_MASTER
        # Two rows were already cleared before the retained import backups. Keep
        # their blank master cells and complete exception evidence, never invent fields.
        if empty:x.h.require(not any(c.value is not None for c in sh[r['row']]),'Previously cleared row changed')
        else:x.h.require([x.norm(sh.cell(r['row'],col).value) for col in (1,2,4,5,6,7)]==[r['name'],r['mobile'],r['product'],r['sku'],r['class_name'],r['planner']], 'Master identity differs')
        status='系统已有一致服务，归档为已导入'+('；总表原行已为空，完整可用资料见原异常记录' if empty else '')
        for col,value in enumerate([r['sheet'],r['row'],status],1):full.cell(idx,col,value)
        for col in range(1,45):clone_cell(sh.cell(r['row'],col),full.cell(idx,col+3))
        for col in range(1,22):clone_cell(original['异常与处理总表'].cell(r['exception_row'],col),history.cell(idx,col))
    for sh in out:x.format_sheet(sh)
    es=exc['异常与处理总表'];expected=[r for r in list(es.values)[1:] if (r[7],r[8]) not in TARGETS]
    for row in range(es.max_row,1,-1):
        if (es.cell(row,8).value,es.cell(row,9).value) in TARGETS:es.delete_rows(row)
    x.h.require(es.max_row==18,'Remaining row count differs');es.tables['ExceptionTable'].ref='A1:U18';es.auto_filter.ref='A1:U18'
    cs=exc['关联记录对比'];removed=0
    for row in range(cs.max_row,1,-1):
        if (cs.cell(row,4).value,cs.cell(row,5).value) in TARGETS:cs.delete_rows(row);removed+=1
    cs.tables['ComparisonTable'].ref=f'A1:K{cs.max_row}';cs.auto_filter.ref=cs.tables['ComparisonTable'].ref
    summary=exc['汇总']
    for row in range(2,summary.max_row+1):
        if summary.cell(row,1).value=='当前保留异常':summary.cell(row,2,17)
        if summary.cell(row,1).value=='当前异常构成':summary.cell(row,2,17);summary.cell(row,3,'3条需分班、2条无目标课程（均涉及多档案）；11条规格差异、1条班主任差异')
    summary.append(['2026-10-28占用档案已有服务清理',3,'已复核一致并归档；未修改系统档案或服务'])
    summary.tables['SummaryTable'].ref=f'A1:C{summary.max_row}';summary.auto_filter.ref=summary.tables['SummaryTable'].ref
    for path,book in ((x.EXCEPTIONS,exc),(OUTPUT,out)):
        temp=path.with_suffix('.tmp.xlsx');book.save(temp)
        with zipfile.ZipFile(temp) as z:x.h.require(z.testzip() is None,'Invalid ZIP')
    saved=openpyxl.load_workbook(x.EXCEPTIONS.with_suffix('.tmp.xlsx'));archive=openpyxl.load_workbook(OUTPUT.with_suffix('.tmp.xlsx'))
    x.h.require(list(saved['异常与处理总表'].values)[1:]==expected,'Retained cells differ')
    x.h.require(list(saved['聊天处理依据'].values)==list(original['聊天处理依据'].values),'History changed')
    for idx,r in enumerate(chosen,2):
        for col in range(1,45):
            a=master[r['sheet']+'-学员导入填写'].cell(r['row'],col);b=archive['原表完整信息'].cell(idx,col+3)
            x.h.require((a.value,a.data_type,a.number_format)==(b.value,b.data_type,b.number_format),'Original source field differs')
        x.h.require(tuple(c.value for c in archive['原异常记录'][idx])==tuple(c.value for c in original['异常与处理总表'][r['exception_row']]),'Original exception differs')
    x.h.require(resolve(chosen,x.h.snapshot())==matches,'Live exact matching changed')
    x.h.require(all(x.sha(p)==hashes[p.name] for p in inputs),'Source changed concurrently')
    x.h.save(ROOT/'snapshot.json',state);x.h.save(ROOT/'matches.json',matches)
    OUTPUT.with_suffix('.tmp.xlsx').replace(OUTPUT);x.EXCEPTIONS.with_suffix('.tmp.xlsx').replace(x.EXCEPTIONS)
    x.h.save(ROOT/'delivery.json',dict(archived=3,remaining=17,comparison_removed=removed,database_writes=0,hashes={p.name:x.sha(p) for p in (*inputs,OUTPUT)}))
    print(x.j(dict(archived=3,remaining=17,comparison_removed=removed,database_writes=0,original_columns=44)))


if __name__=='__main__':run()
