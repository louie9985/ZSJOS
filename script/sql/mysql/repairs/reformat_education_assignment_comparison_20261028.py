#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Convert the existing comparison into original/system row pairs, without DB access."""
import os
from pathlib import Path
import shutil
import zipfile

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
import import_education_exceptions_20261027 as x
from export_education_assignment_comparison_20261028 import OUTPUT, PENDING

ROOT=Path('/opt/zsjos-runtime/backups/education-assignment-paired-20261028')


def run():
    env=os.getenv('ZSJOS_AGENT_ENV')
    if env is None:env=Path('/etc/zsjos/agent-environment').read_text().strip()
    x.h.require(env=='test','Unexpected environment')
    os.umask(0o077);ROOT.mkdir(mode=0o700,parents=True,exist_ok=True)
    if (ROOT/'delivery.json').exists():
        old=x.h.read(ROOT/'delivery.json')
        x.h.require(x.sha(OUTPUT)==old['output_sha256'],'Review file changed; preserve user decisions')
        print('Replay: no changes');return
    hashes={p.name:x.sha(p) for p in (OUTPUT,x.MASTER,x.EXCEPTIONS,PENDING)}
    backup=ROOT/OUTPUT.name
    x.h.require(not backup.exists(),'Backup exists; review required')
    shutil.copy2(OUTPUT,backup);backup.chmod(0o600)
    book=openpyxl.load_workbook(OUTPUT);baseline=openpyxl.load_workbook(OUTPUT)
    old=list(book['逐项对比'].values)[1:];x.h.require(len(old)==22,'Comparison count differs')
    index=book.sheetnames.index('逐项对比');book.remove(book['逐项对比'])
    ws=book.create_sheet('逐行对比',index)
    headers=['序号','数据来源','姓名','手机号','微信号','课程','课程规格','班级','班主任（原表规划师）',
             '服务负责人','服务显示姓名','差异说明','您的处理意见','补充说明','总表工作表','总表行号',
             '服务状态（原值）','接收状态（原值）','交付阶段（原值）','教学是否关闭','服务创建时间']
    ws.append(headers);expected=[]
    full=book['原表完整信息']
    for i,r in enumerate(old,2):
        # The complete source sheet holds metadata4 + original name/mobile/wechat.
        x.h.require(full.cell(i,1).value==r[0],'Source sequence differs')
        original=[r[0],'原表数据',r[6],r[9],full.cell(i,7).value,r[12],r[14],r[16],r[18],None,None,
                  r[3],r[4],r[5],r[1],r[2],None,None,None,None,None]
        system=[r[0],'系统数据',r[7],r[10],r[11],r[13],r[15],r[17],r[19],r[20],r[8],
                r[3],None,None,r[1],r[2],r[26],r[27],r[28],r[29],r[30]]
        expected.extend([tuple(original),tuple(system)]);ws.append(original);ws.append(system)
    x.format_sheet(ws);ws.freeze_panes='E2';ws.sheet_view.zoomScale=85
    table=Table(displayName='AssignmentPaired',ref=ws.dimensions)
    table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=False);ws.add_table(table)
    widths=[8,13,17,18,20,20,48,43,29,29,17,24,32,32,22,12,19,19,19,17,24]
    for i,width in enumerate(widths,1):ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=width
    line=Side(style='thin',color='8397B0')
    for top in range(2,46,2):
        for row in (top,top+1):
            ws.row_dimensions[row].height=58
            for cell in ws[row]:
                cell.fill=PatternFill('solid',fgColor='F2F6FC' if row==top else 'FFFFFF')
                cell.alignment=Alignment(vertical='center',wrap_text=True)
                cell.border=Border(bottom=line if row==top+1 else Side())
            ws.cell(row,2).font=Font(bold=True,color='264478' if row==top else '548235')
            for col in (4,5):ws.cell(row,col).number_format='@'
        for col in range(3,10):
            if ws.cell(top,col).value!=ws.cell(top+1,col).value:
                for row in (top,top+1):
                    cell=ws.cell(row,col);cell.fill=PatternFill('solid',fgColor='FFF2CC');cell.font=Font(color='9C0006')
        for col in (13,14):ws.cell(top,col).fill=PatternFill('solid',fgColor='E2F0D9')
    summary=book['说明与汇总']
    for row in summary:
        if row[0].value=='阅读顺序':row[1].value='逐行对比每组两行：先原表数据，后系统数据；相同字段同列，黄色为差异；原表完整信息保留全部44列'
        if row[0].value=='确认方式':row[1].value='在每组原表数据行的“您的处理意见”填写决定，补充说明栏填写备注'
    summary.append(['未对应字段','原表没有服务负责人、服务显示姓名和系统状态，留空；仅对双方共有字段标黄'])
    summary.auto_filter.ref=summary.dimensions
    book.active=index
    temp=OUTPUT.with_suffix('.tmp.xlsx');book.save(temp)
    with zipfile.ZipFile(temp) as z:x.h.require(z.testzip() is None,'Invalid workbook ZIP')
    saved=openpyxl.load_workbook(temp)
    x.h.require(list(saved['逐行对比'].values)[1:]==expected,'Paired fields differ')
    x.h.require(saved.active.title=='逐行对比' and saved['逐行对比'].max_row==45,'Opening sheet or row count differs')
    for title in ('原表完整信息','原异常记录'):
        for row in baseline[title]:
            for a in row:
                b=saved[title].cell(a.row,a.column)
                x.h.require((a.value,a.data_type,a.number_format)==(b.value,b.data_type,b.number_format),'Protected source cells changed')
    x.h.require(all(x.sha(p)==hashes[p.name] for p in (OUTPUT,x.MASTER,x.EXCEPTIONS,PENDING)),'Input changed concurrently')
    temp.replace(OUTPUT)
    x.h.save(ROOT/'delivery.json',dict(output_sha256=x.sha(OUTPUT),source_hashes=hashes,pairs=22,data_rows=44,database_access=False))
    print('PASS:22 original/system pairs,44 rows; decisions and44source columns preserved; no database access')


if __name__=='__main__':run()
