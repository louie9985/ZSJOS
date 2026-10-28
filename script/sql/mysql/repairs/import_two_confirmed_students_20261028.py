#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Import only two explicitly confirmed source rows, preserving other variants.

Uses the established historical import transaction and private recovery evidence.
No personal data or session credentials are emitted in logs or result reports.
"""
import argparse
from copy import deepcopy
import json
import os
from pathlib import Path
import shutil
import zipfile

import openpyxl
import requests
import import_education_exceptions_20261027 as x

ROOT=Path('/opt/zsjos-runtime/backups/education-two-confirmed-20261028')
MARKER='edu-two-confirmed-20261028'
PENDING=x.ROOT/'output/教务学员_多条报名待确认_20261027.xlsx'
REPORT=x.ROOT/'output/教务学员_两条确认导入结果_20261028.xlsx'
TARGETS={('人社公共营养师',40),('执业中西药师',40)}


def session():
    active=x.h.query('SELECT COUNT(*) n FROM system_oauth2_access_token WHERE tenant_id=1 AND user_id=1 '
                     'AND user_type=2 AND deleted=0 AND expires_time>NOW()')[0]['n']
    if not active:
        tokens=x.h.query('SELECT refresh_token,client_id FROM system_oauth2_refresh_token WHERE tenant_id=1 '
                        'AND user_id=1 AND user_type=2 AND deleted=0 AND expires_time>NOW() ORDER BY expires_time DESC LIMIT 1')
        x.h.require(tokens,'Existing admin session unavailable')
        response=requests.post(x.h.BASE+'/system/auth/refresh-token',headers={'tenant-id':'1'},
            params={'refreshToken':tokens[0]['refresh_token'],'clientId':tokens[0]['client_id']},timeout=30)
        x.h.require(response.status_code==200 and response.json().get('code')==0,'Existing session refresh rejected')
    return x.h.auth()


def prepare():
    x.h.require(not (ROOT/'plan.json').exists(),'Frozen plan exists')
    _,records=x.source_rows(x.EXCEPTIONS)
    selected=[r for r in records if (r['sheet'],r['row']) in TARGETS]
    x.h.require(len(records)==156 and len(selected)==2,'Source scope changed')
    source_hashes={p.name:x.sha(p) for p in (x.MASTER,x.EXCEPTIONS,PENDING)}
    for p in (x.MASTER,x.EXCEPTIONS,PENDING):
        target=ROOT/p.name
        x.h.require(not target.exists(),'Existing backup requires review')
        shutil.copy2(p,target);target.chmod(0o600)
    master=openpyxl.load_workbook(x.MASTER,data_only=True)
    pending=openpyxl.load_workbook(PENDING,data_only=True)['待确认报名']
    headings={c.value:c.column for c in pending[1]}
    snapshot=x.h.snapshot();s=session();options=x.h.api(s,'GET','/zsjos/student-entry/options')
    for r in selected:
        raw=list(master[r['sheet']+'-学员导入填写'].values)[r['row']-1]
        x.h.require([x.norm(v) for v in raw[:7]]==[r['name'],r['mobile'],None,r['product'],r['sku'],r['class_name'],r['planner']],
                    'Master row differs from confirmation')
        matched=[row for row in list(pending.values)[1:] if (row[headings['总表工作表']-1],row[headings['总表行号']-1])==(r['sheet'],r['row'])]
        x.h.require(len(matched)==1 and list(matched[0][5:])==list(raw),'Full-field pending row differs')
        products=[p for p in snapshot['products'] if p['name']==r['product'] and p['status']==0]
        x.h.require(len(products)==1,'Product unavailable')
        r['product_id'],r['product_ref']=products[0]['id'],products[0]['product_ref']
        skus=[k for k in snapshot['skus'] if k['spu_id']==r['product_id'] and k['sku_name']==r['sku'] and k['status']==0]
        x.h.require(len(skus)==1 and skus[0]['sku_ref'] in {k['skuRef'] for k in options['courses']},'SKU unavailable')
        r['sku_id'],r['skuRef']=skus[0]['id'],skus[0]['sku_ref']
        owners=[u for u in snapshot['users'] if u['nickname']==r['planner'] and u['status']==0]
        x.h.require(len(owners)==1,'Planner ambiguous')
        r['owner_id']=owners[0]['id']
        classes=[c for c in snapshot['classes'] if c['class_name']==r['class_name'] and c['homeroom_user_id']==r['owner_id'] and c['status']=='SERVING']
        x.h.require(len(classes)==1 and classes[0]['id'] in {c['id'] for c in options['classes']},'Class unavailable')
        r['class_id']=classes[0]['id']
        people=[p for p in snapshot['people'] if p['mobile']==r['mobile']]
        x.h.require(len(people)<=1 and all(p['name']==r['name'] for p in people),'Mobile/name changed')
        if people:r['person_id']=people[0]['id']
        check=x.h.api(s,'POST','/zsjos/student-entry/check',{'mobile':r['mobile'],'wechatId':None})
        x.h.require(check['status'] in ('new','reuse') and check.get('personId')==r.get('person_id'),'Live identity conflict')
        r['same_services']=[];r['other_services']=[]
        for rel in snapshot['services']:
            if rel['person_id']!=r.get('person_id'):continue
            snap=json.loads(rel['service_snapshot'] or '{}')
            x.h.require(snap and snap.get('productRef')!=r['product_ref'] and snap.get('name',snap.get('productName'))!=r['product'],
                        'Same course already exists; refuse duplicate')
            r['other_services'].append(rel['id'])
        r['issues']=[]
    x.h.require(all(x.sha(p)==source_hashes[p.name] for p in (x.MASTER,x.EXCEPTIONS,PENDING)),'Source changed during preflight')
    x.prepare(selected,snapshot,sync=False,missing_sku_count=0,extra=dict(mobile_only=True,
        fixture_failure_points=[1,2,5,10,15],wrapper_sha256=x.sha(Path(__file__)),pending_sha256=x.sha(PENDING),source_count=156))


def finish(plan):
    c=x.h.connection()
    try:
        with c.cursor() as q:
            q.execute('START TRANSACTION WITH CONSISTENT SNAPSHOT, READ ONLY')
            # The shared daily counter can legitimately advance after commit. The
            # transaction and immediate replay already verified our exact increment.
            # Exclude only later counter advancement from this read-only delivery check.
            committed=x.loads((ROOT/'after.json').read_text())
            verification_plan=deepcopy(plan)
            current_state=x.capture(q,plan)
            current=current_state[x.COUNTER]
            prior=committed[x.COUNTER]
            x.h.require(len(current)==len(prior)==1,'Daily counter scope changed')
            x.h.require(all(current[0][k]==v for k,v in prior[0].items() if k not in ('current_value','update_time')),
                        'Unrelated daily counter metadata changed')
            x.h.require(current[0]['update_time']>=prior[0]['update_time'],'Daily counter time moved backwards')
            new_people=len({r['mobile'] for r in plan['ready'] if not r.get('person_id')})
            x.h.require(prior[0]['current_value']==(plan['before'][x.COUNTER][0]['current_value']+new_people-1)%9999+1,
                        'Committed counter increment differs')
            verification_plan['before'][x.COUNTER][0]['current_value']=(current[0]['current_value']-new_people-1)%9999+1
            # No classes were mutated by this batch. Preserve exact protection for
            # our two destinations while accepting independent edits to other classes.
            target_classes={r['class_id'] for r in plan['ready']}
            observed_classes={r['id']:r for r in current_state[x.C]}
            x.h.require(set(observed_classes)=={r['id'] for r in committed[x.C]},'Class set changed; review required')
            external_classes=[]
            for index,old in enumerate(verification_plan['before'][x.C]):
                if old['id'] not in target_classes:
                    observed=observed_classes[old['id']]
                    if x.h.encoded(old)!=x.h.encoded(observed):
                        external_classes.append(dict(before=old,observed=observed))
                    verification_plan['before'][x.C][index]=observed
            after,result=x.verify(q,verification_plan)
            x.h.save(ROOT/'counter-delivery-check.json',dict(committed=prior,current=current,read_only=True))
            x.h.save(ROOT/'unrelated-class-observations.json',external_classes)
    finally:c.rollback();c.close()
    x.h.save(ROOT/'independent-after.json',after)
    s=session()
    for row in result:
        rid=row['service_id'];context=x.h.api(s,'GET',f'/zsjos/student/service/{rid}/contact-context')
        stages={a['code']:a for a in context['deliveryStages']}
        x.h.require(context['acceptanceStatus']=='accepted' and context['deliveryStage']=='supervision'
            and context['currentTask'] is None and stages['first_contact']['status']==stages['study_plan']['status']=='done'
            and all(a['completed'] for a in context['firstContactProgress']['items']),'API stage incomplete')
        history=x.h.api(s,'GET',f'/zsjos/student/service/{rid}/contact-records?pageNo=1&pageSize=100')
        x.h.require(history['total']==2 and {r['contactType'] for r in history['list']}=={t[0] for t in x.TYPES}
            and all(r['successful'] and r['remark']==x.REMARK for r in history['list']),'API completion record differs')
    x.h.require(x.sha(x.MASTER)==plan['master_sha256'],'Master changed')
    if (ROOT/'files.json').exists():
        old=x.h.read(ROOT/'files.json')
        x.h.require(all(x.sha(p)==old[p.name] for p in (x.EXCEPTIONS,PENDING,REPORT)),'Delivered file changed')
        print('Verified two services; file replay unchanged');return
    x.h.require(x.sha(x.EXCEPTIONS)==plan['exception_sha256'] and x.sha(PENDING)==plan['pending_sha256'],'Review workbook changed')
    exceptions=openpyxl.load_workbook(x.EXCEPTIONS);old_exc=openpyxl.load_workbook(x.EXCEPTIONS)
    sheet=exceptions['异常与处理总表'];removed=[]
    for row in range(sheet.max_row,1,-1):
        if (sheet.cell(row,8).value,sheet.cell(row,9).value) in TARGETS:sheet.delete_rows(row);removed.append(row)
    x.h.require(len(removed)==2 and sheet.max_row==155,'Exception partition differs')
    sheet.tables['ExceptionTable'].ref='A1:U155';sheet.auto_filter.ref='A1:U155'
    comparison=exceptions['关联记录对比']
    for row in range(comparison.max_row,1,-1):
        if (comparison.cell(row,4).value,comparison.cell(row,5).value) in TARGETS:comparison.delete_rows(row)
    comparison.tables['ComparisonTable'].ref=f'A1:K{comparison.max_row}';comparison.auto_filter.ref=comparison.tables['ComparisonTable'].ref
    summary=exceptions['汇总']
    for row in range(2,summary.max_row+1):
        if summary.cell(row,1).value=='当前保留异常':summary.cell(row,2,154)
        if summary.cell(row,1).value=='多条报名待确认':summary.cell(row,2,11);summary.cell(row,3,'原6组中2条已确认导入；仍有11条待确认')
    summary.append(['2026-10-28两条确认导入',2,'已进入常规督学；成功行已从异常主表及待确认表移出'])
    summary.tables['SummaryTable'].ref=f'A1:C{summary.max_row}';summary.auto_filter.ref=summary.tables['SummaryTable'].ref
    pending=openpyxl.load_workbook(PENDING);ps=pending['待确认报名']
    pending_before=list(ps.values)
    for row in range(ps.max_row,1,-1):
        if (ps.cell(row,2).value,ps.cell(row,3).value) in TARGETS:ps.delete_rows(row)
    ps.auto_filter.ref=ps.dimensions
    x.h.require(ps.max_row==12 and ps.max_column==49,'Pending full-field scope differs')
    # Keep source/difference/decision values intact; separate sheet explains resolved siblings.
    title='已确认导入'
    x.h.require(title not in pending.sheetnames,'Unexpected existing result sheet')
    resolved=pending.create_sheet(title)
    resolved.append(['总表工作表','总表行号','确认结果','备注'])
    for r in plan['ready']:resolved.append([r['sheet'],r['row'],'用户已确认，已导入并进入常规督学',x.REMARK])
    x.format_sheet(resolved)
    report=openpyxl.Workbook();report.active.title='导入结果'
    report.active.append(['总表工作表','总表行号','课程','课程规格','班级','处理结果','备注'])
    for r in plan['ready']:report.active.append([r['sheet'],r['row'],r['product'],r['sku'],r['class_name'],'已导入；首联、学习计划完成，常规督学',x.REMARK])
    x.format_sheet(report.active)
    expected_exc=[row for row in list(old_exc['异常与处理总表'].values)[1:] if (row[7],row[8]) not in TARGETS]
    expected_pending=[row for row in pending_before[1:] if (row[1],row[2]) not in TARGETS]
    for path,book in ((x.EXCEPTIONS,exceptions),(PENDING,pending),(REPORT,report)):
        temp=path.with_suffix('.tmp.xlsx');book.save(temp)
        with zipfile.ZipFile(temp) as z:x.h.require(z.testzip() is None,'Invalid workbook ZIP')
    saved_exc=openpyxl.load_workbook(x.EXCEPTIONS.with_suffix('.tmp.xlsx'))
    saved_pending=openpyxl.load_workbook(PENDING.with_suffix('.tmp.xlsx'))
    x.h.require(list(saved_exc['异常与处理总表'].values)[1:]==expected_exc,'Other exception cells changed')
    x.h.require(list(saved_pending['待确认报名'].values)[1:]==expected_pending,'Other full-field pending rows changed')
    for path in (x.EXCEPTIONS,PENDING,REPORT):path.with_suffix('.tmp.xlsx').replace(path)
    x.h.save(ROOT/'files.json',{p.name:x.sha(p) for p in (x.EXCEPTIONS,PENDING,REPORT)})
    x.h.save(ROOT/'verification.json',dict(imported=2,people=2,completed_tasks=4,records=4,remark=x.REMARK,
        remaining_exceptions=154,pending_rows=11,api_contexts=2,api_histories=2,protected_fields=True))
    print('PASS: 2 services, 4 completed tasks/records; 154 exceptions and 11 full-field pending rows remain')


def main():
    env=os.getenv('ZSJOS_AGENT_ENV')
    if env is None:env=Path('/etc/zsjos/agent-environment').read_text().strip()
    x.h.require(env=='test','Test-only execution')
    os.umask(0o077);ROOT.mkdir(mode=0o700,parents=True,exist_ok=True)
    x.EVIDENCE=ROOT;x.MARKER=MARKER
    parser=argparse.ArgumentParser();parser.add_argument('mode',choices=['prepare','fixture','apply','finish']);parser.add_argument('--sha256');args=parser.parse_args()
    if args.mode=='prepare':prepare();return
    plan=x.loads((ROOT/'plan.json').read_text())
    x.h.require(args.sha256==x.h.digest(plan),'Plan digest differs')
    if args.mode!='finish':
        x.h.require(plan['wrapper_sha256']==x.sha(Path(__file__)),'Mutation wrapper changed; do not replay')
    if args.mode=='finish':finish(plan)
    else:getattr(x,args.mode)(plan)


if __name__=='__main__':main()
