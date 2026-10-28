#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Reuse user-confirmed unique WeChat identities in the test historical import.

Preserve contacts and existing services; retain different-SKU enrollment decisions.
Original personal data is confined to private evidence and requested workbooks.
"""
import argparse
from collections import defaultdict
import json
import os
from pathlib import Path
import shutil
import zipfile

import openpyxl
import import_education_exceptions_20261027 as x
from import_two_confirmed_students_20261028 import session
from import_education_identity_20261027 import clean_name
from archive_matching_education_students_20261028 import clone_cell

ROOT=Path('/opt/zsjos-runtime/backups/education-unique-wechat-20261028')
MARKER='edu-unique-wechat-20261028'
REPORT=x.ROOT/'output/教务学员_唯一微信档案处理结果_20261028.xlsx'
PENDING=x.ROOT/'output/教务学员_多条报名待确认_20261027.xlsx'


def prepare():
    x.h.require(not (ROOT/'plan.json').exists(),'Frozen plan exists')
    _,allrows=x.source_rows(x.EXCEPTIONS);x.h.require(len(allrows)==57,'Exception scope changed')
    source_hashes={p.name:x.sha(p) for p in (x.MASTER,x.EXCEPTIONS,PENDING)}
    st=x.h.snapshot();s=session();options=x.h.api(s,'GET','/zsjos/student-entry/options')
    master=openpyxl.load_workbook(x.MASTER,data_only=True);names=defaultdict(set)
    for sh in master.worksheets[:8]:
        for raw in list(sh.values)[1:]:
            if raw[0] and raw[1]:names[x.norm(raw[1])].add(clean_name(raw[0]))
    selected=[];updates={};protected=[]
    for r in allrows:
        if any(p['mobile']==r['mobile'] for p in st['people']):continue
        wechat=[p for p in st['people'] if p['wechat_id']==r['mobile']]
        claims=[c for c in st['claims'] if c['contact_value']==r['mobile']]
        ids={p['id'] for p in wechat}|{c['person_id'] for c in claims if c['person_id']}
        if not ids:continue
        x.h.require(len(ids)==1 and len(wechat)==1,'Nonunique contact identity')
        person=wechat[0]
        x.h.require(not person['mobile'] and all(c['person_id']==person['id'] and c['reservation_key'] is None for c in claims),
                    'Unexpected mobile/claim scope')
        x.h.require(names[r['mobile']]=={r['name']},'Source name disagreement')
        r['person_id']=person['id'];r['candidate_person_ids']=[person['id']]
        raw=list(master[r['sheet']+'-学员导入填写'].values)[r['row']-1]
        x.h.require([x.norm(raw[k]) for k in (0,1,3,4,5,6)]==[r['name'],r['mobile'],r['product'],r['sku'],r['class_name'],r['planner']],
                    'Master/source differs')
        products=[p for p in st['products'] if p['name']==r['product'] and p['status']==0]
        x.h.require(len(products)==1,'Product unavailable')
        r['product_id'],r['product_ref']=products[0]['id'],products[0]['product_ref']
        skus=[k for k in st['skus'] if k['spu_id']==r['product_id'] and k['sku_name']==r['sku'] and k['status']==0]
        x.h.require(len(skus)==1 and skus[0]['sku_ref'] in {k['skuRef'] for k in options['courses']},'SKU unavailable')
        r['sku_id'],r['skuRef']=skus[0]['id'],skus[0]['sku_ref']
        users=[u for u in st['users'] if u['nickname']==r['planner'] and u['status']==0]
        x.h.require(len(users)==1,'Planner unavailable');r['owner_id']=users[0]['id']
        classes=[c for c in st['classes'] if c['class_name']==r['class_name'] and c['homeroom_user_id']==r['owner_id'] and c['status']=='SERVING']
        x.h.require(len(classes)==1 and classes[0]['id'] in {c['id'] for c in options['classes']},'Class unavailable')
        r['class_id']=classes[0]['id']
        check=x.h.api(s,'POST','/zsjos/student-entry/check',{'mobile':r['mobile'],'wechatId':None})
        x.h.require(check['status']=='reuse' and check['personId']==person['id'],'Identity API disagrees')
        services=[rel for rel in st['services'] if rel['person_id']==person['id']]
        same=[rel for rel in services if json.loads(rel['service_snapshot'] or '{}').get('productRef')==r['product_ref']
              or json.loads(rel['service_snapshot'] or '{}').get('name')==r['product']]
        exact=[rel for rel in same if json.loads(rel['service_snapshot'])['skuName']==r['sku'] and rel['class_id']==r['class_id']
               and rel['owner_user_id']==r['owner_id'] and rel['student_name_snapshot']==r['name'] and rel['status']=='active'
               and rel['teaching_delivery'] in (1,b'\1') and rel['teaching_closed'] in (0,b'\0')]
        r['same_services']=[rel['id'] for rel in same];r['other_services']=[rel['id'] for rel in services if rel not in same]
        r['existing_skus']=[json.loads(rel['service_snapshot']).get('skuName') for rel in same]
        r['outcome']='existing' if exact else ('hold' if same else 'new')
        r['issues']=[] if not same else ['已有完全一致服务，归档为已导入' if exact else '档案已确认复用；已有同课程不同规格，仍需确认']
        r['exact_ids']=[rel['id'] for rel in exact]
        if person['name']!=r['name']:
            change=dict(id=person['id'],name=r['name'],sheet=r['sheet'],row=r['row'])
            x.h.require(person['id'] not in updates or updates[person['id']]['name']==r['name'],'Conflicting name changes')
            updates[person['id']]=change
        protected.extend(services);selected.append(r)
    x.h.require(len(selected)==17 and sum(r['outcome']=='new' for r in selected)==8
        and sum(r['outcome']=='existing' for r in selected)==8 and len(updates)==14,'Reviewed scope changed')
    x.h.require(all(x.sha(p)==source_hashes[p.name] for p in (x.MASTER,x.EXCEPTIONS,PENDING)),'Source changed')
    for p in (x.MASTER,x.EXCEPTIONS,PENDING):
        dst=ROOT/p.name;x.h.require(not dst.exists(),'Backup already exists');shutil.copy2(p,dst);dst.chmod(0o600)
    x.prepare(selected,st,sync=False,missing_sku_count=0,extra=dict(mobile_only=False,name_updates=list(updates.values()),
        service_name_updates=[],fixture_failure_points=[1,15,30,60,90],wrapper_sha256=x.sha(Path(__file__)),
        pending_sha256=x.sha(PENDING),protected_existing_services=protected,source_count=57))


def verify_database(plan):
    c=x.h.connection()
    try:
        with c.cursor() as q:
            q.execute('START TRANSACTION WITH CONSISTENT SNAPSHOT, READ ONLY');after,result=x.verify(q,plan)
            for old in plan['protected_existing_services']:
                actual=x.rows(q,f'SELECT * FROM {x.S} WHERE tenant_id=1 AND id=%s',(old['id'],))
                x.h.require(len(actual)==1 and x.h.encoded(actual[0])==x.h.encoded(old),'Existing service changed')
            for r in plan['records']:
                p=x.rows(q,f'SELECT * FROM {x.P} WHERE tenant_id=1 AND id=%s AND deleted=0',(r['person_id'],))[0]
                x.h.require(p['name']==r['name'] and not p['mobile'] and p['wechat_id']==r['mobile'],'Reused person differs')
    finally:c.rollback();c.close()
    x.h.save(ROOT/'independent-after.json',after)
    s=session()
    for row in result:
        rid=row['service_id'];ctx=x.h.api(s,'GET',f'/zsjos/student/service/{rid}/contact-context')
        stages={a['code']:a for a in ctx['deliveryStages']}
        x.h.require(ctx['acceptanceStatus']=='accepted' and ctx['deliveryStage']=='supervision' and ctx['currentTask'] is None
            and stages['first_contact']['status']==stages['study_plan']['status']=='done'
            and all(a['completed'] for a in ctx['firstContactProgress']['items']),'API stage differs')
        history=x.h.api(s,'GET',f'/zsjos/student/service/{rid}/contact-records?pageNo=1&pageSize=100')
        x.h.require(history['total']==2 and all(a['successful'] and a['remark']==x.REMARK for a in history['list']),'API records differ')
    x.h.save(ROOT/'api-verification.json',dict(contexts=8,histories=8,passed=True))


def finish(plan):
    paths=(x.MASTER,x.EXCEPTIONS,PENDING,REPORT)
    if (ROOT/'files.json').exists():
        hashes=x.h.read(ROOT/'files.json');x.h.require(all(x.sha(p)==hashes[p.name] for p in paths),'Delivered file changed')
        print('File replay: no changes');return
    verify_database(plan)
    x.h.require(x.sha(x.MASTER)==plan['master_sha256'] and x.sha(x.EXCEPTIONS)==plan['exception_sha256']
        and x.sha(PENDING)==plan['pending_sha256'],'Workbook changed')
    records=plan['records'];done={(r['sheet'],r['row']) for r in records if r['outcome']!='hold'}
    held=[r for r in records if r['outcome']=='hold'];x.h.require(len(done)==16 and len(held)==1,'Outcome partition changed')
    master=openpyxl.load_workbook(x.MASTER);exc=openpyxl.load_workbook(x.EXCEPTIONS);original=openpyxl.load_workbook(x.EXCEPTIONS)
    report=openpyxl.Workbook();report.active.title='汇总'
    for row in [('处理项','数量','说明'),('复用档案新增服务',8,'完成首联、学习计划，进入常规督学'),('已有一致服务',8,'归档为已导入，现有服务保持'),
                ('规格差异待确认',1,'档案已确认；是否新增另一规格仍待确认'),('姓名按总表修正',14,'不更改手机号或微信号'),('异常表剩余',41,'移出16条已处理记录')]:report.active.append(row)
    full=report.create_sheet('完整原表与处理结果');full.append(['总表工作表','总表行号','处理结果','系统已有规格','说明']+[c.value for c in master[records[0]['sheet']+'-学员导入填写'][1]])
    for idx,r in enumerate(records,2):
        outcome={'new':'复用档案新增服务','existing':'已有一致服务，已导入','hold':'档案复用已确认，规格差异待确认'}[r['outcome']]
        for col,v in enumerate([r['sheet'],r['row'],outcome,'；'.join(r['existing_skus']),x.REMARK if r['outcome']=='new' else '唯一微信档案；原联系方式保留'],1):full.cell(idx,col,v)
        source=master[r['sheet']+'-学员导入填写']
        for col in range(1,45):clone_cell(source.cell(r['row'],col),full.cell(idx,col+5))
    for sh in report:x.format_sheet(sh)
    es=exc['异常与处理总表'];r=held[0];row=r['exception_row']
    es.cell(row,3,'已有课程，规格待确认');es.cell(row,4,'唯一微信档案已确认复用；已有同课程不同规格')
    es.cell(row,6,'表内规格与系统已有规格不同，确认是否保留新增服务')
    es.cell(row,7,str(es.cell(row,7).value or '')+'\n2026-10-28：已确认复用唯一微信档案；姓名按总表修正。表内：'+r['sku']+'；系统：'+'；'.join(r['existing_skus']))
    for row in range(es.max_row,1,-1):
        if (es.cell(row,8).value,es.cell(row,9).value) in done:es.delete_rows(row)
    x.h.require(es.max_row==42,'Exception count differs');es.tables['ExceptionTable'].ref='A1:U42';es.auto_filter.ref='A1:U42'
    cs=exc['关联记录对比']
    for row in range(cs.max_row,1,-1):
        if (cs.cell(row,4).value,cs.cell(row,5).value) in done:cs.delete_rows(row)
    cs.tables['ComparisonTable'].ref=f'A1:K{cs.max_row}';cs.auto_filter.ref=cs.tables['ComparisonTable'].ref
    summary=exc['汇总']
    for row in range(2,summary.max_row+1):
        if summary.cell(row,1).value=='当前保留异常':summary.cell(row,2,41)
        if summary.cell(row,1).value=='当前异常构成':summary.cell(row,2,41);summary.cell(row,3,'唯一微信档案问题已解除；剩余课程差异及多档案等异常')
    summary.append(['2026-10-28唯一微信档案复用',17,'8条新增、8条已有一致服务归档、1条规格待确认；姓名修正14个档案'])
    summary.tables['SummaryTable'].ref=f'A1:C{summary.max_row}';summary.auto_filter.ref=summary.tables['SummaryTable'].ref
    expected=list(es.values)
    for path,book in ((x.EXCEPTIONS,exc),(REPORT,report)):
        temp=path.with_suffix('.tmp.xlsx');book.save(temp)
        with zipfile.ZipFile(temp) as z:x.h.require(z.testzip() is None,'Invalid workbook ZIP')
    saved=openpyxl.load_workbook(x.EXCEPTIONS.with_suffix('.tmp.xlsx'));saved_report=openpyxl.load_workbook(REPORT.with_suffix('.tmp.xlsx'))
    x.h.require(list(saved['异常与处理总表'].values)==expected,'Saved exceptions differ')
    old={(a[7],a[8]):a for a in list(original['异常与处理总表'].values)[1:]}
    for a in list(saved['异常与处理总表'].values)[1:]:
        key=(a[7],a[8]);allow={2,3,5,6} if key==(held[0]['sheet'],held[0]['row']) else set()
        x.h.require(all(v==old[key][i] for i,v in enumerate(a) if i not in allow),'Unrelated exception cells changed')
    x.h.require(list(saved['聊天处理依据'].values)==list(original['聊天处理依据'].values),'History changed')
    for row in list(saved_report['完整原表与处理结果'].rows)[1:]:
        source=master[row[0].value+'-学员导入填写']
        for col in range(1,45):
            a,b=source.cell(row[1].value,col),row[col+4]
            x.h.require((a.value,a.data_type,a.number_format)==(b.value,b.data_type,b.number_format),'Archive field differs')
    x.h.require(x.sha(x.MASTER)==plan['master_sha256'] and x.sha(x.EXCEPTIONS)==plan['exception_sha256'] and x.sha(PENDING)==plan['pending_sha256'],'Source changed before replacement')
    for path in (x.EXCEPTIONS,REPORT):path.with_suffix('.tmp.xlsx').replace(path)
    x.h.save(ROOT/'files.json',{p.name:x.sha(p) for p in paths})
    print('PASS: reused 17 unique identities; 8 new services, 8 existing archived, 1 SKU hold; 41 exceptions remain')


def main():
    env=os.getenv('ZSJOS_AGENT_ENV')
    if env is None:env=Path('/etc/zsjos/agent-environment').read_text().strip()
    x.h.require(env=='test','Test only');os.umask(0o077);ROOT.mkdir(mode=0o700,parents=True,exist_ok=True)
    x.EVIDENCE=ROOT;x.MARKER=MARKER
    parser=argparse.ArgumentParser();parser.add_argument('mode',choices=['prepare','fixture','apply','finish']);parser.add_argument('--sha256');a=parser.parse_args()
    if a.mode=='prepare':prepare();return
    plan=x.loads((ROOT/'plan.json').read_text());x.h.require(a.sha256==x.h.digest(plan),'Plan digest differs')
    if a.mode!='finish':x.h.require(plan['wrapper_sha256']==x.sha(Path(__file__)),'Mutation wrapper changed')
    if a.mode=='finish':finish(plan)
    else:getattr(x,a.mode)(plan)


if __name__=='__main__':main()
