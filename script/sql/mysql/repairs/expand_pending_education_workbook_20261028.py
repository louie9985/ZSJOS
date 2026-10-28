#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Restore all master fields in the user's local pending-review workbook.

Contains personal information only in the explicitly requested workbook and private
backup. Never log cell values. No database access or import operations.
"""
from copy import copy
import hashlib
import os
from pathlib import Path
import shutil
import zipfile

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path('/opt/zsjos')
SOURCE = ROOT / 'output/教务学员导入总表.xlsx'
TARGET = ROOT / 'output/教务学员_多条报名待确认_20261027.xlsx'
BACKUP = Path('/opt/zsjos-runtime/backups/education-pending-full-20261028')
META = ['组号', '总表工作表', '总表行号', '差异说明', '您的处理意见']


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def run():
    env = os.getenv('ZSJOS_AGENT_ENV')
    if env is None:
        env = Path('/etc/zsjos/agent-environment').read_text().strip()
    assert env == 'test', 'Unexpected environment'
    os.umask(0o077)
    source_hash, target_hash = sha(SOURCE), sha(TARGET)
    source = openpyxl.load_workbook(SOURCE)
    current = openpyxl.load_workbook(TARGET)
    old = current['待确认报名']
    headers = {cell.value: cell.column for cell in old[1]}
    records = [[old.cell(row, headers[key]).value for key in META] for row in range(2, old.max_row + 1)]
    assert len(records) == 13 and {r[0] for r in records} == set(range(1, 7)), 'Review scope changed'
    assert len({(r[1], r[2]) for r in records}) == 13, 'Repeated source reference'
    BACKUP.mkdir(mode=0o700, parents=True, exist_ok=True)
    backup = BACKUP / (target_hash[:16] + '.xlsx')
    if not backup.exists():
        shutil.copy2(TARGET, backup)
        backup.chmod(0o600)
    # Keep any other user-maintained sheets and the decision column intact.
    position = current.index(old)
    current.remove(old)
    sheet = current.create_sheet('待确认报名', position)
    source_headers = None
    for out_row, record in enumerate(records, 2):
        source_sheet = source[record[1] + '-学员导入填写']
        assert source_sheet.max_column == 44, 'Source column count changed'
        names = [c.value for c in source_sheet[1]]
        if source_headers is None:
            source_headers = names
            sheet.append(META + names)
        assert names == source_headers, 'Source headers differ'
        for col, value in enumerate(record, 1):
            sheet.cell(out_row, col, value)
        for col in range(1, 45):
            original = source_sheet.cell(record[2], col)
            cell = sheet.cell(out_row, col + 5, original.value)
            cell.data_type = original.data_type
            # Style IDs belong to their workbook; register components in the target.
            cell.font = copy(original.font)
            cell.fill = copy(original.fill)
            cell.border = copy(original.border)
            cell.protection = copy(original.protection)
            cell.number_format = original.number_format
            if original.comment:
                cell.comment = copy(original.comment)
            if original.hyperlink:
                cell.hyperlink = copy(original.hyperlink)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        sheet.row_dimensions[out_row].height = 62
        fill = PatternFill('solid', fgColor='EAF2F8' if record[0] % 2 else 'F4F4F4')
        for cell in sheet[out_row][:5]:
            cell.fill = fill
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        sheet.cell(out_row, 5).fill = PatternFill('solid', fgColor='FFF2CC')
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='264478')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    widths = {1:8, 2:22, 3:12, 4:52, 5:35, 6:18, 7:19, 8:20, 9:20, 10:46, 11:40, 12:30}
    for col in range(1, 50):
        sheet.column_dimensions[get_column_letter(col)].width = widths.get(col, 22)
    sheet.row_dimensions[1].height = 34
    sheet.freeze_panes = 'I2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.zoomScale = 80
    sheet.print_title_rows = '1:1'
    temp = TARGET.with_suffix('.tmp.xlsx')
    current.save(temp)
    with zipfile.ZipFile(temp) as z:
        assert z.testzip() is None, 'Workbook archive invalid'
    saved = openpyxl.load_workbook(temp)['待确认报名']
    assert saved.max_row == 14 and saved.max_column == 49, 'Saved dimensions differ'
    assert [c.value for c in saved[1]] == META + source_headers, 'Headers differ'
    for row, record in enumerate(records, 2):
        assert [saved.cell(row, col).value for col in range(1, 6)] == record, 'Review metadata changed'
        source_sheet = source[record[1] + '-学员导入填写']
        for col in range(1, 45):
            a, b = source_sheet.cell(record[2], col), saved.cell(row, col + 5)
            assert a.value == b.value and a.data_type == b.data_type and a.number_format == b.number_format, 'Original field differs'
    assert sha(SOURCE) == source_hash and sha(TARGET) == target_hash, 'Input changed concurrently'
    temp.replace(TARGET)
    print('PASS: 6 groups, 13 rows, all 44 original fields; review entries preserved; source unchanged.')


if __name__ == '__main__':
    run()
